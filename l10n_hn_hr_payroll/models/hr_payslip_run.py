# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik

import logging
from calendar import monthrange
from collections import defaultdict
from datetime import datetime, time

import pytz

from odoo import api, fields, models, _
from odoo.exceptions import UserError
from odoo.fields import Domain

from odoo.addons.gn_hr_payroll_employee_extras.models.payroll_biweekly_utils import (
    biweekly_end_from_start,
    biweekly_start_from_today,
    period_number_from_day,
)

_logger = logging.getLogger(__name__)

# Estructuras HN donde no se bloquea generate_payslips por _check_if_error en hr.work.entry
# (horas >24 h, solapamientos, ausencias fuera de calendario, etc.) — alineado con l10n_sv_hr_payroll.
_STRUCT_CODES_SKIP_WORK_ENTRY_VALIDATION = frozenset({
    "HN_BIWEEKLY",
    "HN_MONTHLY",
    "HN_13TH_MONTH",
    "HN_14TH_MONTH",
    "HN_VAC_ANN",
    "HN_VAC_DEC",
})

# Lote con tipo propio (13°/14°/vacaciones) pero contratos quincenal o mensual HN.
_STRUCT_CODES_SKIP_PAYROLL_STRUCTURE_TYPE_MATCH = frozenset({
    "HN_13TH_MONTH",
    "HN_14TH_MONTH",
    "HN_VAC_ANN",
    "HN_VAC_DEC",
})


class HrPayslipRun(models.Model):
    """Extend hr.payslip.run to improve schedule_pay auto-loading and bi-weekly period calculation"""
    _inherit = 'hr.payslip.run'

    def _hn_vacation_payroll_valid_version_ids(
        self,
        date_start,
        date_end,
        company_id,
        employee_ids,
        vacation_kind,
    ):
        version_domain = Domain([
            ('company_id', '=', company_id),
            ('employee_id', '!=', False),
            ('contract_date_start', '<=', date_end),
            '|',
            ('contract_date_end', '=', False),
            ('contract_date_end', '>=', date_start),
            ('date_version', '<=', date_end),
            ('structure_type_id', '!=', False),
        ])
        if employee_ids:
            version_domain &= Domain([('employee_id', 'in', employee_ids)])
        all_versions = self.env['hr.version']._read_group(
            domain=version_domain,
            groupby=['employee_id', 'date_version:day'],
            order='date_version:day DESC',
            aggregates=['id:recordset'],
        )
        all_employee_versions = defaultdict(list)
        for employee, _dv, version in all_versions:
            all_employee_versions[employee] += [*version]
        valid_versions = self.env['hr.version']
        for employee_versions in all_employee_versions.values():
            employee_valid_versions = self.env['hr.version']
            for i in range(len(employee_versions)):
                version = employee_versions[i]
                if version.date_version <= date_start or employee_versions[-1] == version:
                    employee_valid_versions |= version
                    break
                if employee_valid_versions:
                    if (employee_valid_versions[-1].contract_date_start > version.contract_date_start
                            and (version.contract_date_start >= version.date_version
                                 or version.contract_date_start > employee_versions[i + 1].contract_date_start)):
                        employee_valid_versions |= version
                elif (version.contract_date_start >= version.date_version
                      or version.contract_date_start > employee_versions[i + 1].contract_date_start):
                    employee_valid_versions |= version
            valid_versions |= employee_valid_versions

        Version = self.env['hr.version']
        kept = Version
        for version in valid_versions:
            emp = version.employee_id.sudo()
            if vacation_kind == 'anniversary_month':
                if not emp.contract_date_start or emp.contract_date_start.month != date_start.month:
                    continue
            elif vacation_kind == 'year_end_december':
                if date_start.year != date_end.year:
                    continue
                if date_start.month != 1 or date_start.day != 1:
                    continue
                dec_last = monthrange(date_end.year, 12)[1]
                if date_end.month != 12 or date_end.day != dec_last:
                    continue
                if emp.contract_date_start and emp.contract_date_start.month == 12:
                    continue
            if emp._hn_vacation_payroll_accrued_cash_days() <= 0:
                continue
            kept |= version
        return kept

    def _get_valid_version_ids(self, date_start=None, date_end=None, structure_id=None, company_id=None, employee_ids=None, schedule_pay=None):
        """Vacaciones HN: filtro por modalidad; resto: estándar + exclusión de archivados salvo excepción."""
        date_start = date_start or self.date_start
        date_end = date_end or self.date_end
        structure = self.env['hr.payroll.structure'].browse(structure_id) if structure_id else self.structure_id
        vacation_kind = structure.hn_vacation_payroll_kind if structure else False

        if vacation_kind:
            company = company_id or self.company_id.id
            versions = self._hn_vacation_payroll_valid_version_ids(
                date_start, date_end, company, employee_ids, vacation_kind,
            )
            version_ids = versions.ids if versions else []
        else:
            version_ids = super()._get_valid_version_ids(
                date_start=date_start, date_end=date_end,
                structure_id=structure_id, company_id=company_id,
                employee_ids=employee_ids, schedule_pay=schedule_pay,
            )
        if not version_ids:
            return version_ids

        versions_browse = self.env['hr.version'].browse(version_ids)

        def _keep_version(version):
            emp = version.employee_id
            if emp.active:
                return True
            if not getattr(emp, 'include_in_current_payroll', False):
                return False
            dep = emp.departure_date
            if not dep:
                return False
            return date_start <= dep <= date_end

        kept = versions_browse.filtered(_keep_version)
        return kept.ids

    def _generate_payslips_hn_vacation(self, version_ids=None, employee_ids=None):
        """Generación de lote nómina solo vacaciones HN (work entries opcionales)."""
        self.ensure_one()
        if employee_ids and not version_ids:
            version_ids = self._get_valid_version_ids(employee_ids=employee_ids)
        if not version_ids:
            raise UserError(self.env._('You must select employee(s) version(s) to generate payslip(s).'))
        valid_versions = self.env['hr.version'].browse(version_ids)
        Payslip = self.env['hr.payslip']
        if self.structure_id.use_worked_day_lines:
            valid_versions.generate_work_entries(self.date_start, self.date_end)
            all_work_entries = dict(self.env['hr.work.entry']._read_group(
                domain=[
                    ('employee_id', 'in', valid_versions.employee_id.ids),
                    ('date', '<=', self.date_end),
                    ('date', '>=', self.date_start),
                ],
                groupby=['version_id'],
                aggregates=['id:recordset'],
            ))
            utc = pytz.utc
            for tz, slips_per_tz in self.slip_ids.grouped(lambda s: s.version_id.tz).items():
                slip_tz = pytz.timezone(tz or utc)
                for slip in slips_per_tz:
                    date_from = slip_tz.localize(datetime.combine(slip.date_from, time.min)).astimezone(utc).replace(tzinfo=None)
                    date_to = slip_tz.localize(datetime.combine(slip.date_to, time.max)).astimezone(utc).replace(tzinfo=None)
                    if version_work_entries := all_work_entries.get(slip.version_id):
                        version_work_entries.filtered_domain([
                            ('date', '<=', date_to),
                            ('date', '>=', date_from),
                        ])
                        version_work_entries._check_undefined_slots(slip.date_from, slip.date_to)
            for work_entries in all_work_entries.values():
                work_entries = work_entries.filtered(lambda we: we.state != 'validated')
                if work_entries._check_if_error():
                    work_entries = work_entries.filtered(lambda we: we.state == 'conflict')
                    conflicts = work_entries._to_intervals()
                    time_intervals_str = ''.join(
                        f'\n - {start} -> {end} ({entry.employee_id.name})'
                        for start, end, entry in conflicts._items
                    )
                    raise UserError(
                        self.env._('Some work entries could not be validated. Time intervals to look for:%s')
                        % time_intervals_str
                    )
        default_values = Payslip.default_get(Payslip.fields_get())
        payslips_vals = []
        for version in valid_versions[::-1]:
            values = default_values | {
                'name': self.env._('New Payslip'),
                'employee_id': version.employee_id.id,
                'payslip_run_id': self.id,
                'date_from': self.date_start,
                'date_to': self.date_end,
                'version_id': version.id,
                'company_id': self.company_id.id,
                'struct_id': self.structure_id.id or version.structure_type_id.default_struct_id.id,
            }
            payslips_vals.append(values)
        self.slip_ids |= Payslip.with_context(tracking_disable=True).create(payslips_vals)
        self.slip_ids._compute_name()
        self.slip_ids.compute_sheet()
        self.state = '01_ready'
        return 1

    def generate_payslips(self, version_ids=None, employee_ids=None):
        """Igual que Enterprise salvo: no exigir type_id del lote = tipo del contrato en beneficios/vacaciones HN;

        y no abortar el lote por conflictos de hr.work.entry en estructuras HN listadas (criterio SV).
        """
        self.ensure_one()

        if self.structure_id and self.structure_id.hn_vacation_payroll_kind:
            return self._generate_payslips_hn_vacation(version_ids=version_ids, employee_ids=employee_ids)

        if employee_ids and not version_ids:
            version_ids = self._get_valid_version_ids(employee_ids=employee_ids)

        if not version_ids:
            raise UserError(
                self.env._("You must select employee(s) version(s) to generate payslip(s).")
            )

        valid_versions = self.env["hr.version"].browse(version_ids)
        Payslip = self.env["hr.payslip"]

        skip_structure_type_match = (
            self.structure_id
            and self.structure_id.code in _STRUCT_CODES_SKIP_PAYROLL_STRUCTURE_TYPE_MATCH
        )
        if self.structure_id and not skip_structure_type_match:
            valid_versions = valid_versions.filtered(
                lambda c: c.structure_type_id.id == self.structure_id.type_id.id
            )

        valid_versions.generate_work_entries(self.date_start, self.date_end)

        all_work_entries = dict(
            self.env["hr.work.entry"]._read_group(
                domain=[
                    ("employee_id", "in", valid_versions.employee_id.ids),
                    ("date", "<=", self.date_end),
                    ("date", ">=", self.date_start),
                ],
                groupby=["version_id"],
                aggregates=["id:recordset"],
            )
        )

        utc = pytz.utc
        for tz, slips_per_tz in self.slip_ids.grouped(lambda s: s.version_id.tz).items():
            slip_tz = pytz.timezone(tz or utc)
            for slip in slips_per_tz:
                date_from = slip_tz.localize(
                    datetime.combine(slip.date_from, time.min)
                ).astimezone(utc).replace(tzinfo=None)
                date_to = slip_tz.localize(
                    datetime.combine(slip.date_to, time.max)
                ).astimezone(utc).replace(tzinfo=None)
                if version_work_entries := all_work_entries.get(slip.version_id):
                    version_work_entries.filtered_domain(
                        [
                            ("date", "<=", date_to),
                            ("date", ">=", date_from),
                        ]
                    )
                    version_work_entries._check_undefined_slots(
                        slip.date_from, slip.date_to
                    )

        if not (
            self.structure_id
            and self.structure_id.code in _STRUCT_CODES_SKIP_WORK_ENTRY_VALIDATION
        ):
            for work_entries in all_work_entries.values():
                work_entries = work_entries.filtered(lambda we: we.state != "validated")
                if work_entries._check_if_error():
                    work_entries = work_entries.filtered(lambda we: we.state == "conflict")
                    conflicts = work_entries._to_intervals()
                    time_intervals_str = "".join(
                        f"\n - {start} -> {end} ({entry.employee_id.name})"
                        for start, end, entry in conflicts._items
                    )
                    raise UserError(
                        self.env._(
                            "Some work entries could not be validated. Time intervals to look for:%s",
                            time_intervals_str,
                        )
                    )

        default_values = Payslip.default_get(Payslip.fields_get())
        payslips_vals = []
        for version in valid_versions[::-1]:
            values = default_values | {
                "name": self.env._("New Payslip"),
                "employee_id": version.employee_id.id,
                "payslip_run_id": self.id,
                "date_from": self.date_start,
                "date_to": self.date_end,
                "version_id": version.id,
                "company_id": self.company_id.id,
                "struct_id": self.structure_id.id
                or version.structure_type_id.default_struct_id.id,
            }
            payslips_vals.append(values)
        self.slip_ids |= Payslip.with_context(tracking_disable=True).create(payslips_vals)
        self.slip_ids._compute_name()
        self.slip_ids.compute_sheet()
        self.state = "01_ready"

        return 1

    payroll_period_id = fields.Many2one('l10n.hn.payroll.period', string='Quincena',
                                        help='Select the payroll period (quincena) to generate payslips for',
                                        ondelete='set null')
    payroll_sheet_id = fields.Many2one('l10n.hn.payroll.sheet', string='Planilla',
                                       help='Planilla HE/TE/BD vinculada (opcional). La nómina considera TODAS las planillas confirmadas del período para HE y TE.',
                                       ondelete='set null')

    @api.depends("structure_id", "structure_id.type_id", "structure_id.type_id.default_schedule_pay")
    def _compute_schedule_pay(self):
        """Override to ensure schedule_pay is loaded when structure_id is selected"""
        for payslip_run in self:
            if payslip_run.structure_id and payslip_run.structure_id.type_id:
                # Load schedule_pay from structure type
                payslip_run.schedule_pay = payslip_run.structure_id.type_id.default_schedule_pay or 'monthly'
            elif not payslip_run.structure_id:
                # If no structure selected, keep current value or default
                if not payslip_run.schedule_pay:
                    payslip_run.schedule_pay = 'monthly'
    
    @api.onchange('structure_id')
    def _onchange_structure_id(self):
        """Update schedule_pay when structure_id changes in the wizard"""
        if self.structure_id and self.structure_id.type_id:
            self.schedule_pay = self.structure_id.type_id.default_schedule_pay or 'monthly'
    
    @api.depends('schedule_pay', 'payroll_period_id', 'country_code')
    def _compute_date_start(self):
        """Quincenas solo para compañías HN; resto delega al estándar / otras localizaciones."""
        hn_bi = self.filtered(
            lambda r: r.schedule_pay == 'bi-weekly' and r.country_code == 'HN'
        )
        for payslip_run in hn_bi:
            if payslip_run.payroll_period_id:
                payslip_run.date_start = payslip_run.payroll_period_id.date_start
            else:
                payslip_run.date_start = biweekly_start_from_today(fields.Date.today())
        other = self - hn_bi
        if other:
            super(HrPayslipRun, other)._compute_date_start()
    
    @api.depends('date_start', 'schedule_pay', 'payroll_period_id', 'country_code')
    def _compute_date_end(self):
        hn_bi = self.filtered(
            lambda r: r.schedule_pay == 'bi-weekly' and r.country_code == 'HN'
        )
        for payslip_run in hn_bi:
            if payslip_run.payroll_period_id:
                payslip_run.date_end = payslip_run.payroll_period_id.date_end
            elif payslip_run.date_start:
                payslip_run.date_end = biweekly_end_from_start(payslip_run.date_start)
        other = self - hn_bi
        if other:
            super(HrPayslipRun, other)._compute_date_end()
    
    @api.model_create_multi
    def create(self, vals_list):
        """Auto-create payroll period when creating bi-weekly payslip runs"""
        # Clean up dict values for Many2one fields (Odoo 19 frontend sometimes sends dicts)
        for vals in vals_list:
            for key in list(vals.keys()):
                value = vals[key]
                if isinstance(value, dict):
                    # Extract id from dict for Many2one fields
                    if 'id' in value:
                        vals[key] = value['id']
                    else:
                        # Remove invalid dict values
                        _logger.warning("Removing invalid dict value for field %s: %s", key, value)
                        del vals[key]
        
        # First create the records
        records = super().create(vals_list)
        
        # Solo HN: período l10n.hn.payroll.period (evita mezclar con SV u otros países)
        for record in records:
            if (
                record.schedule_pay == 'bi-weekly'
                and not record.payroll_period_id
                and record.country_code == 'HN'
            ):
                if record.date_start:
                    year = record.date_start.year
                    month = record.date_start.month
                    day = record.date_start.day
                else:
                    today = fields.Date.today()
                    year = today.year
                    month = today.month
                    day = today.day
                
                period_number = period_number_from_day(day)
                company_id = record.company_id.id if record.company_id else self.env.company.id
                
                period = self.env['l10n.hn.payroll.period'].get_or_create_period(
                    year, month, period_number, company_id)
                if period:
                    record.payroll_period_id = period
        
        return records

    @api.model
    def get_payroll_wizard_options(self, company_ids=None, company_id=None):
        """Return options for the payroll generation wizard: companies, structures, schedule_pay."""
        if not company_ids:
            company_ids = self.env.companies.ids
        companies = self.env['res.company'].search_read(
            [('id', 'in', company_ids)], ['id', 'name'], order='name'
        )
        struct_domain = [('active', '=', True)]
        if company_id:
            struct_domain += ['|', ('company_id', '=', False), ('company_id', '=', company_id)]
        Structure = self.env['hr.payroll.structure']
        struct_fields = ['id', 'name', 'type_id', 'company_id', 'schedule_pay']
        if 'is_default' in Structure._fields:
            struct_fields.append('is_default')
        structures = Structure.search_read(
            struct_domain, struct_fields,
            order='name'
        )
        StructureType = self.env['hr.payroll.structure.type']
        for srec in structures:
            if not srec.get('schedule_pay') and srec.get('type_id'):
                stype = StructureType.browse(srec['type_id'][0])
                if stype:
                    srec['schedule_pay'] = stype.default_schedule_pay

        schedule_pay_selection = self.env['hr.payroll.structure.type']._get_selection_schedule_pay()
        schedule_pay_options = [{'value': v, 'label': l} for v, l in schedule_pay_selection]

        # Estructura por defecto (prioridad: específica de compañía, luego compartida)
        default_structure_id = None
        if 'is_default' in Structure._fields:
            cid_for_default = company_id or self.env.company.id
            if cid_for_default:
                dom_base = [('active', '=', True), ('is_default', '=', True)]
                default_struct = Structure.search(
                    dom_base + [('company_id', '=', cid_for_default)],
                    limit=1
                )
                if not default_struct:
                    default_struct = Structure.search(
                        dom_base + [('company_id', '=', False)],
                        limit=1
                    )
                if default_struct:
                    default_structure_id = default_struct.id

        # Quincena en curso (para bi-weekly)
        current_period = None
        if company_id:
            period = self.env['l10n.hn.payroll.period'].get_current_period(company_id)
            if period:
                current_period = {
                    'id': period.id,
                    'name': period.name,
                    'date_start': period.date_start.isoformat() if period.date_start else None,
                    'date_end': period.date_end.isoformat() if period.date_end else None,
                }

        # Compañía por defecto (para pre-selección al abrir wizard)
        default_company_id = None
        if self.env.company.id in company_ids:
            default_company_id = self.env.company.id

        return {
            'companies': companies,
            'structures': structures,
            'schedule_pay_options': schedule_pay_options,
            'default_structure_id': default_structure_id,
            'current_period': current_period,
            'default_company_id': default_company_id,
        }

    def action_payroll_hr_version_list_view_payrun(self, date_start=None, date_end=None, structure_id=None, company_id=None, schedule_pay=None):
        structure = self.env['hr.payroll.structure'].browse(structure_id) if structure_id else self.structure_id
        if not (structure and structure.hn_vacation_payroll_kind):
            return super().action_payroll_hr_version_list_view_payrun(
                date_start=date_start, date_end=date_end, structure_id=structure_id,
                company_id=company_id, schedule_pay=schedule_pay,
            )
        action = self.env['ir.actions.act_window']._for_xml_id(
            'l10n_hn_payroll_base.action_hr_version_payrun_select'
        )
        valid_version_ids = self._get_valid_version_ids(
            fields.Date.from_string(date_start),
            fields.Date.from_string(date_end),
            structure_id,
            company_id,
            None,
            schedule_pay,
        )
        ds = fields.Date.from_string(date_start) if date_start else self.date_start
        de = fields.Date.from_string(date_end) if date_end else self.date_end
        sid = structure_id if structure_id else (self.structure_id.id if self.structure_id else False)
        payslip_domain = Domain.AND([
            Domain('version_id', 'in', valid_version_ids),
            Domain('date_from', '=', ds),
            Domain('date_to', '=', de),
            Domain('struct_id', '=', sid if sid else False),
            Domain('state', '!=', 'cancel'),
        ])
        existing_version_ids = self.env['hr.payslip'].search(payslip_domain).version_id.ids
        filtered_version_ids = set(valid_version_ids) - set(existing_version_ids)
        action['domain'] = [('id', 'in', list(filtered_version_ids))]
        return action

    @api.model
    def get_payroll_version_list(self, date_start, date_end, structure_id, company_id, schedule_pay, payroll_period_id=None):
        """Return list of valid hr.version records for the payroll generation wizard.
        Excludes versions that already have payslips for the given period."""
        structure = self.env['hr.payroll.structure'].browse(structure_id)
        run = self.browse([])
        date_start_dt = fields.Date.from_string(date_start) if isinstance(date_start, str) else date_start
        date_end_dt = fields.Date.from_string(date_end) if isinstance(date_end, str) else date_end

        valid_version_ids = run._get_valid_version_ids(
            date_start=date_start_dt,
            date_end=date_end_dt,
            structure_id=structure_id,
            company_id=company_id,
            employee_ids=None,
            schedule_pay=schedule_pay,
        )

        if not valid_version_ids:
            return []

        if structure.hn_vacation_payroll_kind:
            payslip_domain = Domain.AND([
                Domain('version_id', 'in', valid_version_ids),
                Domain('date_from', '=', date_start_dt),
                Domain('date_to', '=', date_end_dt),
                Domain('struct_id', '=', structure_id or False),
                Domain('state', '!=', 'cancel'),
            ])
        else:
            payslip_domain = Domain.AND([
                Domain('version_id', 'in', valid_version_ids),
                Domain('date_from', '=', date_start_dt),
                Domain('date_to', '=', date_end_dt),
                Domain('struct_id', '=', structure_id or False),
                Domain('state', '!=', 'cancel'),
                Domain('version_id.schedule_pay', '=', schedule_pay or False),
            ])
        existing_version_ids = self.env['hr.payslip'].search(payslip_domain).version_id.ids
        filtered_version_ids = list(set(valid_version_ids) - set(existing_version_ids))

        return self.env['hr.version'].search_read(
            [('id', 'in', filtered_version_ids)],
            ['id', 'employee_id', 'department_id', 'job_id', 'wage', 'currency_id', 'date_version', 'structure_type_id'],
            order='employee_id, date_version desc',
        )

    @api.onchange('schedule_pay')
    def _onchange_schedule_pay(self):
        """Auto-create payroll period when bi-weekly is selected"""
        if self.schedule_pay == 'bi-weekly':
            if not self.payroll_period_id:
                # Auto-select current period
                period = self.env['l10n.hn.payroll.period'].get_current_period(self.company_id.id if self.company_id else None)
                self.payroll_period_id = period.id
            # Update dates from period
            if self.payroll_period_id:
                self.date_start = self.payroll_period_id.date_start
                self.date_end = self.payroll_period_id.date_end
        else:
            # Clear payroll period for non-bi-weekly schedules
            self.payroll_period_id = False
    
    @api.onchange('payroll_period_id')
    def _onchange_payroll_period_id(self):
        """Update dates when payroll period changes"""
        if self.payroll_period_id and self.schedule_pay == 'bi-weekly':
            self.date_start = self.payroll_period_id.date_start
            self.date_end = self.payroll_period_id.date_end

    def action_diagnose_employees(self):
        """Diagnostic method to check why employees are not being loaded"""
        self.ensure_one()
        if not self.date_start or not self.date_end:
            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': _('Error'),
                    'message': _('Debe configurar las fechas de inicio y fin del período.'),
                    'type': 'danger',
                }
            }
        
        # Find all active employees
        all_employees = self.env['hr.employee'].search([('active', '=', True)])
        
        # Find employees with versions active during the period
        versions = self.env['hr.version'].search([
            ('contract_date_start', '<=', self.date_end),
            '|', ('contract_date_end', '=', False), ('contract_date_end', '>=', self.date_start),
        ])
        
        employees_with_versions = versions.mapped('employee_id')
        
        # Find employees with correct schedule_pay
        if self.schedule_pay:
            versions_with_schedule = versions.filtered(
                lambda v: v.schedule_pay == self.schedule_pay
            )
            employees_with_schedule = versions_with_schedule.mapped('employee_id')
        else:
            employees_with_schedule = employees_with_versions
        
        # Find employees with correct structure
        employees_with_structure = self.env['hr.employee']
        if self.structure_id:
            for version in versions:
                if version.structure_type_id == self.structure_id.type_id:
                    employees_with_structure |= version.employee_id
        
        message = _('Diagnóstico de Empleados:\n\n')
        message += _('- Total de empleados activos: %d\n') % len(all_employees)
        message += _('- Empleados con versiones activas en el período: %d\n') % len(employees_with_versions)
        if self.schedule_pay:
            message += _('- Empleados con schedule_pay = "%s": %d\n') % (self.schedule_pay, len(employees_with_schedule))
        if self.structure_id:
            message += _('- Empleados con estructura "%s": %d\n') % (self.structure_id.name, len(employees_with_structure))
        
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Diagnóstico'),
                'message': message,
                'type': 'info',
                'sticky': True,
            }
        }
    
    def action_compute_sheet(self):
        """Override to ensure employees are found for bi-weekly payrolls"""
        # For bi-weekly payrolls, temporarily update versions' schedule_pay
        # to match the payroll run so employees are included
        if self.schedule_pay == 'bi-weekly':
            # Find all versions active during the period that match the structure
            domain = [
                ('contract_date_start', '<=', self.date_end),
                '|', ('contract_date_end', '=', False), ('contract_date_end', '>=', self.date_start),
            ]
            
            # If structure is selected, filter by structure type
            if self.structure_id and self.structure_id.type_id:
                domain.append(('structure_type_id', '=', self.structure_id.type_id.id))
            
            versions = self.env['hr.version'].search(domain)
            
            # Temporarily update schedule_pay for versions that don't match
            versions_to_update = versions.filtered(lambda v: v.schedule_pay != 'bi-weekly')
            original_schedule_pays = {}
            original_periodicities = {}
            for version in versions_to_update:
                original_schedule_pays[version.id] = version.schedule_pay
                original_periodicities[version.id] = version.periodicity
                # Update both schedule_pay and periodicity to ensure consistency
                version.schedule_pay = 'bi-weekly'
                version.periodicity = 'biweekly'
            
            try:
                result = super().action_compute_sheet()
            finally:
                # Restore original values
                for version in versions_to_update:
                    if version.id in original_schedule_pays:
                        version.schedule_pay = original_schedule_pays[version.id]
                    if version.id in original_periodicities:
                        version.periodicity = original_periodicities[version.id]
            
            return result
        
        return super().action_compute_sheet()

    def action_export_payroll_xlsx(self):
        """Exportar reporte de nómina a Excel agrupado por departamento (estructura imagen)."""
        self.ensure_one()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))

        wb = Workbook()
        wb.remove(wb.active)  # Eliminar hoja por defecto; crearemos hojas por proyecto

        # Borde gris 35%
        gray_border = Border(
            left=Side(style='thin', color='A5A5A5'),
            right=Side(style='thin', color='A5A5A5'),
            top=Side(style='thin', color='A5A5A5'),
            bottom=Side(style='thin', color='A5A5A5'),
        )

        # Moneda de la compañía
        currency = self.company_id.currency_id or self.env.company.currency_id
        symbol = (currency.symbol or '').replace('"', '""')
        currency_fmt = '"%s"#,##0.00' % symbol if symbol else '#,##0.00'

        # Colores y fuentes
        header_fill = PatternFill(start_color='A52A2A', end_color='A52A2A', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        dept_header_fill = PatternFill(start_color='C0C0C0', end_color='C0C0C0', fill_type='solid')
        dept_header_font = Font(bold=True, color='000000')
        subtotal_fill = PatternFill(start_color='808080', end_color='808080', fill_type='solid')
        subtotal_font = Font(bold=True, color='FFFFFF')
        bold_font = Font(bold=True)
        red_font = Font(color='FF0000')
        red_bold_font = Font(bold=True, color='FF0000')

        # Obtener payslips antes de construir headers (necesarios para estructuras)
        slips = self.slip_ids.filtered(lambda s: s.employee_id)
        if not slips:
            raise UserError(_('No hay recibos de nómina en este período.'))

        # ── Columnas dinámicas basadas en reglas salariales ──────────
        structures = slips.mapped('struct_id')
        all_rules = structures.mapped('rule_ids').filtered(lambda r: r.active)
        _exclude = {'GROSS', 'NET'}

        _seen = set()
        income_rules = []
        for r in all_rules.sorted(key=lambda r: (r.sequence, r.id)):
            if (r.category_id and r.category_id.code in ('BASIC', 'ALW')
                    and r.code not in _exclude and r.code not in _seen):
                _seen.add(r.code)
                income_rules.append((r.code, r.name))

        _seen_ded = set()
        deduction_rules = []
        for r in all_rules.sorted(key=lambda r: (r.sequence, r.id)):
            if (r.category_id and r.category_id.code == 'DED'
                    and r.code not in _seen_ded):
                _seen_ded.add(r.code)
                deduction_rules.append((r.code, r.name))

        # Pre-computar totales por código desde line_ids (única fuente de verdad).
        # input_line_ids son entradas crudas que las reglas consumen para generar
        # line_ids; sumarlos duplicaría ingresos y cancelaría deducciones a cero.
        _slip_lines = {}
        for slip in slips:
            totals = {}
            for line in slip.line_ids:
                if line.code:
                    totals[line.code] = totals.get(line.code, 0.0) + line.total
            _slip_lines[slip.id] = totals

        # Detectar códigos extra en line_ids no cubiertos por reglas de estructura
        _known = (set(c for c, _ in income_rules)
                  | set(c for c, _ in deduction_rules) | _exclude)
        extra_income = {}
        extra_deductions = {}
        for slip in slips:
            for line in slip.line_ids:
                if not line.code or line.code in _known:
                    continue
                cat = (line.salary_rule_id.category_id.code
                       if line.salary_rule_id and line.salary_rule_id.category_id
                       else None)
                name = ((line.salary_rule_id and line.salary_rule_id.name)
                        or line.name or line.code)
                if cat in ('BASIC', 'ALW'):
                    extra_income.setdefault(line.code, name)
                elif cat == 'DED' or line.total < 0:
                    extra_deductions.setdefault(line.code, name)
                elif line.total > 0:
                    extra_income.setdefault(line.code, name)

        # Filtrar columnas cuyo total absoluto en el lote sea cero
        def _batch_abs_total(code):
            return abs(sum(_slip_lines[s.id].get(code, 0.0) for s in slips))

        income_cols = [(c, n) for c, n in income_rules
                       if _batch_abs_total(c) >= 0.01]
        income_cols += sorted(
            [(c, n) for c, n in extra_income.items()
             if _batch_abs_total(c) >= 0.01],
            key=lambda x: x[1])

        deduction_cols = [(c, n) for c, n in deduction_rules
                          if _batch_abs_total(c) >= 0.01]
        deduction_cols += sorted(
            [(c, n) for c, n in extra_deductions.items()
             if _batch_abs_total(c) >= 0.01],
            key=lambda x: x[1])

        # ── Headers dinámicos ──────────────────────────────────────────
        fixed_headers = [
            _('Departamento'), _('Empleado'), _('Asignación'), _('Proyecto'),
            _('Horas Normales'), _('Días lab'), _('Horas extras'),
        ]
        headers = (
            fixed_headers
            + [n for _, n in income_cols]
            + ['Total Ingresos']
            + [n for _, n in deduction_cols]
            + ['Total Deducciones', 'TOTAL NETO']
        )

        n_fixed = len(fixed_headers)
        n_income = len(income_cols)
        col_total_ingresos = n_fixed + n_income + 1
        n_ded = len(deduction_cols)
        col_total_deducciones = col_total_ingresos + n_ded + 1
        col_neto = col_total_deducciones + 1
        CURRENCY_COLS = list(range(n_fixed + 1, col_neto + 1))

        def _style_cell(cell, fill=None, font=None, border=None, number_format=None):
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if border:
                cell.border = border
            if number_format:
                cell.number_format = number_format

        def _sanitize_sheet_name(name):
            """Excel sheet names: max 31 chars, no : \\ / ? * [ ]"""
            if not name:
                return _('Sin proyecto')
            for c in ':\\/?*[]':
                name = name.replace(c, '_')
            return (name[:31] if len(name) > 31 else name) or _('Sin proyecto')

        def _get_slip_row_values(slip, dept_name):
            """Valores de fila para un slip usando line_ids pre-computados."""
            lt = _slip_lines[slip.id]
            days_worked = (slip._get_report_days_worked()
                           if hasattr(slip, '_get_report_days_worked')
                           else (slip.sum_worked_hours or 0) / 8)
            hours_worked = (slip._get_report_hours_worked()
                            if hasattr(slip, '_get_report_hours_worked')
                            else (slip.sum_worked_hours or 0))
            he = (slip._get_report_he_breakdown()
                  if hasattr(slip, '_get_report_he_breakdown') else {})
            horas_extras = sum(he.get(k, 0)
                               for k in ('he_25', 'he_50', 'he_75', 'he_100'))

            inc_vals = [round(lt.get(c, 0.0), 2) for c, _ in income_cols]
            total_ingresos = sum(lt.get(c, 0.0) for c, _ in income_cols)

            ded_vals = [round(abs(lt.get(c, 0.0)), 2) for c, _ in deduction_cols]
            total_deducciones = sum(abs(lt.get(c, 0.0))
                                    for c, _ in deduction_cols)

            total_neto = total_ingresos - total_deducciones
            pago_quincenal = lt.get('BASIC', 0.0)

            emp = slip.employee_id
            asignacion = ((emp.assignment_id and emp.assignment_id.name)
                          or _('Sin asignación'))
            proyecto = ((emp.project_id and emp.project_id.name)
                        or _('Sin proyecto'))
            values = (
                [dept_name, emp.name or '', asignacion, proyecto,
                 round(hours_worked, 2), round(days_worked, 2),
                 round(horas_extras, 2)]
                + inc_vals + [round(total_ingresos, 2)]
                + ded_vals + [round(total_deducciones, 2),
                              round(total_neto, 2)]
            )
            return values, total_neto, pago_quincenal, total_ingresos, total_deducciones

        # Agrupar por proyecto y por asignación
        by_project = {}
        for slip in slips:
            emp = slip.employee_id
            proj_name = (emp.project_id and emp.project_id.name) or _('Sin proyecto')
            asig_name = (emp.assignment_id and emp.assignment_id.name) or _('Sin asignación')
            if proj_name not in by_project:
                by_project[proj_name] = {}
            if asig_name not in by_project[proj_name]:
                by_project[proj_name][asig_name] = []
            by_project[proj_name][asig_name].append(slip)

        by_dept = {}
        for slip in slips:
            dept_name = (slip.department_id and slip.department_id.name) or _('Sin departamento')
            if dept_name not in by_dept:
                by_dept[dept_name] = []
            by_dept[dept_name].append(slip)

        grand_totals = {i: 0.0 for i in range(1, len(headers) + 1)}
        summary_asignacion = []

        # --- Hojas por proyecto (agrupadas por asignación) ---
        sheet_index = 0
        for proj_name in sorted(by_project.keys()):
            ws_proj = wb.create_sheet(
                _sanitize_sheet_name(proj_name) if proj_name else _('Sin proyecto'),
                sheet_index,
            )
            sheet_index += 1
            for c, h in enumerate(headers, 1):
                cell = ws_proj.cell(row=1, column=c, value=h)
                _style_cell(cell, fill=header_fill, font=header_font, border=gray_border)

            row = 2
            proj_totals = {i: 0.0 for i in range(1, len(headers) + 1)}

            for asig_name in sorted(by_project[proj_name].keys()):
                asig_slips = by_project[proj_name][asig_name]

                for c in range(1, len(headers) + 1):
                    cell = ws_proj.cell(row=row, column=c, value=asig_name if c == 3 else (proj_name if c == 4 else ''))
                    _style_cell(cell, fill=dept_header_fill, font=dept_header_font, border=gray_border)
                row += 1

                asig_totals = {i: 0.0 for i in range(1, len(headers) + 1)}
                asig_pago = asig_ingresos = asig_ded = asig_neto = 0.0

                for slip in asig_slips:
                    d = (slip.department_id and slip.department_id.name) or _('Sin departamento')
                    values, total_neto, pago_q, tot_ing, tot_ded = _get_slip_row_values(slip, d)
                    for c, val in enumerate(values, 1):
                        cell = ws_proj.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                        if c in CURRENCY_COLS:
                            cell.number_format = currency_fmt
                        if isinstance(val, (int, float)) and val == 0 and c in CURRENCY_COLS:
                            cell.font = red_bold_font if c in (col_total_ingresos, col_neto) else red_font
                        elif c in (col_total_ingresos, col_neto):
                            cell.font = bold_font
                        asig_totals[c] += val if isinstance(val, (int, float)) else 0
                        if c > 4:
                            proj_totals[c] += val if isinstance(val, (int, float)) else 0
                    asig_pago += pago_q
                    asig_ingresos += tot_ing
                    asig_ded += tot_ded
                    asig_neto += total_neto
                    row += 1

                summary_asignacion.append(
                    (asig_name, proj_name, asig_pago, asig_ingresos, asig_ded, asig_neto))

                for c in range(1, len(headers) + 1):
                    val = len(asig_slips) if c == 3 else round(asig_totals.get(c, 0), 2)
                    cell = ws_proj.cell(row=row, column=c, value=val)
                    _style_cell(cell, fill=subtotal_fill, font=subtotal_font, border=gray_border)
                    if c in CURRENCY_COLS:
                        cell.number_format = currency_fmt
                row += 1

            for c in range(1, len(headers) + 1):
                val = _('TOTAL') if c == 1 else round(proj_totals.get(c, 0), 2)
                cell = ws_proj.cell(row=row, column=c, value=val)
                _style_cell(cell, fill=subtotal_fill, font=subtotal_font, border=gray_border)
                if c in CURRENCY_COLS:
                    cell.number_format = currency_fmt
                if c > 4:
                    grand_totals[c] += proj_totals.get(c, 0)
            row += 1

            for c in range(1, len(headers) + 1):
                ws_proj.column_dimensions[get_column_letter(c)].width = 14

        # --- Hoja Resumen por asignación ---
        ws_resumen_asig = wb.create_sheet(_('Resumen por asignación'), sheet_index)
        sheet_index += 1
        resumen_asig_headers = [
            _('Asignación'), _('Proyecto'), _('Total pago quincenal'), _('Total ingresos'),
            _('Total deducciones'), _('Total Neto'),
        ]
        for c, h in enumerate(resumen_asig_headers, 1):
            cell = ws_resumen_asig.cell(row=1, column=c, value=h)
            _style_cell(cell, fill=header_fill, font=header_font, border=gray_border)
        for r, (asig, proj, pago, ingresos, ded, neto) in enumerate(summary_asignacion, 2):
            for c, val in enumerate([asig, proj, pago, ingresos, ded, neto], 1):
                cell = ws_resumen_asig.cell(row=r, column=c, value=val)
                cell.border = gray_border
                if c >= 3:
                    cell.number_format = currency_fmt
        total_row = len(summary_asignacion) + 2
        for c in range(1, 7):
            cell = ws_resumen_asig.cell(row=total_row, column=c)
            _style_cell(cell, fill=subtotal_fill, font=subtotal_font, border=gray_border)
            if c == 1:
                cell.value = _('TOTAL')
            elif c == 2:
                cell.value = ''
            else:
                cell.value = round(sum(s[c - 1] for s in summary_asignacion), 2)
                cell.number_format = currency_fmt
        for c in range(1, 7):
            ws_resumen_asig.column_dimensions[get_column_letter(c)].width = 18

        summary_data = []
        for dept_name in sorted(by_dept.keys()):
            dept_slips = by_dept[dept_name]
            dept_sums = {'pago': 0.0, 'ingresos': 0.0, 'ded': 0.0, 'neto': 0.0}
            for slip in dept_slips:
                values, total_neto, pago_q, tot_ing, tot_ded = _get_slip_row_values(slip, dept_name)
                dept_sums['pago'] += pago_q
                dept_sums['ingresos'] += tot_ing
                dept_sums['ded'] += tot_ded
                dept_sums['neto'] += total_neto
            summary_data.append((
                dept_name,
                round(dept_sums['pago'], 2),
                round(dept_sums['ingresos'], 2),
                round(dept_sums['ded'], 2),
                round(dept_sums['neto'], 2),
            ))

        # --- Hoja Resumen por departamento ---
        ws_resumen = wb.create_sheet(_('Resumen por departamento'), sheet_index)
        sheet_index += 1
        resumen_headers = ['Departamento', 'Total pago quincenal', 'Total ingresos', 'Total deducciones', 'Total Neto']
        for c, h in enumerate(resumen_headers, 1):
            cell = ws_resumen.cell(row=1, column=c, value=h)
            _style_cell(cell, fill=header_fill, font=header_font, border=gray_border)
        resumen_currency_cols = (2, 3, 4, 5)
        for r, (dept, pago, ingresos, deducciones, neto) in enumerate(summary_data, 2):
            for c, val in enumerate([dept, pago, ingresos, deducciones, neto], 1):
                cell = ws_resumen.cell(row=r, column=c, value=val)
                cell.border = gray_border
                if c in resumen_currency_cols:
                    cell.number_format = currency_fmt
        # Fila total general
        total_row = len(summary_data) + 2
        for c in range(1, 6):
            cell = ws_resumen.cell(row=total_row, column=c)
            _style_cell(cell, fill=subtotal_fill, font=subtotal_font, border=gray_border)
            if c == 1:
                cell.value = _('TOTAL')
            else:
                idx = c - 1
                cell.value = round(sum(s[idx] for s in summary_data), 2)
                cell.number_format = currency_fmt
        for c in range(1, 6):
            ws_resumen.column_dimensions[get_column_letter(c)].width = 18

        # --- Hoja Detalle (todos los empleados con Asignación y Proyecto) ---
        ws_detalle = wb.create_sheet(_('Detalle'), sheet_index)
        for c, h in enumerate(headers, 1):
            cell = ws_detalle.cell(row=1, column=c, value=h)
            _style_cell(cell, fill=header_fill, font=header_font, border=gray_border)
        detalle_row = 2
        for dept_name in sorted(by_dept.keys()):
            dept_slips = by_dept[dept_name]
            for slip in dept_slips:
                values, _unused_neto, _unused_pago, _unused_ing, _unused_ded = _get_slip_row_values(slip, dept_name)
                for c, val in enumerate(values, 1):
                    cell = ws_detalle.cell(row=detalle_row, column=c, value=val)
                    cell.border = gray_border
                    if c in CURRENCY_COLS:
                        cell.number_format = currency_fmt
                    if isinstance(val, (int, float)) and val == 0 and c in CURRENCY_COLS:
                        cell.font = red_bold_font if c in (col_total_ingresos, col_neto) else red_font
                    elif c in (col_total_ingresos, col_neto):
                        cell.font = bold_font
                detalle_row += 1

        # Fila total general - Hoja Detalle
        for c in range(1, len(headers) + 1):
            val = _('TOTAL') if c == 1 else (len(slips) if c == 2 else round(grand_totals.get(c, 0), 2))
            cell = ws_detalle.cell(row=detalle_row, column=c, value=val)
            _style_cell(cell, fill=subtotal_fill, font=subtotal_font, border=gray_border)
            if c in CURRENCY_COLS:
                cell.number_format = currency_fmt

        # Ajustar anchos de columna - Detalle
        for c in range(1, len(headers) + 1):
            ws_detalle.column_dimensions[get_column_letter(c)].width = 14

        # Descargar
        import base64
        import io
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        # Nombre del archivo = nombre de la quincena (run)
        import re
        safe_name = re.sub(r'[^\w\s\-]', '', (self.name or 'report'))
        filename = '%s.xlsx' % safe_name.replace(' ', '_')
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(buffer.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'self',
        }
