# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik

import base64
import io
from datetime import datetime
from calendar import monthrange

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError
from odoo.fields import Domain


class L10nHnPayrollSheet(models.Model):
    """Planilla de Horas Extra, Turnos Extra y Beneficios/Deducciones"""
    _name = 'l10n.hn.payroll.sheet'
    _description = 'Planilla HE/TE/BD'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'id desc'

    name = fields.Char(string='Nombre', required=True, tracking=True)
    payroll_period_id = fields.Many2one('l10n.hn.payroll.period', string='Quincena', required=True,
                                        ondelete='restrict', tracking=True)
    date_start = fields.Date(string='Fecha Inicio', related='payroll_period_id.date_start', readonly=True)
    date_end = fields.Date(string='Fecha Fin', related='payroll_period_id.date_end', readonly=True)
    cutoff_date_from = fields.Date(string='Corte HE Desde', readonly=True)
    cutoff_date_to = fields.Date(string='Corte HE Hasta', readonly=True)
    user_id = fields.Many2one('res.users', string='Responsable', default=lambda self: self.env.user, required=True, tracking=True)
    department_id = fields.Many2one('hr.department', string='Departamento',
                                    help='Si se configura, Obtener Empleados filtra por este departamento', tracking=True)
    filter_my_department_only = fields.Boolean(string='Solo mi departamento', default=False, tracking=True)
    company_id = fields.Many2one('res.company', string='Compañía', default=lambda self: self.env.company, required=True, tracking=True)
    state = fields.Selection([
        ('draft', 'Borrador'),
        ('done', 'Confirmada'),
    ], string='Estado', default='draft', required=True, tracking=True)
    line_ids = fields.One2many('l10n.hn.payroll.line.he', 'parent_id', string='Horas Extra')
    line_te_ids = fields.One2many('l10n.hn.payroll.line.te', 'parent_id', string='Turnos Extra')
    line_bd_ids = fields.One2many('l10n.hn.payroll.line.bd', 'parent_id', string='Beneficios y Deducciones')

    # Resumen de totales (para cabecera y reportes)
    total_he = fields.Float(string='Total HE', digits=(12, 2), compute='_compute_totals', store=True)
    total_te = fields.Float(string='Total TE', digits=(12, 2), compute='_compute_totals', store=True)
    total_bd = fields.Float(string='Total BD', digits=(12, 2), compute='_compute_totals', store=True)

    payslip_run_ids = fields.One2many('hr.payslip.run', 'payroll_sheet_id', string='Nóminas',
                                       readonly=True, help='Nóminas que usaron esta planilla')

    # Empleados en líneas (para filtro jerárquico)
    line_employee_ids = fields.Many2many('hr.employee', string='Empleados en Planilla',
                                          compute='_compute_line_employees', store=True,
                                          help='Empleados que aparecen en las líneas HE/TE/BD')

    @api.depends('line_ids.amount_total', 'line_te_ids.amount_other_income', 'line_bd_ids.amount')
    def _compute_totals(self):
        for sheet in self:
            sheet.total_he = sum(sheet.line_ids.mapped('amount_total'))
            sheet.total_te = sum(sheet.line_te_ids.mapped('amount_other_income'))
            sheet.total_bd = sum(sheet.line_bd_ids.mapped('amount'))

    @api.depends('line_ids.employee_id', 'line_te_ids.employee_id', 'line_bd_ids.employee_id')
    def _compute_line_employees(self):
        for sheet in self:
            employees = self.env['hr.employee']
            employees |= sheet.line_ids.mapped('employee_id')
            employees |= sheet.line_te_ids.mapped('employee_id')
            employees |= sheet.line_bd_ids.mapped('employee_id')
            sheet.line_employee_ids = employees

    @api.model
    def _get_name_for_period(self, payroll_period):
        """Nombre según quincena: Planilla para 1ra Quincena MMM-YYYY"""
        if not payroll_period or not payroll_period.year or not payroll_period.month:
            return _('Planilla')
        month_names = {
            1: 'enero', 2: 'febrero', 3: 'marzo', 4: 'abril',
            5: 'mayo', 6: 'junio', 7: 'julio', 8: 'agosto',
            9: 'septiembre', 10: 'octubre', 11: 'noviembre', 12: 'diciembre'
        }
        ord_text = '1ra' if payroll_period.period_number == '1' else '2da'
        month_name = month_names.get(payroll_period.month, '')
        return f'Planilla para {ord_text} Quincena {month_name}-{payroll_period.year}'

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if 'name' not in vals or not vals.get('name'):
                period = self.env['l10n.hn.payroll.period'].browse(vals.get('payroll_period_id'))
                if period:
                    vals['name'] = self._get_name_for_period(period)
            if payroll_period_id := vals.get('payroll_period_id'):
                period = self.env['l10n.hn.payroll.period'].browse(payroll_period_id)
                cutoff_from, cutoff_to = self.env['l10n.hn.payroll.cutoff.config'].get_cutoff_dates(period)
                vals['cutoff_date_from'] = cutoff_from
                vals['cutoff_date_to'] = cutoff_to
        return super().create(vals_list)

    @api.onchange('payroll_period_id')
    def _onchange_payroll_period_id(self):
        if self.payroll_period_id:
            self.name = self._get_name_for_period(self.payroll_period_id)
            cutoff_from, cutoff_to = self.env['l10n.hn.payroll.cutoff.config'].get_cutoff_dates(self.payroll_period_id)
            self.cutoff_date_from = cutoff_from
            self.cutoff_date_to = cutoff_to

    def _get_employees_domain(self):
        self.ensure_one()
        base_domain = Domain([
            ('company_id', '=', self.company_id.id),
        ])
        if self.department_id:
            base_domain &= Domain([('department_id', '=', self.department_id.id)])
        elif self.filter_my_department_only and self.env.user.employee_id:
            if self.env.user.employee_id.department_id:
                base_domain &= Domain([('department_id', '=', self.env.user.employee_id.department_id.id)])

        # Activos O archivados con flag de incluir en quincena en curso (departure_date en período)
        if self.date_start and self.date_end:
            active_domain = Domain([('active', '=', True)])
            archived_ok_domain = Domain([
                ('active', '=', False),
                ('include_in_current_payroll', '=', True),
                ('departure_date', '>=', self.date_start),
                ('departure_date', '<=', self.date_end),
            ])
            domain = (active_domain | archived_ok_domain) & base_domain
        else:
            domain = Domain([('active', '=', True)]) & base_domain
        return list(domain)

    def _get_versions_for_employees(self, employees):
        """Versiones de contrato (hr.version) activas en el período. Prefiere current_version_id para consistencia con payslip."""
        self.ensure_one()
        if not self.date_start or not self.date_end:
            return self.env['hr.version']
        versions = self.env['hr.version'].search([
            ('employee_id', 'in', employees.ids),
            ('contract_date_start', '<=', self.date_end),
            '|', ('contract_date_end', '=', False), ('contract_date_end', '>=', self.date_start),
        ])
        # Preferir current_version_id por empleado para consistencia con payslip
        result = self.env['hr.version']
        for emp in employees:
            current = emp.current_version_id.sudo() if emp.current_version_id else False
            if current and current in versions:
                result |= current
            else:
                emp_versions = versions.filtered(lambda v: v.employee_id == emp)
                if emp_versions:
                    result |= emp_versions[0]
        return result

    def _get_version_for_employee(self, employee):
        """Versión de contrato activa para un empleado en el período de la planilla."""
        if not employee:
            return self.env['hr.version']
        versions = self._get_versions_for_employees(employee)
        return versions[:1]

    def action_get_employees_he(self):
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se pueden obtener empleados en planillas en borrador.'))
        employees = self.env['hr.employee'].search(self._get_employees_domain())
        versions = self._get_versions_for_employees(employees)
        existing_employee_ids = self.line_ids.mapped('employee_id').ids
        new_lines = []
        for version in versions:
            emp = version.employee_id
            if emp.id not in existing_employee_ids:
                new_lines.append((0, 0, {
                    'employee_id': emp.id,
                    'parent_id': self.id,
                }))
                existing_employee_ids.append(emp.id)
        if new_lines:
            self.write({'line_ids': new_lines})
            self.message_post(body=_('Se obtuvieron %d empleados en Horas Extra.') % len(new_lines))
        return True

    def action_get_employees_te(self):
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se pueden obtener empleados en planillas en borrador.'))
        employees = self.env['hr.employee'].search(self._get_employees_domain())
        versions = self._get_versions_for_employees(employees)
        existing_employee_ids = self.line_te_ids.mapped('employee_id').ids
        new_lines = []
        for version in versions:
            emp = version.employee_id
            if emp.id not in existing_employee_ids:
                new_lines.append((0, 0, {
                    'employee_id': emp.id,
                    'parent_id': self.id,
                }))
                existing_employee_ids.append(emp.id)
        if new_lines:
            self.write({'line_te_ids': new_lines})
            self.message_post(body=_('Se obtuvieron %d empleados en Turnos Extra.') % len(new_lines))
        return True

    def action_get_employees_all(self):
        """Obtener empleados en las tres secciones (HE, TE, BD) en un solo clic"""
        self.ensure_one()
        self.action_get_employees_he()
        self.action_get_employees_te()
        self.action_get_employees_bd()
        return True

    def action_get_employees_bd(self):
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se pueden obtener empleados en planillas en borrador.'))
        employees = self.env['hr.employee'].search(self._get_employees_domain())
        versions = self._get_versions_for_employees(employees)
        existing_employee_ids = self.line_bd_ids.mapped('employee_id').ids
        new_lines = []
        for version in versions:
            emp = version.employee_id
            if emp.id not in existing_employee_ids:
                new_lines.append((0, 0, {
                    'employee_id': emp.id,
                    'parent_id': self.id,
                    'amount': 0.0,
                }))
                existing_employee_ids.append(emp.id)
        if new_lines:
            self.write({'line_bd_ids': new_lines})
            self.message_post(body=_('Se obtuvieron %d empleados en Beneficios y Deducciones.') % len(new_lines))
        return True

    def action_confirm(self):
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se puede confirmar una planilla en borrador.'))
        if not self.line_ids and not self.line_te_ids and not self.line_bd_ids:
            raise UserError(_('No puede confirmar una planilla vacía. Obtenga empleados primero.'))
        self._apply_he_to_contracts()
        self._apply_te_to_contracts()
        self._apply_bd_to_contracts()
        self.write({'state': 'done'})
        self.message_post(body=_('Planilla confirmada. HE, TE y BD aplicados a contratos.'))
        return True

    def _apply_he_to_contracts(self):
        for line in self.line_ids:
            # Usar version_id (versión activa en el período de la planilla) para que el payslip la encuentre
            version = line.version_id
            if line.amount_total and version:
                self.env['l10n.hn.extra.hours.register'].create({
                    'version_id': version.id,
                    'payroll_sheet_id': self.id,
                    'he_25': line.he_25,
                    'he_50': line.he_50,
                    'he_75': line.he_75,
                    'he_100': line.he_100,
                    'amount_25': line.amount_25,
                    'amount_50': line.amount_50,
                    'amount_75': line.amount_75,
                    'amount_100': line.amount_100,
                })

    def _apply_te_to_contracts(self):
        for line in self.line_te_ids:
            # Usar version_id (versión activa en el período de la planilla) para que el payslip la encuentre
            # SUMAR para soportar múltiples planillas en el mismo período (no sobrescribir)
            version = line.version_id
            if line.amount_other_income and version:
                version.amount_other_income = (version.amount_other_income or 0) + line.amount_other_income

    def _apply_bd_to_contracts(self):
        for line in self.line_bd_ids:
            if line.concept_id and line.amount and line.version_id:
                employee = line.version_id.employee_id
                if not employee:
                    continue
                include_in = line.include_in or ('first' if self.date_start.day <= 15 else 'second')
                concept = line.concept_id
                existing = self.env['hr.employee.recurring.concept'].search([
                    ('employee_id', '=', employee.id),
                    ('name', '=', concept.name),
                    ('active', '=', True),
                ], limit=1)
                vals = {
                    'employee_id': employee.id,
                    'company_id': employee.company_id.id if employee.company_id else self.company_id.id,
                    'name': concept.name,
                    'code': concept.code or concept.name[:20].upper().replace(' ', '_'),
                    'concept_type': concept.concept_type,
                    'amount_type': 'fixed',
                    'amount': abs(line.amount) if concept.concept_type == 'deduction' else line.amount,
                    'date_start': self.date_start,
                    'date_end': self.date_end,
                    'apply_biweek': include_in,
                    'active': True,
                    'payroll_sheet_line_bd_id': line.id,
                }
                if existing:
                    existing.write(vals)
                else:
                    self.env['hr.employee.recurring.concept'].create(vals)

    def action_set_to_draft(self):
        self.ensure_one()
        if self.state != 'done':
            raise UserError(_('Solo se puede reestablecer a borrador una planilla confirmada.'))
        if not self.env.user.has_group('l10n_hn_hr_payroll.group_hr_payroll_admin'):
            raise UserError(_('Solo el administrador de nómina puede reestablecer una planilla a borrador.'))
        self.env['l10n.hn.extra.hours.register'].search([('payroll_sheet_id', '=', self.id)]).unlink()
        for line in self.line_te_ids:
            version = line.version_id
            if version and line.amount_other_income:
                # Restar solo lo que aportó esta planilla (soporta múltiples planillas)
                version.amount_other_income = max(0, (version.amount_other_income or 0) - line.amount_other_income)
        # Desactivar conceptos creados desde BD usando trazabilidad
        for line in self.line_bd_ids:
            self.env['hr.employee.recurring.concept'].search([
                ('payroll_sheet_line_bd_id', '=', line.id),
            ]).write({'active': False, 'payroll_sheet_line_bd_id': False})
        self.write({'state': 'draft'})
        self.message_post(body=_('Planilla reestablecida a borrador. Registros HE/TE/BD revertidos.'))
        return True

    def action_open_import_wizard(self):
        """Abrir wizard de importación HHEE/TE con esta planilla en contexto."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Importar HHEE y TE'),
            'res_model': 'hr.he.te.import.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_payroll_sheet_id': self.id},
        }

    def action_open_bd_import_wizard(self):
        """Abrir wizard de importación BD con esta planilla en contexto."""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Importar Beneficios y Deducciones'),
            'res_model': 'hr.bd.import.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_payroll_sheet_id': self.id},
        }

    def action_clean_all(self):
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se pueden limpiar líneas en planillas en borrador.'))
        self.env['l10n.hn.extra.hours.register'].search([('payroll_sheet_id', '=', self.id)]).unlink()
        self.write({
            'line_ids': [(5, 0, 0)],
            'line_te_ids': [(5, 0, 0)],
            'line_bd_ids': [(5, 0, 0)],
        })
        self.message_post(body=_('Se eliminaron todas las líneas de la planilla.'))
        return True

    def action_clean_zero_totals(self):
        """Eliminar líneas de empleados cuyo total HE, TE y BD sea cero."""
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Solo se pueden limpiar líneas en planillas en borrador.'))
        # Empleados con totales en cero (HE=0, TE=0, BD=0)
        emp_he = {l.employee_id.id: l.amount_total or 0 for l in self.line_ids}
        emp_te = {l.employee_id.id: l.amount_other_income or 0 for l in self.line_te_ids}
        emp_bd = {}
        for line in self.line_bd_ids:
            emp_bd[line.employee_id.id] = emp_bd.get(line.employee_id.id, 0) + (line.amount or 0)
        all_emp_ids = set(emp_he.keys()) | set(emp_te.keys()) | set(self.line_bd_ids.mapped('employee_id').ids)
        to_remove = [
            eid for eid in all_emp_ids
            if (emp_he.get(eid, 0) == 0 and emp_te.get(eid, 0) == 0 and emp_bd.get(eid, 0) == 0)
        ]
        if not to_remove:
            return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {
                'title': _('Sin cambios'),
                'message': _('No hay empleados con totales en cero.'),
                'type': 'info',
            }}
        self.line_ids.filtered(lambda l: l.employee_id.id in to_remove).unlink()
        self.line_te_ids.filtered(lambda l: l.employee_id.id in to_remove).unlink()
        self.line_bd_ids.filtered(lambda l: l.employee_id.id in to_remove).unlink()
        self.message_post(body=_('Se eliminaron %d empleado(s) con totales en cero.') % len(to_remove))
        return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {
            'title': _('Líneas eliminadas'),
            'message': _('Se eliminaron %d empleado(s) con totales en cero.') % len(to_remove),
            'type': 'success',
        }}

    def action_export_xlsx(self):
        """Exportar planilla a Excel agrupada por departamento"""
        self.ensure_one()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))

        wb = Workbook()
        header_fill = PatternFill(start_color='8B0000', end_color='8B0000', fill_type='solid')  # Rojo oscuro
        header_font = Font(bold=True, color='FFFFFF')
        dept_subheader_fill = PatternFill(start_color='ADD8E6', end_color='ADD8E6', fill_type='solid')  # Azul claro
        dept_summary_fill = PatternFill(start_color='404040', end_color='404040', fill_type='solid')  # Gris oscuro
        summary_font = Font(bold=True, color='FFFFFF')

        # Hoja Planilla HE agrupada por departamento (formato imagen)
        ws_he = wb.active
        ws_he.title = 'Horas Extra'
        he_headers = [
            'Departamento', 'Empleado', 'Salario Quincenal', 'Sueldo x Hora',
            'Horas al 25%', 'Horas al 50%', 'Horas al 75%', 'Horas al 100%', 'Total Extras',
            'Monto 25%', 'Monto 50%', 'Monto 75%', 'Monto 100%', 'Total',
        ]
        for c, h in enumerate(he_headers, 1):
            cell = ws_he.cell(row=1, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font

        # Agrupar líneas por departamento
        lines_by_dept = {}
        for line in self.line_ids:
            dept = line.department_id or self.env['hr.department']
            dept_name = dept.name or _('Sin departamento')
            if dept_name not in lines_by_dept:
                lines_by_dept[dept_name] = []
            lines_by_dept[dept_name].append(line)

        # Totales generales para fila final
        g_salario = g_sueldo_hora = g_he25 = g_he50 = g_he75 = g_he100 = 0.0
        g_extras = g_m25 = g_m50 = g_m75 = g_m100 = g_monto = 0.0
        g_empleados = 0

        row = 2
        for dept_name in sorted(lines_by_dept.keys()):
            lines = lines_by_dept[dept_name]
            # Fila subencabezado departamento (azul claro)
            ws_he.cell(row=row, column=1, value=dept_name).fill = dept_subheader_fill
            for c in range(2, 15):
                ws_he.cell(row=row, column=c).fill = dept_subheader_fill
            row += 1

            total_salario = total_sueldo_hora = 0.0
            total_he25 = total_he50 = total_he75 = total_he100 = 0.0
            total_extras = 0.0
            total_m25 = total_m50 = total_m75 = total_m100 = total_monto = 0.0

            for line in lines:
                wage = (line.version_id and line.version_id.wage) or 0.0
                salario_quincenal = wage / 2.0 if wage else 0.0
                sueldo_hora = wage / (30 * 8) if wage else 0.0
                total_extras_line = (line.he_25 or 0) + (line.he_50 or 0) + (line.he_75 or 0) + (line.he_100 or 0)

                ws_he.cell(row=row, column=1, value=dept_name)
                ws_he.cell(row=row, column=2, value=line.employee_id.name or '')
                ws_he.cell(row=row, column=3, value=round(salario_quincenal, 2))
                ws_he.cell(row=row, column=4, value=round(sueldo_hora, 2))
                ws_he.cell(row=row, column=5, value=line.he_25 or 0)
                ws_he.cell(row=row, column=6, value=line.he_50 or 0)
                ws_he.cell(row=row, column=7, value=line.he_75 or 0)
                ws_he.cell(row=row, column=8, value=line.he_100 or 0)
                ws_he.cell(row=row, column=9, value=round(total_extras_line, 2))
                ws_he.cell(row=row, column=10, value=line.amount_25 or 0)
                ws_he.cell(row=row, column=11, value=line.amount_50 or 0)
                ws_he.cell(row=row, column=12, value=line.amount_75 or 0)
                ws_he.cell(row=row, column=13, value=line.amount_100 or 0)
                ws_he.cell(row=row, column=14, value=line.amount_total or 0)
                row += 1

                total_salario += salario_quincenal
                total_sueldo_hora += sueldo_hora
                total_he25 += line.he_25 or 0
                total_he50 += line.he_50 or 0
                total_he75 += line.he_75 or 0
                total_he100 += line.he_100 or 0
                total_extras += total_extras_line
                total_m25 += line.amount_25 or 0
                total_m50 += line.amount_50 or 0
                total_m75 += line.amount_75 or 0
                total_m100 += line.amount_100 or 0
                total_monto += line.amount_total or 0

            # Fila totales departamento (gris oscuro)
            ws_he.cell(row=row, column=1, value=dept_name).fill = dept_summary_fill
            ws_he.cell(row=row, column=1).font = summary_font
            ws_he.cell(row=row, column=2, value=len(lines)).fill = dept_summary_fill
            ws_he.cell(row=row, column=2).font = summary_font
            ws_he.cell(row=row, column=3, value=round(total_salario, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=3).font = summary_font
            ws_he.cell(row=row, column=4, value=round(total_sueldo_hora, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=4).font = summary_font
            ws_he.cell(row=row, column=5, value=0).fill = dept_summary_fill
            ws_he.cell(row=row, column=6, value=0).fill = dept_summary_fill
            ws_he.cell(row=row, column=7, value=0).fill = dept_summary_fill
            ws_he.cell(row=row, column=8, value=0).fill = dept_summary_fill
            ws_he.cell(row=row, column=9, value=round(total_extras, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=9).font = summary_font
            ws_he.cell(row=row, column=10, value=round(total_m25, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=10).font = summary_font
            ws_he.cell(row=row, column=11, value=round(total_m50, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=11).font = summary_font
            ws_he.cell(row=row, column=12, value=round(total_m75, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=12).font = summary_font
            ws_he.cell(row=row, column=13, value=round(total_m100, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=13).font = summary_font
            ws_he.cell(row=row, column=14, value=round(total_monto, 2)).fill = dept_summary_fill
            ws_he.cell(row=row, column=14).font = summary_font
            row += 1

            g_salario += total_salario
            g_sueldo_hora += total_sueldo_hora
            g_he25 += total_he25
            g_he50 += total_he50
            g_he75 += total_he75
            g_he100 += total_he100
            g_extras += total_extras
            g_m25 += total_m25
            g_m50 += total_m50
            g_m75 += total_m75
            g_m100 += total_m100
            g_monto += total_monto
            g_empleados += len(lines)

        # Fila total general HE
        total_fill = PatternFill(start_color='2F4F4F', end_color='2F4F4F', fill_type='solid')
        ws_he.cell(row=row, column=1, value=_('TOTAL')).fill = total_fill
        ws_he.cell(row=row, column=1).font = summary_font
        ws_he.cell(row=row, column=2, value=g_empleados).fill = total_fill
        ws_he.cell(row=row, column=2).font = summary_font
        ws_he.cell(row=row, column=3, value=round(g_salario, 2)).fill = total_fill
        ws_he.cell(row=row, column=3).font = summary_font
        ws_he.cell(row=row, column=4, value=round(g_sueldo_hora, 2)).fill = total_fill
        ws_he.cell(row=row, column=4).font = summary_font
        ws_he.cell(row=row, column=5, value=round(g_he25, 2)).fill = total_fill
        ws_he.cell(row=row, column=5).font = summary_font
        ws_he.cell(row=row, column=6, value=round(g_he50, 2)).fill = total_fill
        ws_he.cell(row=row, column=6).font = summary_font
        ws_he.cell(row=row, column=7, value=round(g_he75, 2)).fill = total_fill
        ws_he.cell(row=row, column=7).font = summary_font
        ws_he.cell(row=row, column=8, value=round(g_he100, 2)).fill = total_fill
        ws_he.cell(row=row, column=8).font = summary_font
        ws_he.cell(row=row, column=9, value=round(g_extras, 2)).fill = total_fill
        ws_he.cell(row=row, column=9).font = summary_font
        ws_he.cell(row=row, column=10, value=round(g_m25, 2)).fill = total_fill
        ws_he.cell(row=row, column=10).font = summary_font
        ws_he.cell(row=row, column=11, value=round(g_m50, 2)).fill = total_fill
        ws_he.cell(row=row, column=11).font = summary_font
        ws_he.cell(row=row, column=12, value=round(g_m75, 2)).fill = total_fill
        ws_he.cell(row=row, column=12).font = summary_font
        ws_he.cell(row=row, column=13, value=round(g_m100, 2)).fill = total_fill
        ws_he.cell(row=row, column=13).font = summary_font
        ws_he.cell(row=row, column=14, value=round(g_monto, 2)).fill = total_fill
        ws_he.cell(row=row, column=14).font = summary_font
        row += 1

        # Hoja TE agrupada por departamento
        ws_te = wb.create_sheet('Turnos Extra')
        te_headers = ['Departamento', 'Empleado', 'Turno A', 'Turno B', 'Turno C', 'Otros Ingresos']
        for c, h in enumerate(te_headers, 1):
            cell = ws_te.cell(row=1, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
        te_by_dept = {}
        for line in self.line_te_ids:
            dept_name = (line.department_id and line.department_id.name) or _('Sin departamento')
            te_by_dept.setdefault(dept_name, []).append(line)
        g_te_turn_a = g_te_turn_b = g_te_turn_c = g_te_monto = 0.0
        g_te_empleados = 0
        row_te = 2
        for dept_name in sorted(te_by_dept.keys()):
            lines_te = te_by_dept[dept_name]
            ws_te.cell(row=row_te, column=1, value=dept_name).fill = dept_subheader_fill
            for c in range(2, 7):
                ws_te.cell(row=row_te, column=c).fill = dept_subheader_fill
            row_te += 1
            total_te = total_ta = total_tb = total_tc = 0.0
            for line in lines_te:
                ws_te.cell(row=row_te, column=1, value=dept_name)
                ws_te.cell(row=row_te, column=2, value=line.employee_id.name or '')
                ta, tb, tc = line.turn_a or 0, line.turn_b or 0, line.turn_c or 0
                amt = line.amount_other_income or 0
                ws_te.cell(row=row_te, column=3, value=ta)
                ws_te.cell(row=row_te, column=4, value=tb)
                ws_te.cell(row=row_te, column=5, value=tc)
                ws_te.cell(row=row_te, column=6, value=amt)
                total_te += amt
                total_ta += ta
                total_tb += tb
                total_tc += tc
                row_te += 1
            ws_te.cell(row=row_te, column=1, value=dept_name).fill = dept_summary_fill
            ws_te.cell(row=row_te, column=1).font = summary_font
            ws_te.cell(row=row_te, column=2, value=len(lines_te)).fill = dept_summary_fill
            ws_te.cell(row=row_te, column=2).font = summary_font
            ws_te.cell(row=row_te, column=3, value=total_ta).fill = dept_summary_fill
            ws_te.cell(row=row_te, column=4, value=total_tb).fill = dept_summary_fill
            ws_te.cell(row=row_te, column=5, value=total_tc).fill = dept_summary_fill
            ws_te.cell(row=row_te, column=6, value=round(total_te, 2)).fill = dept_summary_fill
            ws_te.cell(row=row_te, column=6).font = summary_font
            row_te += 1
            g_te_turn_a += total_ta
            g_te_turn_b += total_tb
            g_te_turn_c += total_tc
            g_te_monto += total_te
            g_te_empleados += len(lines_te)

        # Fila total general TE
        total_fill_te = PatternFill(start_color='2F4F4F', end_color='2F4F4F', fill_type='solid')
        ws_te.cell(row=row_te, column=1, value=_('TOTAL')).fill = total_fill_te
        ws_te.cell(row=row_te, column=1).font = summary_font
        ws_te.cell(row=row_te, column=2, value=g_te_empleados).fill = total_fill_te
        ws_te.cell(row=row_te, column=2).font = summary_font
        ws_te.cell(row=row_te, column=3, value=g_te_turn_a).fill = total_fill_te
        ws_te.cell(row=row_te, column=3).font = summary_font
        ws_te.cell(row=row_te, column=4, value=g_te_turn_b).fill = total_fill_te
        ws_te.cell(row=row_te, column=4).font = summary_font
        ws_te.cell(row=row_te, column=5, value=g_te_turn_c).fill = total_fill_te
        ws_te.cell(row=row_te, column=5).font = summary_font
        ws_te.cell(row=row_te, column=6, value=round(g_te_monto, 2)).fill = total_fill_te
        ws_te.cell(row=row_te, column=6).font = summary_font

        # Hoja BD agrupada por departamento (aplica a la quincena de la planilla)
        ws_bd = wb.create_sheet('Beneficios y Deducciones')
        bd_headers = ['Departamento', 'Empleado', 'Concepto', 'Monto']
        for c, h in enumerate(bd_headers, 1):
            cell = ws_bd.cell(row=1, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
        bd_by_dept = {}
        for line in self.line_bd_ids:
            dept_name = (line.department_id and line.department_id.name) or _('Sin departamento')
            bd_by_dept.setdefault(dept_name, []).append(line)
        g_bd_lines = 0
        g_bd_monto = 0.0
        row_bd = 2
        for dept_name in sorted(bd_by_dept.keys()):
            lines_bd = bd_by_dept[dept_name]
            ws_bd.cell(row=row_bd, column=1, value=dept_name).fill = dept_subheader_fill
            for c in range(2, 5):
                ws_bd.cell(row=row_bd, column=c).fill = dept_subheader_fill
            row_bd += 1
            total_bd = 0.0
            for line in lines_bd:
                ws_bd.cell(row=row_bd, column=1, value=dept_name)
                ws_bd.cell(row=row_bd, column=2, value=line.employee_id.name or '')
                ws_bd.cell(row=row_bd, column=3, value=line.concept_id.name or '')
                ws_bd.cell(row=row_bd, column=4, value=line.amount or 0)
                total_bd += line.amount or 0
                row_bd += 1
            ws_bd.cell(row=row_bd, column=1, value=dept_name).fill = dept_summary_fill
            ws_bd.cell(row=row_bd, column=1).font = summary_font
            ws_bd.cell(row=row_bd, column=2, value=len(lines_bd)).fill = dept_summary_fill
            ws_bd.cell(row=row_bd, column=2).font = summary_font
            ws_bd.cell(row=row_bd, column=4, value=round(total_bd, 2)).fill = dept_summary_fill
            ws_bd.cell(row=row_bd, column=4).font = summary_font
            row_bd += 1
            g_bd_lines += len(lines_bd)
            g_bd_monto += total_bd

        # Fila total general BD
        total_fill_bd = PatternFill(start_color='2F4F4F', end_color='2F4F4F', fill_type='solid')
        for c in range(1, 5):
            ws_bd.cell(row=row_bd, column=c).fill = total_fill_bd
        ws_bd.cell(row=row_bd, column=1, value=_('TOTAL')).font = summary_font
        ws_bd.cell(row=row_bd, column=2, value=g_bd_lines).font = summary_font
        ws_bd.cell(row=row_bd, column=4, value=round(g_bd_monto, 2)).font = summary_font

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f'Planilla_{self.name.replace(" ", "_")}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }


class L10nHnPayrollLineHe(models.Model):
    """Línea de Horas Extra en Planilla"""
    _name = 'l10n.hn.payroll.line.he'
    _description = 'Planilla Línea HE'

    parent_id = fields.Many2one('l10n.hn.payroll.sheet', string='Planilla', required=True, ondelete='cascade')
    employee_id = fields.Many2one('hr.employee', string='Empleado', required=True, ondelete='cascade')
    version_id = fields.Many2one(
        'hr.version', string='Contrato',
        compute='_compute_version_id', store=True, readonly=True,
        help='Contrato activo del empleado en el período (calculado automáticamente)',
    )
    department_id = fields.Many2one('hr.department', related='employee_id.department_id', readonly=True)
    he_25 = fields.Float(string='HE 25%', digits=(10, 2), default=0)
    he_50 = fields.Float(string='HE 50%', digits=(10, 2), default=0)
    he_75 = fields.Float(string='HE 75%', digits=(10, 2), default=0)
    he_100 = fields.Float(string='HE 100%', digits=(10, 2), default=0)
    amount_25 = fields.Float(string='Monto 25%', digits=(12, 2), compute='_compute_amounts', store=True)
    amount_50 = fields.Float(string='Monto 50%', digits=(12, 2), compute='_compute_amounts', store=True)
    amount_75 = fields.Float(string='Monto 75%', digits=(12, 2), compute='_compute_amounts', store=True)
    amount_100 = fields.Float(string='Monto 100%', digits=(12, 2), compute='_compute_amounts', store=True)
    amount_total = fields.Float(string='Total HE', digits=(12, 2), compute='_compute_amounts', store=True)

    @api.depends('employee_id', 'parent_id', 'parent_id.date_start', 'parent_id.date_end')
    def _compute_version_id(self):
        for line in self:
            if line.employee_id and line.parent_id:
                line.version_id = line.parent_id._get_version_for_employee(line.employee_id)
            else:
                line.version_id = False

    @api.depends('he_25', 'he_50', 'he_75', 'he_100', 'version_id', 'version_id.wage')
    def _compute_amounts(self):
        for line in self:
            wage = 0.0
            if line.version_id:
                wage = line.version_id.wage or 0.0
            wage_by_hours = wage / (30 * 8) if wage else 0.0
            line.amount_25 = round((wage_by_hours * 1.25) * line.he_25, 2)
            line.amount_50 = round((wage_by_hours * 1.50) * line.he_50, 2)
            line.amount_75 = round((wage_by_hours * 1.75) * line.he_75, 2)
            line.amount_100 = round((wage_by_hours * 2.0) * line.he_100, 2)
            line.amount_total = line.amount_25 + line.amount_50 + line.amount_75 + line.amount_100

    @api.constrains('employee_id', 'parent_id')
    def _check_unique_employee(self):
        for line in self:
            dup = self.search([
                ('parent_id', '=', line.parent_id.id),
                ('employee_id', '=', line.employee_id.id),
                ('id', '!=', line.id),
            ], limit=1)
            if dup:
                raise ValidationError(_('No puede haber dos líneas con el mismo empleado en la misma planilla.'))


class L10nHnPayrollLineTe(models.Model):
    """Línea de Turnos Extra en Planilla"""
    _name = 'l10n.hn.payroll.line.te'
    _description = 'Planilla Línea TE'

    parent_id = fields.Many2one('l10n.hn.payroll.sheet', string='Planilla', required=True, ondelete='cascade')
    employee_id = fields.Many2one('hr.employee', string='Empleado', required=True, ondelete='cascade')
    version_id = fields.Many2one(
        'hr.version', string='Contrato',
        compute='_compute_version_id', store=True, readonly=True,
        help='Contrato activo del empleado en el período (calculado automáticamente)',
    )
    department_id = fields.Many2one('hr.department', related='employee_id.department_id', readonly=True)
    turn_a = fields.Integer(string='Turno A', default=0)
    turn_b = fields.Integer(string='Turno B', default=0)
    turn_c = fields.Integer(string='Turno C', default=0)
    amount_other_income = fields.Float(string='Otros Ingresos', digits=(12, 2), compute='_compute_amount', store=True)

    @api.depends('employee_id', 'parent_id', 'parent_id.date_start', 'parent_id.date_end')
    def _compute_version_id(self):
        for line in self:
            if line.employee_id and line.parent_id:
                line.version_id = line.parent_id._get_version_for_employee(line.employee_id)
            else:
                line.version_id = False

    @api.depends('turn_a', 'turn_b', 'turn_c', 'version_id', 'version_id.wage', 'parent_id', 'parent_id.company_id')
    def _compute_amount(self):
        """TE = (Días por turno) × (Sueldo diario) × Recargo del turno"""
        TurnConfig = self.env['l10n.hn.turn.extra.config']
        for line in self:
            wage = 0.0
            if line.version_id:
                wage = line.version_id.wage or 0.0
            sueldo_diario = wage / 30.0 if wage else 0.0
            company = line.parent_id.company_id if line.parent_id else self.env.company
            recargo_a, recargo_b, recargo_c = TurnConfig.get_factors(company)
            # (Días por turno) * (Sueldo diario) * Recargo del turno
            total = (
                (line.turn_a or 0) * sueldo_diario * recargo_a
                + (line.turn_b or 0) * sueldo_diario * recargo_b
                + (line.turn_c or 0) * sueldo_diario * recargo_c
            )
            line.amount_other_income = round(total, 2)


class L10nHnPayrollLineBd(models.Model):
    """Línea de Beneficios y Deducciones en Planilla"""
    _name = 'l10n.hn.payroll.line.bd'
    _description = 'Planilla Línea BD'

    parent_id = fields.Many2one('l10n.hn.payroll.sheet', string='Planilla', required=True, ondelete='cascade')
    employee_id = fields.Many2one('hr.employee', string='Empleado', required=True, ondelete='cascade')
    version_id = fields.Many2one(
        'hr.version', string='Contrato',
        compute='_compute_version_id', store=True, readonly=True,
        help='Contrato activo del empleado en el período (calculado automáticamente)',
    )
    department_id = fields.Many2one('hr.department', related='employee_id.department_id', readonly=True)
    concept_id = fields.Many2one(
        'hr.employee.recurring.concept',
        string='Concepto',
        domain="['&', '&', ('for_payroll', '=', True), ('employee_id', '=', False), '|', "
        "('company_id', '=', False), ('company_id', '=', parent.company_id)]",
        help='Plantilla compartida (sin empleado) marcada para planilla: misma compañía que la planilla o sin compañía (todas).',
    )
    amount = fields.Float(string='Monto', digits=(12, 2), required=True, default=0.0)
    include_in = fields.Selection([
        ('first', 'Primera Quincena'),
        ('second', 'Segunda Quincena'),
        ('both', 'Ambas'),
    ], string='Incluir En', compute='_compute_include_in', store=True, readonly=True,
       help='Se establece automáticamente según la quincena de la planilla.')

    @api.depends('employee_id', 'parent_id', 'parent_id.date_start', 'parent_id.date_end')
    def _compute_version_id(self):
        for line in self:
            if line.employee_id and line.parent_id:
                line.version_id = line.parent_id._get_version_for_employee(line.employee_id)
            else:
                line.version_id = False

    @api.depends('parent_id', 'parent_id.payroll_period_id', 'parent_id.payroll_period_id.period_number')
    def _compute_include_in(self):
        for line in self:
            if line.parent_id and line.parent_id.payroll_period_id:
                period = line.parent_id.payroll_period_id
                line.include_in = 'first' if period.period_number == '1' else 'second'
            else:
                line.include_in = 'both'
