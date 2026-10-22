# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik

import base64
import io
import uuid
from datetime import date, datetime
from dateutil.relativedelta import relativedelta

from odoo import api, fields, models, _
from odoo.exceptions import ValidationError, UserError

from odoo.addons.gn_hr_payroll_employee_extras.models.payroll_biweekly_utils import (
    biweekly_date_range,
    period_number_from_day,
)


def _l10n_hn_xlsx_apply_table(ws, num_cols, header_row=1):
    """Formato tabla de Excel (openpyxl) sobre el rango de datos."""
    try:
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError:
        return
    if ws.max_row < header_row:
        return
    last_col = get_column_letter(num_cols)
    last_row = ws.max_row
    ref = 'A%d:%s%d' % (header_row, last_col, last_row)
    tab = Table(displayName='tbl_%s' % uuid.uuid4().hex[:12], ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium2',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)


class HrEmployeeContractHistory(models.Model):
    """Contract History - Track salary changes and contract versions"""
    _name = 'hr.employee.contract.history'
    _description = 'Employee Contract History'
    _order = 'date_start desc'

    name = fields.Char(string='Name', compute='_compute_name', store=True)
    employee_id = fields.Many2one('hr.employee', string='Employee', required=True, ondelete='cascade')
    
    # Contract period
    date_start = fields.Date(string='Start Date', required=True)
    date_end = fields.Date(string='End Date',
                           help='Leave empty for current contract')
    
    # Salary information
    wage = fields.Float(string='Wage', digits=(12, 2), required=True)
    currency_id = fields.Many2one('res.currency', string='Currency', required=True,
                                   default=lambda self: self.env.company.currency_id)
    secondary_currency_id = fields.Many2one('res.currency', string='Secondary Currency',
                                             help='Optional secondary currency')
    secondary_wage = fields.Float(string='Secondary Wage', digits=(12, 2))
    
    # Payment periodicity
    periodicity = fields.Selection([
        ('weekly', 'Weekly'),
        ('biweekly', 'Biweekly'),
        ('monthly', 'Monthly'),
    ], string='Payment Periodicity', required=True, default='monthly')
    
    # Job information
    job_id = fields.Many2one('hr.job', string='Job Position')
    department_id = fields.Many2one('hr.department', string='Department')
    
    # Notes
    notes = fields.Text(string='Notes')
    
    active = fields.Boolean(string='Active', default=True)
    
    @api.depends('employee_id', 'date_start', 'wage')
    def _compute_name(self):
        for rec in self:
            rec.name = f'{rec.employee_id.name or ""} - {rec.date_start} - {rec.wage}'


class HrEmployee(models.Model):
    """Extend hr.employee for Honduras payroll localization"""
    _inherit = 'hr.employee'

    disabled = fields.Boolean(
        readonly=False,
        related='version_id.disabled',
        inherited=True,
        groups='l10n_hn_hr_payroll.group_hr_payroll_user',
        string='Discapacidad',
    )
    disability_line_ids = fields.One2many(
        'hr.employee.disability.line',
        'employee_id',
        string='Detalle de discapacidad',
        groups='l10n_hn_hr_payroll.group_hr_payroll_user',
    )

    @api.model
    def _selection_yes_no(self):
        """Sí/No traducible (evita _('…') en lista estática del campo al importar)."""
        return [('yes', _('Sí')), ('not', _('No'))]

    seniority_display = fields.Char(
        string='Antigüedad (desde contrato)',
        compute='_compute_seniority_display',
        store=False,
    )

    @api.depends('version_ids', 'version_ids.contract_date_start')
    def _compute_seniority_display(self):
        today = fields.Date.today()
        for emp in self:
            start = emp.sudo()._get_first_contract_date(no_gap=False)
            if not start:
                emp.seniority_display = False
                continue
            diff = relativedelta(today, start)
            parts = []
            if diff.years:
                parts.append('%d año%s' % (diff.years, 's' if diff.years != 1 else ''))
            if diff.months:
                parts.append('%d mes%s' % (diff.months, 'es' if diff.months != 1 else ''))
            if diff.days or not parts:
                parts.append('%d día%s' % (diff.days, 's' if diff.days != 1 else ''))
            emp.seniority_display = ', '.join(parts)

    hn_company_country_code = fields.Char(
        string='Código país empresa (HN)',
        compute='_compute_hn_company_country_code',
        readonly=True,
    )

    @api.depends(
        'company_id',
        'company_id.partner_id.country_id',
        'company_id.partner_id.country_id.code',
    )
    def _compute_hn_company_country_code(self):
        for emp in self:
            partner = emp.company_id.partner_id
            country = partner.country_id if partner else False
            emp.hn_company_country_code = country.code if country else ''

    l10n_hn_monthly_salary_history_ids = fields.One2many(
        'l10n.hn.employee.monthly.salary.history',
        'employee_id',
        string='Historial salarial mensual',
    )

    @api.depends('l10n_hn_monthly_salary_history_ids')
    def _compute_l10n_hn_monthly_salary_history_count(self):
        for emp in self:
            emp.l10n_hn_monthly_salary_history_count = len(emp.l10n_hn_monthly_salary_history_ids)

    def action_open_l10n_hn_monthly_salary_history(self):
        self.ensure_one()
        action = self.env['ir.actions.act_window']._for_xml_id(
            'l10n_hn_hr_payroll.action_l10n_hn_employee_monthly_salary_history'
        )
        action['domain'] = [('employee_id', '=', self.id)]
        action['context'] = {'default_employee_id': self.id, 'search_default_employee_id': self.id}
        return action

    current_date_version = fields.Date(
        string='Fecha de firma de contrato',
        compute='_compute_current_version_fields',
        inverse='_inverse_current_date_version',
        groups='hr.group_hr_user',
    )
    current_contract_date_end = fields.Date(
        string='Fecha de vencimiento',
        compute='_compute_current_version_fields',
        inverse='_inverse_current_contract_date_end',
        groups='hr.group_hr_user',
    )

    @api.depends('current_version_id')
    def _compute_current_version_fields(self):
        for emp in self:
            version = emp.current_version_id.sudo() if emp.current_version_id else False
            emp.current_date_version = version.date_version if version else False
            emp.current_contract_date_end = version.contract_date_end if version else False

    def _inverse_current_date_version(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().date_version = emp.current_date_version

    def _inverse_current_contract_date_end(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().contract_date_end = emp.current_contract_date_end

    l10n_hn_show_edit_initial_contract_wizard = fields.Boolean(
        string="Mostrar asistente contrato inicial (HN)",
        compute="_compute_l10n_hn_show_edit_initial_contract_wizard",
    )

    @api.depends("hn_company_country_code", "version_ids", "company_id", "company_id.country_id")
    def _compute_l10n_hn_show_edit_initial_contract_wizard(self):
        Payslip = self.env["hr.payslip"].sudo()
        for emp in self:
            if (
                emp.hn_company_country_code == "HN"
                and len(emp.version_ids) == 1
            ):
                emp.l10n_hn_show_edit_initial_contract_wizard = not bool(
                    Payslip.search_count([("employee_id", "=", emp.id)])
                )
            else:
                emp.l10n_hn_show_edit_initial_contract_wizard = False

    def _l10n_edit_initial_contract_eligible(self, country_code):
        self.ensure_one()
        if not self.env.user.has_group(
            "gn_hr_payroll_employee_extras.group_gn_edit_initial_contract"
        ):
            return False
        if self.env["hr.payslip"].sudo().search_count([("employee_id", "=", self.id)]):
            return False
        if len(self.version_ids) != 1:
            return False
        if country_code == "HN":
            return self.hn_company_country_code == "HN"
        return False

    def l10n_hn_action_open_edit_initial_contract_wizard(self):
        return self._l10n_action_open_edit_initial_contract_wizard("HN")

    def _l10n_action_open_edit_initial_contract_wizard(self, country_code):
        self.ensure_one()
        if not self._l10n_edit_initial_contract_eligible(country_code):
            raise UserError(
                _(
                    "Solo se puede abrir el asistente con un único contrato, sin recibos de nómina, "
                    "compañía del país adecuado, y con el permiso «Nómina: editar contrato inicial (asistente)»."
                )
            )
        return {
            "type": "ir.actions.act_window",
            "name": _("Editar contrato inicial"),
            "res_model": "l10n.hn.edit.initial.contract.wizard",
            "view_mode": "form",
            "target": "new",
            "context": {**self.env.context, "default_employee_id": self.id},
        }

    include_in_current_payroll = fields.Boolean(
        string='Incluir en nómina de quincena en curso',
        default=False,
        help='Si el empleado fue dado de baja a mitad de quincena, marcar para incluirlo en la '
             'nómina del período en curso. Se limpia automáticamente al generar su recibo.',
    )

    contract_expires_this_month = fields.Boolean(
        string='Contrato vence este mes',
        compute='_compute_contract_expires_this_month',
        search='_search_contract_expires_this_month',
    )

    @api.depends('current_version_id', 'current_version_id.contract_date_end')
    def _compute_contract_expires_this_month(self):
        today = fields.Date.context_today(self)
        start_of_month = today.replace(day=1)
        end_of_month = start_of_month + relativedelta(months=1) - relativedelta(days=1)
        for emp in self:
            version = emp.current_version_id.sudo() if emp.current_version_id else False
            end_date = version.contract_date_end if version else None
            emp.contract_expires_this_month = bool(
                end_date and start_of_month <= end_date <= end_of_month
            )

    @api.model
    def _search_contract_expires_this_month(self, operator, value):
        """Dominio para filtrar empleados cuyo contrato vence este mes."""
        today = fields.Date.context_today(self)
        start_of_month = today.replace(day=1)
        end_of_month = start_of_month + relativedelta(months=1) - relativedelta(days=1)
        if (operator, value) in [('=', True), ('!=', False)]:
            return [
                ('current_contract_date_end', '!=', False),
                ('current_contract_date_end', '>=', start_of_month),
                ('current_contract_date_end', '<=', end_of_month),
            ]
        return [
            '|',
            ('current_contract_date_end', '=', False),
            '|',
            ('current_contract_date_end', '<', start_of_month),
            ('current_contract_date_end', '>', end_of_month),
        ]

    @api.model
    def _cron_trial_period_alert(self):
        """Cron diario (HN): envía correo cuando un empleado cumple los días de período de prueba."""
        country = self.env.ref("base.hn", raise_if_not_found=False)
        if not country:
            return
        # country_id en res.company es computed (no almacenado): filtrar por partner.
        companies = self.env["res.company"].search([("partner_id.country_id", "=", country.id)])
        if not companies:
            return
        template = self.env.ref(
            "l10n_hn_hr_payroll.mail_template_trial_period_alert",
            raise_if_not_found=False,
        )
        if not template:
            return
        today = fields.Date.context_today(self)
        for company in companies:
            trial_days = company._get_localization_trial_period_days()
            target_date = today - relativedelta(days=trial_days)
            employees = self.search([
                ("company_id", "=", company.id),
                ("active", "=", True),
                ("trial_period_notified", "=", False),
            ])
            for emp in employees:
                start = emp._get_first_contract_date(no_gap=False)
                if not start or start != target_date:
                    continue
                recipients = []
                if emp.parent_id and emp.parent_id.work_email:
                    recipients.append(emp.parent_id.work_email)
                for user in company.l10n_hn_hr_alert_recipient_ids:
                    if user.email and user.email not in recipients:
                        recipients.append(user.email)
                if not recipients:
                    continue
                try:
                    template.with_context(
                        trial_days=trial_days,
                        email_to=",".join(recipients),
                    ).send_mail_payroll_tracked(emp.id, force_send=True)
                except Exception:
                    import logging
                    logging.getLogger(__name__).warning(
                        "Failed to send trial period alert for employee %s",
                        emp.display_name, exc_info=True
                    )
                emp.trial_period_notified = True

    def _get_hr_notification_user_id(self, notification_type):
        """
        Obtiene el user_id al que asignar una notificación del tipo indicado.
        Prioridad: hr_responsible_id > hr.notification.recipient > hr_payroll_manager_id > uid
        """
        self.ensure_one()
        if self.hr_responsible_id:
            return self.hr_responsible_id.id
        company = self.company_id
        if not company:
            return self.env.uid
        recipient_ids = self.env['hr.notification.recipient'].get_recipient_user_ids(
            company.id, notification_type)
        if recipient_ids:
            return recipient_ids[0]
        if hasattr(company, 'hr_payroll_manager_id') and company.hr_payroll_manager_id:
            return company.hr_payroll_manager_id.id
        return self.env.uid

    @api.model
    def notify_expiring_contract_work_permit(self):
        """Override para usar hr.notification.recipient cuando no hay hr_responsible_id."""
        companies = self.env['res.company'].search([])
        employees_contract_expiring = self.env['hr.employee']
        employees_work_permit_expiring = self.env['hr.employee']

        for company in companies:
            employees_contract_expiring += self.env['hr.employee'].search([
                ('company_id', '=', company.id),
                ('contract_date_start', '!=', False),
                ('contract_date_start', '<', fields.Date.today()),
                ('contract_date_end', '=', fields.Date.today() + relativedelta(days=company.contract_expiration_notice_period)),
            ])

            employees_work_permit_expiring += self.env['hr.employee'].search([
                ('company_id', '=', company.id),
                ('work_permit_expiration_date', '!=', False),
                ('work_permit_expiration_date', '=', fields.Date.today() + relativedelta(days=company.work_permit_expiration_notice_period)),
            ])

        for employee in employees_contract_expiring:
            user_id = employee._get_hr_notification_user_id('contract_expiring')
            employee.with_context(mail_activity_quick_update=True).activity_schedule(
                'mail.mail_activity_data_todo', employee.contract_date_end,
                _("The contract of %s is about to expire.", employee.name),
                user_id=user_id)

        for employee in employees_work_permit_expiring:
            user_id = employee._get_hr_notification_user_id('work_permit_expiring')
            employee.with_context(mail_activity_quick_update=True).activity_schedule(
                'mail.mail_activity_data_todo', employee.work_permit_expiration_date,
                _("The work permit of %s is about to expire.", employee.name),
                user_id=user_id)

        return True

    def _prevent_portal_leave_manager(self, vals):
        """If parent_id or leave_manager_id points to a portal user, prevent
        hr_holidays from auto-assigning them as leave_manager_id (which would
        add group_user and conflict with group_portal).  Returns modified vals."""
        portal_group = self.env.ref('base.group_portal', raise_if_not_found=False)
        if not portal_group:
            return vals

        # Check leave_manager_id directly (form onchange may send it)
        if vals.get('leave_manager_id'):
            user = self.env['res.users'].browse(vals['leave_manager_id'])
            if portal_group in user.group_ids:
                vals = dict(vals, leave_manager_id=False)
                return vals

        # Check parent_id → manager user
        if vals.get('parent_id'):
            manager_user = self.env['hr.employee'].browse(vals['parent_id']).user_id
            if manager_user and portal_group in manager_user.group_ids:
                vals = dict(vals, leave_manager_id=False)

        return vals

    def write(self, vals):
        vals = self._prevent_portal_leave_manager(vals)
        if vals.get('has_allergies') is False or vals.get('relatives_in_company') == 'not':
            vals = dict(vals)
            if vals.get('has_allergies') is False:
                vals['allergy_ids'] = [(5, 0, 0)]
            if vals.get('relatives_in_company') == 'not':
                vals['in_company_relative_ids'] = [(5, 0, 0)]
        res = super().write(vals)
        for emp in self:
            if not emp.disabled and emp.disability_line_ids:
                emp.disability_line_ids.unlink()
        return res

    @api.model_create_multi
    def create(self, vals_list):
        """Create employee and set up default deductions when first version is created.

        Filter version_ids to avoid duplicate date_version: when the form sends both
        the main contract fields (wage, contract_date_start) and version_ids with a
        new line (0, 0, {...}) that has the same date_version (e.g. today), we would
        create two versions with the same date. We remove the redundant version_ids
        line (typically the empty inline with wage=0) to prevent the unique constraint violation.
        """
        vals_list = [self._prevent_portal_leave_manager(vals) for vals in vals_list]

        today = fields.Date.today()
        for vals in vals_list:
            version_ids = vals.get('version_ids') or []
            if not version_ids:
                continue

            # Filter out (0, 0, {...}) or (0, 'virtual_XXX', {...}) entries that would create duplicate date_version
            # and appear to be the empty default line (wage=0) from the Honduras tab.
            # Odoo web form uses (0, 'virtual_XXX', {...}) for new inline records, not (0, 0, {...}).
            filtered_ids = []
            for cmd in version_ids:
                is_new_create = cmd[0] == 0 and (cmd[1] == 0 or isinstance(cmd[1], str)) and isinstance(cmd[2], dict) if len(cmd) >= 3 else False
                if is_new_create:
                    sub_vals = cmd[2]
                    sub_date = sub_vals.get('date_version')
                    sub_wage = sub_vals.get('wage', 0) or 0
                    if sub_date and sub_wage == 0:
                        sub_date = fields.Date.from_string(sub_date) if isinstance(sub_date, str) else sub_date
                        if sub_date == today:
                            # Skip empty duplicate line - main form version will be used
                            continue
                filtered_ids.append(cmd)

            if filtered_ids != version_ids:
                vals['version_ids'] = filtered_ids

        employees = super().create(vals_list)

        # When a version is created for these employees, deductions will be automatically linked
        # See hr_version.py _create_default_deductions method
        # The linking happens automatically when the version is created

        return employees
    
    def action_create_default_deductions(self):
        """Action to manually create default deductions for current version"""
        self.ensure_one()
        if not self.current_version_id:
            raise ValidationError(_('El empleado debe tener una versión de contrato activa.'))
        
        self.current_version_id._create_default_deductions()
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Deducciones Creadas'),
                'message': _('Las deducciones se han vinculado a la versión del contrato.'),
                'type': 'success',
                'sticky': False,
            }
        }

    def _write_employees_sheet(self, ws, employees, SEX_LABELS, MARITAL_LABELS,
                               header_fill, header_font, header_alignment):
        """Escribe una hoja con el listado de empleados (formato simple)."""
        headers = [
            'DNI', 'Nombre Completo', 'Sexo', 'Lugar de Residencia', 'Estado Civil',
            'Profesión u oficio', 'Número de Teléfono', 'Correo Electrónico', 'Área',
            'Nombre Completo del Jefe Inmediato', 'Salario', 'Cargo', 'Fecha de Inicio',
            'Fecha de terminación', 'Fecha de firma contrato', 'Tipo de contrato',
            'Días para vencer',
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        for row, emp in enumerate(employees, 2):
            version = emp.current_version_id.sudo() if emp.current_version_id else False
            ws.cell(row=row, column=1, value=emp.identification_id or '')
            ws.cell(row=row, column=2, value=emp.name or '')
            ws.cell(row=row, column=3, value=SEX_LABELS.get(emp.sex, emp.sex or ''))
            ws.cell(row=row, column=4, value=emp.private_city or '')
            ws.cell(row=row, column=5, value=MARITAL_LABELS.get(emp.marital, emp.marital or ''))
            ws.cell(row=row, column=6, value=emp.job_title or (emp.job_id.name if emp.job_id else ''))
            ws.cell(row=row, column=7, value=emp.work_phone or '')
            ws.cell(row=row, column=8, value=emp.work_email or '')
            ws.cell(row=row, column=9, value=emp.department_id.name if emp.department_id else '')
            ws.cell(row=row, column=10, value=emp.parent_id.name if emp.parent_id else '')
            wage = version.wage if version else 0
            ws.cell(row=row, column=11, value=float(wage))
            ws.cell(row=row, column=12, value=emp.job_id.name if emp.job_id else '')
            ws.cell(row=row, column=13, value=version.contract_date_start.strftime('%Y-%m-%d') if version and version.contract_date_start else '')
            ws.cell(row=row, column=14, value=version.contract_date_end.strftime('%Y-%m-%d') if version and version.contract_date_end else '')
            ws.cell(row=row, column=15, value=emp.current_date_version.strftime('%Y-%m-%d') if emp.current_date_version else '')
            ws.cell(row=row, column=16, value=version.contract_type_id.name if version and version.contract_type_id else '')
            ws.cell(row=row, column=17, value=emp.days_to_contract_expiration if emp.days_to_contract_expiration else '')

        from openpyxl.utils import get_column_letter
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        _l10n_hn_xlsx_apply_table(ws, len(headers), header_row=1)

    def action_export_employees_excel(self):
        """Export employee list to Excel file.
        Hoja 1: Todos los empleados. Hojas adicionales: una por tipo de contrato (solo si hay empleados).
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))

        active_ids = self.env.context.get('active_ids', [])
        if active_ids:
            employees = self.browse(active_ids).exists()
        else:
            employees = self.search([('company_id', 'in', self.env.companies.ids)], order='name')

        if not employees:
            raise UserError(_('No hay empleados para exportar.'))

        SEX_LABELS = {'male': 'Masculino', 'female': 'Femenino', 'other': 'Otro'}
        MARITAL_LABELS = {
            'single': 'Soltero(a)', 'married': 'Casado(a)', 'cohabitant': 'Unión Libre',
            'widower': 'Viudo(a)', 'divorced': 'Divorciado(a)', 'other': 'Otro',
        }

        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        header_alignment = Alignment(horizontal='center', vertical='center')

        wb = Workbook()
        ws_all = wb.active
        ws_all.title = 'Empleados'
        self._write_employees_sheet(ws_all, employees, SEX_LABELS, MARITAL_LABELS,
                                    header_fill, header_font, header_alignment)

        by_contract_type = {}
        for emp in employees:
            version = emp.current_version_id.sudo() if emp.current_version_id else False
            ctype_name = version.contract_type_id.name if version and version.contract_type_id else 'Sin tipo'
            if ctype_name not in by_contract_type:
                by_contract_type[ctype_name] = self.env['hr.employee']
            by_contract_type[ctype_name] |= emp

        for ctype_name, emps in sorted(by_contract_type.items(), key=lambda x: x[0]):
            if not emps:
                continue
            sheet_name = self._sanitize_sheet_name(ctype_name)
            ws = wb.create_sheet(title=sheet_name)
            self._write_employees_sheet(ws, emps, SEX_LABELS, MARITAL_LABELS,
                                       header_fill, header_font, header_alignment)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'Empleados_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': 'hr.employee',
            'res_id': self[:1].id if self else 0,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true&filename={filename}',
            'target': 'self',
        }

    def _sanitize_sheet_name(self, name):
        """Excel sheet names: max 31 chars, no : \\ / ? * [ ]"""
        if not name:
            return 'Sin tipo'
        invalid = ':\\/?*[]'
        for c in invalid:
            name = name.replace(c, '_')
        return (name[:31] if len(name) > 31 else name) or 'Sin tipo'

    def _write_registro_contratos_sheet(self, ws, employees, SEX_LABELS, MARITAL_LABELS,
                                        MONTHS_ES_FULL, orange_fill, blue_fill, white_font,
                                        black_font, center_align, row_alt_fill):
        """Escribe una hoja con el formato Registro de contratos."""
        def _fmt_date(d):
            return f'{d.day} de {MONTHS_ES_FULL[d.month - 1]} del {d.year}' if d else ''

        personal_headers = ['DNI', 'Nombre Completo', 'Sexo', 'Lugar de Residencia', 'Estado Civil',
                           'Profesión u oficio', 'Número de Teléfono', 'Correo Electrónico']
        labor_headers = ['Area', 'Nombre Completo del Jefe Inmediato', 'Salario', 'Cargo',
                        'Fecha de Inicio', 'Fecha de Terminación', 'Fecha de firma contrato', 'Tipo de Contrato',
                        'Días para vencer']
        all_headers = personal_headers + labor_headers

        ws.merge_cells('A1:Q1')
        ws.cell(row=1, column=1, value='INFORMACIÓN DE NUEVOS CONTRATOS')
        ws.cell(row=1, column=1).font = black_font
        ws.cell(row=1, column=1).alignment = center_align

        ws.cell(row=2, column=1, value='DATOS PERSONALES')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
        ws.cell(row=2, column=1).fill = orange_fill
        ws.cell(row=2, column=1).font = white_font
        ws.cell(row=2, column=1).alignment = center_align

        ws.cell(row=2, column=9, value='DATOS LABORALES')
        ws.merge_cells(start_row=2, start_column=9, end_row=2, end_column=17)
        ws.cell(row=2, column=9).fill = orange_fill
        ws.cell(row=2, column=9).font = white_font
        ws.cell(row=2, column=9).alignment = center_align

        for col, header in enumerate(all_headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = blue_fill
            cell.font = white_font
            cell.alignment = center_align

        for row_idx, emp in enumerate(employees, 4):
            version = emp.current_version_id.sudo() if emp.current_version_id else False
            row_fill = row_alt_fill if (row_idx - 4) % 2 == 1 else None
            addr_parts = [p for p in (emp.private_street, emp.private_street2, emp.private_city,
                                       emp.private_state_id.name if emp.private_state_id else None,
                                       emp.private_country_id.name if emp.private_country_id else None) if p]
            lugar_residencia = ', '.join(addr_parts) if addr_parts else (emp.private_city or emp.private_street or '')
            vals = [
                emp.identification_id or '',
                emp.name or '',
                SEX_LABELS.get(emp.sex, emp.sex or ''),
                lugar_residencia,
                MARITAL_LABELS.get(emp.marital, emp.marital or ''),
                emp.job_title or (emp.job_id.name if emp.job_id else ''),
                emp.work_phone or '',
                emp.work_email or 'N/A' if not emp.work_email else emp.work_email,
                emp.department_id.name if emp.department_id else '',
                emp.parent_id.name if emp.parent_id else '',
                float(version.wage if version else 0),
                emp.job_id.name if emp.job_id else '',
                _fmt_date(version.contract_date_start) if version else '',
                _fmt_date(version.contract_date_end) if version else '',
                _fmt_date(emp.current_date_version),
                version.contract_type_id.name if version and version.contract_type_id else '',
                emp.days_to_contract_expiration if emp.days_to_contract_expiration else '',
            ]
            for col, val in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                if row_fill:
                    cell.fill = row_fill

        from openpyxl.utils import get_column_letter
        for col in range(1, len(all_headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22
        _l10n_hn_xlsx_apply_table(ws, len(all_headers), header_row=3)

    def action_export_registro_contratos(self):
        """Export Registro de contratos Excel.
        Filename: Registro contratos a la fecha de DD-MMM-YYYY (fecha del día de generación).
        Hoja 1: Todos los empleados. Hojas adicionales: una por tipo de contrato (solo si hay empleados).
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))

        active_ids = self.env.context.get('active_ids', [])
        if active_ids:
            employees = self.browse(active_ids).exists()
        else:
            employees = self.search([('company_id', 'in', self.env.companies.ids)], order='name')

        if not employees:
            raise UserError(_('No hay empleados para exportar.'))

        SEX_LABELS = {'male': 'Masculino', 'female': 'Femenino', 'other': 'Otro'}
        MARITAL_LABELS = {
            'single': 'Soltero(a)', 'married': 'Casado(a)', 'cohabitant': 'Unión Libre',
            'widower': 'Viudo(a)', 'divorced': 'Divorciado(a)', 'other': 'Otro',
        }
        MONTHS_ES = ('Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic')
        MONTHS_ES_FULL = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
                          'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')

        today = datetime.now().date()
        date_str = f'{today.day:02d}-{MONTHS_ES[today.month - 1]}-{today.year}'
        filename = f'Registro contratos a la fecha de {date_str}.xlsx'

        orange_fill = PatternFill(start_color='E67E22', end_color='E67E22', fill_type='solid')
        blue_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        white_font = Font(bold=True, color='FFFFFF')
        black_font = Font(bold=True, color='000000', size=14)
        center_align = Alignment(horizontal='center', vertical='center')
        row_alt_fill = PatternFill(start_color='FFFACD', end_color='FFFACD', fill_type='solid')

        wb = Workbook()
        # Hoja 1: Todos los empleados
        ws_all = wb.active
        ws_all.title = self._sanitize_sheet_name('Todos')
        self._write_registro_contratos_sheet(
            ws_all, employees, SEX_LABELS, MARITAL_LABELS, MONTHS_ES_FULL,
            orange_fill, blue_fill, white_font, black_font, center_align, row_alt_fill
        )

        # Agrupar por tipo de contrato (solo tipos con empleados)
        by_contract_type = {}
        for emp in employees:
            version = emp.current_version_id.sudo() if emp.current_version_id else False
            ctype_name = version.contract_type_id.name if version and version.contract_type_id else 'Sin tipo'
            if ctype_name not in by_contract_type:
                by_contract_type[ctype_name] = self.env['hr.employee']
            by_contract_type[ctype_name] |= emp

        # Una hoja por tipo de contrato (excluir "Todos" que ya está)
        for ctype_name, emps in sorted(by_contract_type.items(), key=lambda x: x[0]):
            if not emps:
                continue
            sheet_name = self._sanitize_sheet_name(ctype_name)
            ws = wb.create_sheet(title=sheet_name)
            self._write_registro_contratos_sheet(
                ws, emps, SEX_LABELS, MARITAL_LABELS, MONTHS_ES_FULL,
                orange_fill, blue_fill, white_font, black_font, center_align, row_alt_fill
            )

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': 'hr.employee',
            'res_id': self[:1].id if self else 0,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true&filename={filename}',
            'target': 'self',
        }
    
    # Identification (Honduras specific)
    employee_code = fields.Char(string='Código de Empleado',
                                help='Código interno del empleado')
    biometric_code = fields.Char(
        string='Código Biométrico',
        help='Número con que el empleado está registrado en el dispositivo biométrico.',
    )
    l10n_hn_receives_commission = fields.Boolean(
        string='Recibe comisiones',
        default=False,
        help='Historial salarial mensual: si hay comisiones y está marcado, el total mostrado suma neto + comisiones.',
    )
    l10n_hn_add_commission_to_monthly_display = fields.Boolean(
        string='Sumar comisiones al total mostrado (historial mensual)',
        default=False,
        help='Si activo, siempre suma líneas de reglas marcadas como comisión al total mostrado. '
             'Peligro de doble conteo si el neto del recibo ya incluye comisiones.',
    )
    l10n_hn_monthly_salary_history_count = fields.Integer(
        compute='_compute_l10n_hn_monthly_salary_history_count',
        string='# Historial mensual',
    )

    identification_id = fields.Char(string='Identification Number', size=13,
                                    help='13-digit identification number (without dashes)')
    rtn = fields.Char(string='RTN', size=14,
                      help='Registro Tributario Nacional (14 dígitos)')
    
    # Clothing sizes
    shirt_size = fields.Selection([
        ('xs', 'XS'),
        ('s', 'S'),
        ('m', 'M'),
        ('l', 'L'),
        ('xl', 'XL'),
        ('xxl', 'XXL'),
        ('xxxl', 'XXXL'),
    ], string='Talla de Camisa')
    pants_size = fields.Char(string='Talla de Pantalón',
                              help='Ej: 32, 34, 36')
    shoe_size = fields.Char(string='Talla de Calzado',
                             help='Ej: 40, 42, 9, 10')
    
    # Misma definición que reclutamiento (varchar yes/not); Boolean aquí rompe -u de nómina.
    has_license = fields.Selection(
        selection='_selection_yes_no',
        string='¿Tiene Licencia de Conducir?',
        default='not',
    )
    
    # Family Members
    family_ids = fields.One2many(
        'hr.employee.family',
        'employee_id',
        string='Datos Familiares'
    )
    beneficiary_ids = fields.One2many(
        'hr.employee.beneficiary',
        'employee_id',
        string='Beneficiarios',
    )
    relatives_in_company = fields.Selection(
        selection='_selection_yes_no',
        string='¿Tiene parientes trabajando en esta empresa?',
        default='not',
    )
    in_company_relative_ids = fields.One2many(
        'hr.employee.in.company.relative',
        'employee_id',
        string='Parientes en la empresa',
    )
    has_allergies = fields.Boolean(string='¿Tiene alergias?')
    allergy_ids = fields.Many2many(
        'hr.allergy',
        'hr_employee_allergy_rel',
        'employee_id',
        'allergy_id',
        string='Indique alergias',
    )
    available_to_travel = fields.Boolean(
        string='¿Tiene disponibilidad de viajar?',
        help='Disponibilidad para viajar por trabajo u otros motivos laborales.',
    )
    parent_type = fields.Selection([
        ('father', 'Padre'),
        ('mother', 'Madre'),
    ], string='Padre/Madre',
       help='Indica si el empleado es padre o madre para efectos de reportes.')
    _CHILD_RELATIONS = ('Hijo', 'Hija', 'Hijastro', 'Hijastra')

    @api.depends('parent_type', 'sex', 'family_ids', 'family_ids.relation_id')
    def _compute_warn_parent(self):
        WarningType = self.env['hr.employee.warning.type']
        if not WarningType._is_warning_active('PARENT_TYPE'):
            for emp in self:
                emp.parent_type_warning = False
                emp.has_parent_type_warning = False
            return
        for emp in self:
            warnings = []
            if emp.parent_type:
                if emp.parent_type == 'mother' and emp.sex == 'male':
                    warnings.append(_('Empleado masculino marcado como madre.'))
                elif emp.parent_type == 'father' and emp.sex == 'female':
                    warnings.append(_('Empleada femenina marcada como padre.'))
                has_children = any(
                    f.relation_id and f.relation_id.name in self._CHILD_RELATIONS
                    for f in emp.family_ids
                )
                if not has_children:
                    warnings.append(_('Marcado como %s pero sin hijos registrados en datos familiares.')
                                    % dict(emp._fields['parent_type'].selection).get(emp.parent_type, ''))
            emp.parent_type_warning = ' | '.join(warnings) if warnings else False
            emp.has_parent_type_warning = bool(warnings)

    @api.onchange('relatives_in_company')
    def _onchange_relatives_in_company(self):
        if self.relatives_in_company != 'yes':
            self.in_company_relative_ids = [(5, 0, 0)]

    @api.onchange('has_allergies')
    def _onchange_has_allergies(self):
        if not self.has_allergies:
            self.allergy_ids = [(6, 0, [])]

    @api.onchange('disabled')
    def _onchange_disabled_disability_lines(self):
        if not self.disabled:
            self.disability_line_ids = [(5, 0, 0)]

    # Emergency contacts (multiple)
    emergency_contact_ids = fields.One2many(
        'hr.employee.emergency.contact',
        'employee_id',
        string='Contactos de Emergencia'
    )
    
    # Legacy emergency contact fields (deprecated - kept for backward compatibility)
    emergency_contact = fields.Char(string='Emergency Contact Name (Legacy)')
    emergency_relation = fields.Char(string='Emergency Contact Relation (Legacy)',
                                     help='Relationship with emergency contact')
    emergency_phone = fields.Char(string='Emergency Contact Phone (Legacy)')
    
    # Recurring concepts / ubicación laboral
    project_id = fields.Many2one(
        'hr.project',
        string='Proyecto',
        help='Proyecto asignado al empleado.',
        check_company=True,
    )
    assignment_id = fields.Many2one(
        'hr.assignment',
        string='Asignación',
        help='Asignación del empleado.',
        check_company=True,
    )

    # Contract information (using native Odoo 19 hr.version)
    # Note: current_version_id is already available from hr.employee (Odoo 19 native)
    current_periodicity = fields.Selection([
        ('weekly', 'Weekly'),
        ('biweekly', 'Biweekly'),
        ('monthly', 'Monthly'),
    ], string='Current Periodicity', 
       compute='_compute_current_periodicity', store=False,
       help='Payment periodicity from current contract version or structure default')
    
    @api.depends('current_version_id', 'current_version_id.periodicity')
    def _compute_current_periodicity(self):
        """Get current periodicity from active contract version.
        
        Note: The periodicity value in hr.version is automatically set from
        structure_type_id.default_schedule_pay when the version is created.
        See hr_version.py for the default value logic.
        """
        for employee in self:
            version = employee.current_version_id.sudo() if employee.current_version_id else False
            if version and version.periodicity:
                employee.current_periodicity = version.periodicity
            else:
                employee.current_periodicity = False
    
    # Legacy contract history (deprecated - kept for backward compatibility)
    # Use hr.version (version_ids) instead for contract history
    contract_history_ids = fields.One2many('hr.employee.contract.history', 'employee_id',
                                           string='Contract History (Legacy)')
    current_contract_id = fields.Many2one('hr.employee.contract.history', string='Current Contract (Legacy)',
                                           compute='_compute_current_contract', store=False)
    
    # Benefits
    thirteenth_month_ids = fields.One2many('l10n.hn.thirteenth.month', 'employee_id',
                                          string='13th Month Records')
    fourteenth_month_ids = fields.One2many('l10n.hn.fourteenth.month', 'employee_id',
                                           string='14th Month Records')

    # Planilla lines (HE, TE, BD) - from confirmed planillas
    payroll_line_he_ids = fields.One2many(
        'l10n.hn.payroll.line.he', 'employee_id',
        string='Horas Extra (Planillas)',
        help='Líneas de Horas Extra de planillas confirmadas'
    )
    payroll_line_te_ids = fields.One2many(
        'l10n.hn.payroll.line.te', 'employee_id',
        string='Turnos Extra (Planillas)',
        help='Líneas de Turnos Extra de planillas confirmadas'
    )
    payroll_line_bd_ids = fields.One2many(
        'l10n.hn.payroll.line.bd', 'employee_id',
        string='Beneficios y Deducciones (Planillas)',
        help='Líneas de Beneficios y Deducciones de planillas confirmadas'
    )
    
    # Overtime and shifts
    overtime_hours = fields.Float(string='Overtime Hours', digits=(10, 2))
    overtime_type = fields.Selection([
        ('daytime', 'Daytime'),
        ('nighttime', 'Nighttime'),
        ('rest_day', 'Rest Day'),
        ('holiday', 'Holiday'),
    ], string='Overtime Type')
    
    lactation_hours = fields.Float(string='Lactation Hours', digits=(10, 2),
                                   help='Paid lactation hours')
    
    # Health Status fields
    height = fields.Float(string='Altura (cm)', digits=(5, 2))
    weight = fields.Float(string='Peso (kg)', digits=(5, 2))
    blood_type_id = fields.Many2one(
        'l10n_hn_hr_payroll.blood.type',
        string='Tipo de sangre',
        ondelete='set null',
    )
    health_status = fields.Selection([
        ('good', 'Buena'),
        ('regular', 'Regular'),
        ('bad', 'Mala'),
    ], string='Estado de Salud')
    
    # Sports
    practice_sports = fields.Boolean(string='¿Practica Deportes?', default=False)
    sport_ids = fields.Many2many('hr.sport', 'hr_employee_sport_rel', 
                                  'employee_id', 'sport_id', string='Deportes')
    
    # Social club
    visit_social_club = fields.Boolean(string='¿Visita un Club Social o Deportivo?', default=False)
    
    # Hobbies
    hobby_ids = fields.Many2many('hr.hobby', 'hr_employee_hobby_rel',
                                  'employee_id', 'hobby_id', string='Pasatiempos Favoritos')
    
    # Diseases
    has_chronic_disease = fields.Boolean(string='¿Enfermedad Crónica?', default=False)
    chronic_disease_line_ids = fields.One2many(
        'hr.employee.chronic.disease.line',
        'employee_id',
        string='Enfermedades crónicas',
    )
    
    # Medications
    continuous_medication = fields.Boolean(string='¿Medicamentos de Uso Continuo?', default=False)
    medication_ids = fields.Many2many('hr.medication', 'hr_employee_medication_rel',
                                       'employee_id', 'medication_id', string='Medicamentos')
    
    # Surgeries
    surgery_ids = fields.Many2many('hr.surgery', 'hr_employee_surgery_rel',
                                    'employee_id', 'surgery_id', string='Cirugías')
    
    # IHSS configuration
    ihss_excluded = fields.Boolean(string='Exclude from IHSS', default=False,
                                    help='Exclude this employee from IHSS deductions')
    ihss_settings_id = fields.Many2one('l10n.hn.ihss.settings', string='IHSS Settings',
                                        compute='_compute_ihss_settings', store=False)
    
    # RAP configuration
    rap_excluded = fields.Boolean(string='Exclude from RAP', default=False,
                                  help='Exclude this employee from RAP deductions')
    rap_settings_id = fields.Many2one('l10n.hn.rap.settings', string='RAP Settings',
                                        compute='_compute_rap_settings', store=False)
    
    # ISR configuration
    isr_excluded = fields.Boolean(string='Exclude from ISR', default=False,
                                  help='Exclude this employee from ISR deductions')
    isr_history_ids = fields.One2many('l10n.hn.isr.history', 'employee_id', string='ISR History')
    isr_employee_deduction_ids = fields.One2many('l10n.hn.isr.employee_deduction', 'employee_id',
                                                  string='ISR Deductions')
    isr_ytd = fields.Float(string='ISR YTD', compute='_compute_isr_ytd', store=False,
                           help='ISR accumulated year to date')
    
    # Deduction amounts from current contract version (for display in employee form)
    ihss_amount_monthly = fields.Float(string='IHSS Mensual', digits=(12, 2),
                                       compute='_compute_deduction_amounts', store=False,
                                       help='Monto mensual de IHSS basado en la versión actual del contrato')
    ihss_amount_biweekly = fields.Float(string='IHSS Quincenal', digits=(12, 2),
                                        compute='_compute_deduction_amounts', store=False,
                                        help='Monto quincenal de IHSS')
    ihss_ivm_amount = fields.Float(string='IVM Mensual', digits=(12, 2),
                                   compute='_compute_deduction_amounts', store=False,
                                   help='Monto mensual de IVM (Invalidez, Vejez y Muerte)')
    ihss_em_amount = fields.Float(string='ENF-MAT Mensual', digits=(12, 2),
                                  compute='_compute_deduction_amounts', store=False,
                                  help='Monto mensual de ENF-MAT (Enfermedad y Maternidad)')
    
    rap_foviif_amount_monthly = fields.Float(string='RAP Mensual', digits=(12, 2),
                                             compute='_compute_deduction_amounts', store=False,
                                             help='Monto mensual de RAP basado en la versión actual del contrato')
    rap_foviif_amount_biweekly = fields.Float(string='RAP Quincenal', digits=(12, 2),
                                              compute='_compute_deduction_amounts', store=False,
                                              help='Monto quincenal de RAP')
    
    isr_amount_monthly = fields.Float(string='ISR Mensual Estimado', digits=(12, 2),
                                     compute='_compute_deduction_amounts', store=False,
                                     help='Monto mensual estimado de ISR basado en la versión actual del contrato')
    isr_amount_biweekly = fields.Float(string='ISR Quincenal Estimado', digits=(12, 2),
                                      compute='_compute_deduction_amounts', store=False,
                                      help='Monto quincenal estimado de ISR')
    
    ihss_manual = fields.Boolean(
        string='IHSS Manual', compute='_compute_manual_deduction_fields',
        inverse='_inverse_ihss_manual', store=False)
    amount_ihss_manual = fields.Float(
        string='Monto IHSS Manual', compute='_compute_manual_deduction_fields',
        inverse='_inverse_amount_ihss_manual', store=False)
    rap_manual = fields.Boolean(
        string='RAP Manual', compute='_compute_manual_deduction_fields',
        inverse='_inverse_rap_manual', store=False)
    amount_rap_manual = fields.Float(
        string='Monto RAP Manual', compute='_compute_manual_deduction_fields',
        inverse='_inverse_amount_rap_manual', store=False)
    amount_isr_calculated = fields.Float(
        string='ISR Calculado', compute='_compute_manual_deduction_fields',
        store=False, readonly=True)
    isr_manual = fields.Boolean(
        string='ISR Manual', compute='_compute_manual_deduction_fields',
        inverse='_inverse_isr_manual', store=False)
    amount_isr_manual = fields.Float(
        string='Monto ISR Manual', compute='_compute_manual_deduction_fields',
        inverse='_inverse_amount_isr_manual', store=False)

    @api.depends('current_version_id')
    def _compute_manual_deduction_fields(self):
        for emp in self:
            v = emp.current_version_id.sudo() if emp.current_version_id else False
            emp.ihss_manual = v.ihss_manual if v else False
            emp.amount_ihss_manual = v.amount_ihss_manual if v else 0.0
            emp.rap_manual = v.rap_manual if v else False
            emp.amount_rap_manual = v.amount_rap_manual if v else 0.0
            emp.amount_isr_calculated = v.amount_isr_calculated if v else 0.0
            emp.isr_manual = v.isr_manual if v else False
            emp.amount_isr_manual = v.amount_isr_manual if v else 0.0

    def _inverse_ihss_manual(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().ihss_manual = emp.ihss_manual

    def _inverse_amount_ihss_manual(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().amount_ihss_manual = emp.amount_ihss_manual

    def _inverse_rap_manual(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().rap_manual = emp.rap_manual

    def _inverse_amount_rap_manual(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().amount_rap_manual = emp.amount_rap_manual

    def _inverse_isr_manual(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().isr_manual = emp.isr_manual

    def _inverse_amount_isr_manual(self):
        for emp in self:
            if emp.current_version_id:
                emp.current_version_id.sudo().amount_isr_manual = emp.amount_isr_manual
    
    @api.depends('current_version_id', 'current_version_id.ihss_amount_monthly',
                 'current_version_id.ihss_amount_biweekly', 'current_version_id.ihss_ivm_amount',
                 'current_version_id.ihss_em_amount', 'current_version_id.rap_foviif_amount_monthly',
                 'current_version_id.rap_foviif_amount_biweekly', 'current_version_id.wage',
                 'current_version_id.periodicity', 'current_version_id.amount_isr_calculated',
                 'current_version_id.isr_manual', 'current_version_id.amount_isr_manual',
                 'isr_excluded', 'ihss_excluded', 'rap_excluded')
    def _compute_deduction_amounts(self):
        """Compute deduction amounts from current contract version"""
        # Pre-fetch recent payslips for all employees to avoid N+1
        all_payslips = self.env['hr.payslip'].search([
            ('employee_id', 'in', self.ids),
            ('state', 'in', ['done', 'paid']),
        ], order='date_to desc')
        payslip_map = {}
        for slip in all_payslips:
            payslip_map.setdefault(slip.employee_id.id, slip)
        # Pre-fetch ISR line amounts
        isr_amounts = {}
        payslip_ids = [s.id for s in payslip_map.values()]
        if payslip_ids:
            isr_lines = self.env['hr.payslip.line'].search([
                ('slip_id', 'in', payslip_ids),
                ('code', '=', 'ISR'),
            ])
            for line in isr_lines:
                isr_amounts[line.slip_id.id] = abs(line.total)

        for employee in self:
            if not employee.current_version_id:
                employee.ihss_amount_monthly = 0.0
                employee.ihss_amount_biweekly = 0.0
                employee.ihss_ivm_amount = 0.0
                employee.ihss_em_amount = 0.0
                employee.rap_foviif_amount_monthly = 0.0
                employee.rap_foviif_amount_biweekly = 0.0
                employee.isr_amount_monthly = 0.0
                employee.isr_amount_biweekly = 0.0
                continue
            
            version = employee.current_version_id.sudo()
            
            # IHSS amounts (from version)
            if not employee.ihss_excluded:
                employee.ihss_amount_monthly = version.ihss_amount_monthly or 0.0
                employee.ihss_amount_biweekly = version.ihss_amount_biweekly or 0.0
                employee.ihss_ivm_amount = version.ihss_ivm_amount or 0.0
                employee.ihss_em_amount = version.ihss_em_amount or 0.0
            else:
                employee.ihss_amount_monthly = 0.0
                employee.ihss_amount_biweekly = 0.0
                employee.ihss_ivm_amount = 0.0
                employee.ihss_em_amount = 0.0
            
            # RAP amounts (from version)
            if not employee.rap_excluded:
                employee.rap_foviif_amount_monthly = version.rap_foviif_amount_monthly or 0.0
                employee.rap_foviif_amount_biweekly = version.rap_foviif_amount_biweekly or 0.0
            else:
                employee.rap_foviif_amount_monthly = 0.0
                employee.rap_foviif_amount_biweekly = 0.0
            
            # ISR amounts: payslip history > manual > calculado > 0
            if not employee.isr_excluded:
                isr_amount = 0.0
                periodicity = version.periodicity or 'monthly'
                recent_payslip = payslip_map.get(employee.id)
                if recent_payslip:
                    isr_amount = isr_amounts.get(recent_payslip.id, 0.0)

                if not isr_amount and version.isr_manual and version.amount_isr_manual:
                    isr_amount = abs(version.amount_isr_manual)

                if not isr_amount and version.amount_isr_calculated:
                    isr_amount = abs(version.amount_isr_calculated)

                if isr_amount:
                    if periodicity in ('bi-weekly', 'biweekly'):
                        employee.isr_amount_biweekly = isr_amount
                        employee.isr_amount_monthly = isr_amount * 2
                    elif periodicity == 'weekly':
                        employee.isr_amount_biweekly = isr_amount * 2
                        employee.isr_amount_monthly = isr_amount * 4
                    else:
                        employee.isr_amount_monthly = isr_amount
                        employee.isr_amount_biweekly = isr_amount / 2
                else:
                    employee.isr_amount_monthly = 0.0
                    employee.isr_amount_biweekly = 0.0
            else:
                employee.isr_amount_monthly = 0.0
                employee.isr_amount_biweekly = 0.0
    
    @api.depends('company_id')
    def _compute_ihss_settings(self):
        """Get active IHSS settings for employee's company"""
        current_year = fields.Date.today().year
        all_settings = self.env['l10n.hn.ihss.settings'].search([
            ('company_id', 'in', self.mapped('company_id').ids),
            ('state', '=', 'validated'),
            ('year', '=', current_year),
        ], order='create_date desc')
        settings_map = {}
        for s in all_settings:
            settings_map.setdefault(s.company_id.id, s)
        for employee in self:
            employee.ihss_settings_id = settings_map.get(employee.company_id.id, False)
    
    @api.depends('company_id')
    def _compute_rap_settings(self):
        """Get active RAP settings for employee's company"""
        current_year = fields.Date.today().year
        all_settings = self.env['l10n.hn.rap.settings'].search([
            ('company_id', 'in', self.mapped('company_id').ids),
            ('state', '=', 'validated'),
            ('year', '=', current_year),
        ], order='create_date desc')
        settings_map = {}
        for s in all_settings:
            settings_map.setdefault(s.company_id.id, s)
        for employee in self:
            employee.rap_settings_id = settings_map.get(employee.company_id.id, False)
    
    def _compute_isr_ytd(self):
        """Calculate ISR YTD for current year"""
        current_year = fields.Date.today().year
        all_history = self.env['l10n.hn.isr.history'].search([
            ('employee_id', 'in', self.ids),
            ('year', '=', current_year),
        ], order='period_end desc')
        history_map = {}
        for h in all_history:
            history_map.setdefault(h.employee_id.id, h)
        for employee in self:
            hist = history_map.get(employee.id)
            employee.isr_ytd = hist.isr_ytd if hist else 0.0
    
    @api.depends('contract_history_ids', 'contract_history_ids.date_start', 'contract_history_ids.date_end')
    def _compute_current_contract(self):
        """Get current active contract"""
        today = fields.Date.today()
        for employee in self:
            current = employee.contract_history_ids.filtered(
                lambda c: c.active and c.date_start and c.date_start <= today
                          and (not c.date_end or c.date_end >= today)
            ).sorted('date_start', reverse=True)[:1]
            employee.current_contract_id = current or False
    
    @api.constrains('identification_id')
    def _check_identification(self):
        """Validate identification number format (13 digits)"""
        for employee in self:
            if employee.company_id.country_id.code and employee.company_id.country_id.code != 'HN':
                continue
            if employee.identification_id:
                if not employee.identification_id.isdigit() or len(employee.identification_id) != 13:
                    raise ValidationError(_('Identification number must be exactly 13 digits (without dashes).'))
    
    @api.constrains('rtn')
    def _check_rtn(self):
        """Validate RTN format (14 digits)"""
        for employee in self:
            if employee.rtn:
                # Remove any dashes or spaces
                rtn_clean = employee.rtn.replace('-', '').replace(' ', '')
                if not rtn_clean.isdigit() or len(rtn_clean) != 14:
                    raise ValidationError(_('El RTN debe tener exactamente 14 dígitos.'))
    
    def create_contract_version(self, wage, date_start, periodicity='monthly', job_id=False, department_id=False, date_end=False):
        """Create a new contract version using native Odoo 19 hr.version"""
        self.ensure_one()
        # Use native Odoo 19 version system
        version_vals = {
            'employee_id': self.id,
            'wage': wage,
            'contract_date_start': date_start,
            'contract_date_end': date_end,
            'periodicity': periodicity,  # Our custom field
            'date_version': date_start,  # Effective date for version
        }
        if job_id:
            version_vals['job_id'] = job_id
        if department_id:
            version_vals['department_id'] = department_id
        elif self.department_id:
            version_vals['department_id'] = self.department_id.id
        
        return self.env['hr.version'].create(version_vals)

    def _apply_low_net_pay_warning_hn(self):
        if not self._warning_type_active("LOW_NET_PAY_QUINCENA"):
            return
        today = fields.Date.context_today(self)
        ratio = self._gn_hr_warning_low_net_pay_ratio()
        for emp in self.filtered(lambda e: e.hn_company_country_code == "HN" and e.current_version_id):
            gross = emp._l10n_gn_gross_quincena_estimate()
            if gross <= 0:
                continue
            statutory = 0.0
            if not emp.ihss_excluded:
                statutory += emp.ihss_amount_biweekly or 0.0
            if not emp.rap_excluded:
                statutory += emp.rap_foviif_amount_biweekly or 0.0
            if not emp.isr_excluded:
                statutory += emp.isr_amount_biweekly or 0.0
            pnum = period_number_from_day(today.day)
            d_from, d_to = biweekly_date_range(today.year, today.month, pnum)
            recurring = emp.gn_estimate_recurring_deductions_for_payslip_period(d_from, d_to)
            total_ded = round(statutory + recurring, 2)
            net = round(gross - total_ded, 2)
            threshold_net = round(gross * ratio, 2)
            if net <= threshold_net + 0.01:
                emp.has_low_net_pay_quincena_warning = True
                emp.low_net_pay_quincena_warning_message = _(
                    "El líquido estimado para la quincena en curso (%(net).2f) es ≤ %(pct)g%% del bruto "
                    "(%(gross).2f), tras IHSS, RAP, ISR y deducciones recurrentes (estimación)."
                ) % {
                    "net": net,
                    "pct": round(ratio * 100.0, 2),
                    "gross": gross,
                }

    @api.depends(
        "hn_company_country_code",
        "ihss_amount_biweekly",
        "rap_foviif_amount_biweekly",
        "isr_amount_biweekly",
        "ihss_excluded",
        "rap_excluded",
        "isr_excluded",
        "recurring_concept_ids",
        "recurring_concept_ids.active",
        "recurring_concept_ids.is_template",
        "recurring_concept_ids.concept_type",
        "recurring_concept_ids.amount",
        "recurring_concept_ids.amount_type",
        "recurring_concept_ids.date_start",
        "recurring_concept_ids.date_end",
        "recurring_concept_ids.apply_biweek",
        "recurring_concept_ids.is_loan",
        "recurring_concept_ids.installments",
        "recurring_concept_ids.paid_installments",
        "recurring_concept_ids.loan_split_quota_across_biweeks",
        "recurring_concept_ids.loan_first_biweek_amount",
        "recurring_concept_ids.first_quota_differentiated",
        "recurring_concept_ids.first_quota_amount",
        "current_version_id",
        "current_version_id.wage",
        "current_version_id.periodicity",
        "current_version_id.schedule_pay",
    )
    def _compute_low_net_pay_quincena_warning(self):
        super()._compute_low_net_pay_quincena_warning()
        self._apply_low_net_pay_warning_hn()

    def action_unarchive(self):
        res = super().action_unarchive()
        self.write({'departure_specific_reason_id': False})
        return res

    def action_open_rehire_wizard(self):
        self.ensure_one()
        if self.active:
            raise UserError(_('La recontratación solo aplica a empleados archivados.'))
        return {
            'type': 'ir.actions.act_window',
            'name': _('Recontratar'),
            'res_model': 'l10n.hn.employee.rehire.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_employee_id': self.id},
        }

