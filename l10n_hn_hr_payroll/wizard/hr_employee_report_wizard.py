# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik

import base64
import io
import re
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime
from dateutil.relativedelta import relativedelta
import unicodedata

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


class HrEmployeeReportWizard(models.TransientModel):
    """Wizard para generar reportes de empleados a Excel"""
    _name = 'hr.employee.report.wizard'
    _description = 'Reporte de Empleados'

    _MONTH_ABBR_ES = (
        'ene', 'feb', 'mar', 'abr', 'may', 'jun',
        'jul', 'ago', 'sep', 'oct', 'nov', 'dic',
    )

    employee_ids = fields.Many2many(
        'hr.employee',
        string='Empleados',
        help='Dejar vacío para incluir todos los empleados activos'
    )
    company_id = fields.Many2one(
        'res.company',
        string='Empresa',
        required=True,
        default=lambda self: self.env.company
    )
    # Tipos de reporte seleccionables
    report_tallas = fields.Boolean(string='Tallas (camisa, pantalón, calzado)', default=False)
    report_emergency_contacts = fields.Boolean(string='Contactos de emergencia', default=False)
    report_licenses = fields.Boolean(string='Licencias de conducir', default=False)
    report_birthdays = fields.Boolean(string='Cumpleaños', default=False)
    report_family = fields.Boolean(string='Datos familiares', default=False)
    report_documents = fields.Boolean(string='Documentos de entrada', default=False)
    report_loans = fields.Boolean(string='Préstamos (detalle)', default=False)
    report_contracts = fields.Boolean(string='Registro de contratos', default=False)
    report_salary_increases = fields.Boolean(string='Historial de aumentos salariales', default=False)
    report_seniority = fields.Boolean(string='Antigüedad', default=False)
    report_headcount = fields.Boolean(string='Headcount (depto/asignación/proyecto)', default=False)
    report_references = fields.Boolean(string='Referencias', default=False)
    report_health_status = fields.Boolean(string='Estado de salud', default=False)
    report_beneficiaries = fields.Boolean(string='Beneficiarios', default=False)
    report_warnings = fields.Boolean(string='Empleados con alertas activas', default=False)
    report_trial_period = fields.Boolean(string='Empleados en período de prueba', default=False)
    report_hires_departures = fields.Boolean(string='Altas y Bajas', default=False)

    date_from = fields.Date(
        string='Desde',
        default=lambda self: fields.Date.today().replace(day=1),
        help='Fecha inicio del período para reportes que requieren rango de fechas (ej. Altas y Bajas)',
    )
    date_to = fields.Date(
        string='Hasta',
        default=lambda self: fields.Date.today(),
        help='Fecha fin del período para reportes que requieren rango de fechas',
    )
    birthday_range_from = fields.Date(
        string='Cumpleaños desde',
        help='Opcional: limitar el reporte a empleados cuyo cumpleaños (día y mes) cae en este intervalo de fechas del calendario.',
    )
    birthday_range_to = fields.Date(
        string='Cumpleaños hasta',
        help='Opcional: última fecha del intervalo (inclusive). Debe usarse junto con «Cumpleaños desde».',
    )

    @api.constrains('report_birthdays', 'birthday_range_from', 'birthday_range_to')
    def _check_birthday_range(self):
        for wiz in self:
            if not wiz.report_birthdays:
                continue
            a, b = wiz.birthday_range_from, wiz.birthday_range_to
            if (a or b) and not (a and b):
                raise ValidationError(_(
                    'Para filtrar por cumpleaños indique ambas fechas («Cumpleaños desde» y «Cumpleaños hasta») '
                    'o deje ambas vacías.'
                ))
            if a and b and a > b:
                raise ValidationError(_('«Cumpleaños hasta» debe ser mayor o igual que «Cumpleaños desde».'))

    def _birthday_anniversary_on_year(self, birthday, year):
        """Fecha de cumpleaños en un año civil (ajusta 29-feb en años no bisiestos)."""
        if not birthday:
            return None
        m, d = birthday.month, birthday.day
        last = monthrange(year, m)[1]
        dd = min(d, last)
        return date(year, m, dd)

    def _employee_birthday_occurs_in_calendar_range(self, birthday, d0, d1):
        """True si el aniversario (mes/día) coincide con alguna fecha en [d0, d1]."""
        if not birthday or not d0 or not d1 or d0 > d1:
            return False
        for y in range(d0.year - 1, d1.year + 2):
            occ = self._birthday_anniversary_on_year(birthday, y)
            if occ and d0 <= occ <= d1:
                return True
        return False

    def _filter_employees_birthday_range(self, employees):
        """Filtra por rango de fechas del calendario; sin rango definido devuelve el mismo recordset."""
        self.ensure_one()
        if not (self.birthday_range_from and self.birthday_range_to):
            return employees
        d0, d1 = self.birthday_range_from, self.birthday_range_to
        return employees.filtered(
            lambda e: e.birthday
            and self._employee_birthday_occurs_in_calendar_range(e.birthday, d0, d1),
        )

    def _get_employees(self):
        """Obtener empleados según filtro"""
        self.ensure_one()
        domain = [
            ('company_id', '=', self.company_id.id),
            ('active', '=', True),
        ]
        if self.employee_ids:
            domain.append(('id', 'in', self.employee_ids.ids))
        return self.env['hr.employee'].search(domain, order='name')

    def _hd_version_employee_domain(self):
        if self.employee_ids:
            return [('employee_id', 'in', self.employee_ids.ids)]
        return []

    def _hd_count_contract_starts(self, d_start, d_end):
        """Empleados distintos con inicio de contrato en el período (hr.version)."""
        dom = [
            ('employee_id.company_id', '=', self.company_id.id),
            ('contract_date_start', '!=', False),
            ('contract_date_start', '>=', d_start),
            ('contract_date_start', '<=', d_end),
        ]
        dom += self._hd_version_employee_domain()
        versions = self.env['hr.version'].sudo().search(dom)
        return len(set(versions.mapped('employee_id').ids))

    def _hd_employee_active_on_snapshot(self, emp, snapshot_date):
        """Contratado y no de baja al cierre del día snapshot_date (incluye ese día)."""
        emp = emp.sudo()
        versions = emp.version_ids.filtered(lambda v: v.date_version <= snapshot_date)
        if not versions:
            return False
        ver = versions.sorted('date_version', reverse=True)[:1]
        v = ver[0]
        if not v.contract_date_start or v.contract_date_start > snapshot_date:
            return False
        if v.contract_date_end and v.contract_date_end < snapshot_date:
            return False
        if emp.departure_date and emp.departure_date <= snapshot_date:
            return False
        return True

    def _hd_count_active_headcount_on_date(self, snapshot_date):
        """Población empleados activos al cierre del día indicado (p. ej. 31-dic para enero siguiente)."""
        dom = [('company_id', '=', self.company_id.id)]
        if self.employee_ids:
            dom.append(('id', 'in', self.employee_ids.ids))
        employees = self.env['hr.employee'].with_context(active_test=False).search(dom)
        return sum(1 for emp in employees if self._hd_employee_active_on_snapshot(emp, snapshot_date))

    def _hd_count_contract_ends_union(self, d_start, d_end):
        """Bajas en el mes: unión por fecha de baja del empleado y/o fin de contrato en versión."""
        emp_ids = set()
        Emp = self.env['hr.employee'].with_context(active_test=False)
        dom_e = [
            ('company_id', '=', self.company_id.id),
            ('departure_date', '!=', False),
            ('departure_date', '>=', d_start),
            ('departure_date', '<=', d_end),
        ]
        if self.employee_ids:
            dom_e.append(('id', 'in', self.employee_ids.ids))
        emp_ids.update(Emp.search(dom_e).ids)
        dom_v = [
            ('employee_id.company_id', '=', self.company_id.id),
            ('contract_date_end', '!=', False),
            ('contract_date_end', '>=', d_start),
            ('contract_date_end', '<=', d_end),
        ]
        dom_v += self._hd_version_employee_domain()
        emp_ids.update(
            self.env['hr.version'].sudo().search(dom_v).mapped('employee_id').ids
        )
        return len(emp_ids)

    def _hd_count_departures_by_reason_year(self, year, reason_id):
        """reason_id None = sin motivo de salida configurado."""
        y_start = date(year, 1, 1)
        y_end = date(year, 12, 31)
        dom = [
            ('company_id', '=', self.company_id.id),
            ('departure_date', '>=', y_start),
            ('departure_date', '<=', y_end),
        ]
        if reason_id is None:
            dom.append(('departure_reason_id', '=', False))
        else:
            dom.append(('departure_reason_id', '=', reason_id))
        if self.employee_ids:
            dom.append(('id', 'in', self.employee_ids.ids))
        return self.env['hr.employee'].with_context(active_test=False).search_count(dom)

    def _common_cols(self, emp):
        """Columnas comunes: Nombre, Departamento, Asignación (hr.assignment), DNI"""
        return [
            emp.name or '',
            emp.department_id.name if emp.department_id else '',
            emp.assignment_id.name if emp.assignment_id else '',
            emp.identification_id or '',
        ]

    _XLSX_MAX_STR = 32767
    _XLSX_ILLEGAL_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

    def _xlsx_clean_value(self, val):
        """Texto seguro para OOXML: controles, suplentes inválidos y longitud máxima de Excel."""
        if not isinstance(val, str):
            return val
        s = unicodedata.normalize('NFC', val)
        s = self._XLSX_ILLEGAL_CHARS.sub('', s)
        s = s.replace('\ufeff', '')
        out = []
        for ch in s:
            cat = unicodedata.category(ch)
            if cat == 'Cc' and ch not in '\t\n\r':
                continue
            o = ord(ch)
            if 0xD800 <= o <= 0xDFFF or o in (0xFFFE, 0xFFFF):
                continue
            out.append(ch)
        s = ''.join(out)
        if len(s) > self._XLSX_MAX_STR:
            s = s[: self._XLSX_MAX_STR]
        return s

    def _sanitize_workbook_cells(self, wb):
        """Sanea todas las celdas de texto antes de guardar (datos de prod. vs local)."""
        if not wb:
            return
        for ws in wb.worksheets:
            if isinstance(ws.title, str):
                t = self._xlsx_clean_value(ws.title)
                ws.title = (t[:31] if len(t) > 31 else t) or 'Hoja'
            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    val = cell.value
                    if isinstance(val, str):
                        cell.value = self._xlsx_clean_value(val)

    def _birthdays_sheet_headers(self):
        return [
            _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
            _('Fecha Cumpleaños'), _('Año'), _('Mes'), _('Día'), _('Edad'),
        ]

    def _birthday_month_abbr_es(self, bd):
        if not bd:
            return ''
        return self._MONTH_ABBR_ES[bd.month - 1]

    def _write_birthdays_sheet(self, ws, employees, header_fill, header_font, gray_border):
        """Rellena una hoja de cumpleaños (encabezado + filas + borde). No añade tabla aquí."""
        headers = self._birthdays_sheet_headers()
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=self._xlsx_clean_value(h))
            cell.fill = header_fill
            cell.font = header_font
            cell.border = gray_border
        today = fields.Date.context_today(self)
        row = 2
        for emp in employees:
            edad = ''
            if emp.birthday:
                diff = relativedelta(today, emp.birthday)
                edad = diff.years
            fecha_s = ''
            anio = mes_abr = ''
            dia = ''
            if emp.birthday:
                fecha_s = emp.birthday.strftime('%Y-%m-%d')
                anio = emp.birthday.year
                mes_abr = self._birthday_month_abbr_es(emp.birthday)
                dia = emp.birthday.day
            vals = [self._xlsx_clean_value(x) for x in self._common_cols(emp)] + [
                self._xlsx_clean_value(fecha_s),
                anio,
                self._xlsx_clean_value(mes_abr),
                dia,
                edad,
            ]
            for c, val in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=val)
                cell.border = gray_border
            row += 1
        return ws.max_row

    def _apply_excel_table(self, ws, num_cols, display_name, header_row=1):
        """Sin tablas Excel ( elemento ``<table>`` en OOXML ).

        openpyxl genera definiciones de tabla que **Microsoft Excel en Windows** suele marcar como
        dañadas y abre con «Recuperar libro», aunque LibreOffice / otros no fallen. Los estilos de
        celda (bordes, cabecera, rayas) se siguen aplicando sin usar `Table`.
        """
        return

    def _seniority_fecha_ingreso(self, emp):
        """Primera fecha de inicio de contrato (hr.version) del empleado."""
        versions = emp.version_ids.filtered(lambda v: v.contract_date_start)
        if not versions:
            return None
        return min(v.contract_date_start for v in versions)

    def _seniority_anniversary_month_mmm(self, fecha_ingreso):
        """Mes del aniversario laboral (mes de fecha de ingreso) en formato abreviado tipo mmm."""
        if not fecha_ingreso:
            return ''
        return self._MONTH_ABBR_ES[fecha_ingreso.month - 1]

    def _seniority_human_label(self, fecha_ingreso, ref_date):
        """Antigüedad legible, p. ej. «1 año 2 meses + 10 días»."""
        if not fecha_ingreso:
            return ''
        diff = relativedelta(ref_date, fecha_ingreso)
        y, m, d = diff.years, diff.months, diff.days
        parts = []
        if y:
            parts.append(_('%d año') % y if y == 1 else _('%d años') % y)
        if m:
            parts.append(_('%d mes') % m if m == 1 else _('%d meses') % m)
        middle = ' '.join(parts)
        if d:
            day_s = (_('%d día') % d if d == 1 else _('%d días') % d)
            if middle:
                return '%s + %s' % (middle, day_s)
            return day_s
        return middle if middle else _('0 días')

    def _seniority_sheet_headers(self):
        return [
            _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
            _('Fecha Ingreso'), _('Mes aniversario'), _('Años'), _('Meses'),
            _('Antigüedad'), _('Cargo Actual'),
        ]

    def _write_seniority_sheet(self, ws, employees, header_fill, header_font, gray_border):
        """Antigüedad: años y meses completos (relativedelta); texto detallado y mes aniversario (mmm)."""
        headers = self._seniority_sheet_headers()
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=self._xlsx_clean_value(h))
            cell.fill = header_fill
            cell.font = header_font
            cell.border = gray_border
        today = fields.Date.context_today(self)
        row = 2
        for emp in employees:
            fecha_ingreso = self._seniority_fecha_ingreso(emp)
            mes_mmm = self._xlsx_clean_value(self._seniority_anniversary_month_mmm(fecha_ingreso))
            anti_txt = self._xlsx_clean_value(self._seniority_human_label(fecha_ingreso, today))
            anos = ''
            meses = ''
            if fecha_ingreso:
                diff = relativedelta(today, fecha_ingreso)
                anos = diff.years
                meses = diff.months
            vals = [self._xlsx_clean_value(x) for x in self._common_cols(emp)] + [
                self._xlsx_clean_value(
                    fecha_ingreso.strftime('%Y-%m-%d') if fecha_ingreso else ''
                ),
                mes_mmm,
                anos,
                self._xlsx_clean_value(meses),
                anti_txt,
                self._xlsx_clean_value(emp.job_id.name if emp.job_id else ''),
            ]
            for c, val in enumerate(vals, 1):
                cell = ws.cell(row=row, column=c, value=val)
                cell.border = gray_border
            row += 1
        return ws.max_row

    def _get_report_types(self):
        """Lista de (clave, nombre_hoja, activo)"""
        return [
            ('tallas', _('Tallas'), self.report_tallas),
            ('emergency', _('Contactos Emergencia'), self.report_emergency_contacts),
            ('licenses', _('Licencias Conducir'), self.report_licenses),
            ('birthdays', _('Cumpleaños'), self.report_birthdays),
            ('family', _('Datos Familiares'), self.report_family),
            ('beneficiaries', _('Beneficiarios'), self.report_beneficiaries),
            ('documents', _('Documentos Entrada'), self.report_documents),
            ('loans', _('Préstamos'), self.report_loans),
            ('contracts', _('Todos'), self.report_contracts),
            ('salary_increases', _('Historial Aumentos'), self.report_salary_increases),
            ('seniority', _('Antigüedad'), self.report_seniority),
            ('headcount', _('Headcount'), self.report_headcount),
            ('references', _('Referencias'), self.report_references),
            ('health_status', _('Estado de Salud'), self.report_health_status),
            ('warnings', _('Alertas Activas'), self.report_warnings),
            ('trial_period', _('Período de Prueba'), self.report_trial_period),
            ('hires_departures', _('Altas y Bajas'), self.report_hires_departures),
        ]

    def action_export_excel(self):
        """Generar reporte Excel"""
        self.ensure_one()
        if not Workbook:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))

        report_types = [(k, name, active) for k, name, active in self._get_report_types() if active]
        if not report_types:
            raise UserError(_('Debe seleccionar al menos un tipo de reporte.'))

        employees = self._get_employees()
        if not employees:
            raise UserError(_('No hay empleados que coincidan con el filtro.'))

        def _sanitize_sheet_name(name):
            if not name:
                return _('Hoja')
            name = str(name)
            for c in ':\\/?*[]':
                name = name.replace(c, '_')
            return (name[:31] if len(name) > 31 else name) or _('Hoja')

        gray_border = Border(
            left=Side(style='thin', color='A5A5A5'),
            right=Side(style='thin', color='A5A5A5'),
            top=Side(style='thin', color='A5A5A5'),
            bottom=Side(style='thin', color='A5A5A5'),
        )
        header_fill = PatternFill(start_color='A52A2A', end_color='A52A2A', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')

        wb = Workbook()
        wb.remove(wb.active)

        SHIRT_LABELS = {
            'xs': 'XS', 's': 'S', 'm': 'M', 'l': 'L',
            'xl': 'XL', 'xxl': 'XXL', 'xxxl': 'XXXL',
        }

        for report_key, sheet_name, _active in report_types:
            ws = wb.create_sheet(_sanitize_sheet_name(sheet_name), len(wb.worksheets))

            if report_key == 'tallas':
                headers = [_('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                          _('Talla Camisa'), _('Talla Pantalón'), _('Talla Calzado')]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                for row, emp in enumerate(employees, 2):
                    vals = self._common_cols(emp) + [
                        SHIRT_LABELS.get(emp.shirt_size, emp.shirt_size or ''),
                        emp.pants_size or '',
                        emp.shoe_size or '',
                    ]
                    for c, val in enumerate(vals, 1):
                        cell = ws.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                self._apply_excel_table(ws, 8, '')

            elif report_key == 'emergency':
                headers = [_('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                          _('Contacto'), _('Parentesco'), _('Teléfono'), _('Teléfono Sec.')]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for emp in employees:
                    contacts = emp.emergency_contact_ids
                    if not contacts:
                        vals = self._common_cols(emp) + ['', '', '', '']
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for contact in contacts:
                            vals = self._common_cols(emp) + [
                                contact.name or '',
                                contact.relation_id.name if contact.relation_id else '',
                                contact.phone or '',
                                contact.phone_secondary or '',
                            ]
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                self._apply_excel_table(ws, 8, '')

            elif report_key == 'licenses':
                headers = [_('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                          _('Tipo Licencia'), _('Número'), _('Fecha Vencimiento'),
                          _('Días para Vencer'), _('Estado')]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                today = fields.Date.today()
                row = 2
                for emp in employees:
                    licenses = emp.license_ids
                    if not licenses:
                        vals = self._common_cols(emp) + ['', '', '', '', '']
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for lic in licenses:
                            estado = _('Vencida') if lic.is_expired else _('Vigente')
                            dias_vencer = ''
                            if lic.expiration_date:
                                dias_vencer = (lic.expiration_date - today).days
                            vals = self._common_cols(emp) + [
                                lic.license_type_id.name if lic.license_type_id else '',
                                lic.license_number or '',
                                lic.expiration_date.strftime('%Y-%m-%d') if lic.expiration_date else '',
                                dias_vencer,
                                estado,
                            ]
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                self._apply_excel_table(ws, 9, '')

            elif report_key == 'birthdays':
                today = fields.Date.context_today(self)
                BDAY_NCOLS = 9
                if self.birthday_range_from and self.birthday_range_to:
                    bday_main = self._filter_employees_birthday_range(employees)
                    if not bday_main:
                        raise UserError(_(
                            'No hay empleados con cumpleaños en el rango seleccionado.',
                        ))
                else:
                    bday_main = employees
                self._write_birthdays_sheet(ws, bday_main, header_fill, header_font, gray_border)
                if ws.max_row > 1:
                    self._apply_excel_table(ws, BDAY_NCOLS, '')
                for c in range(1, BDAY_NCOLS + 1):
                    ws.column_dimensions[get_column_letter(c)].width = 16
                if not (self.birthday_range_from and self.birthday_range_to):
                    month_emps = employees.filtered(
                        lambda e: e.birthday and e.birthday.month == today.month,
                    ).sorted(key=lambda e: (e.birthday.day, e.name or ''))
                    ws_m = wb.create_sheet(
                        _sanitize_sheet_name(_('Cumpleaños del mes')),
                        len(wb.worksheets),
                    )
                    self._write_birthdays_sheet(ws_m, month_emps, header_fill, header_font, gray_border)
                    if ws_m.max_row > 1:
                        self._apply_excel_table(
                            ws_m, BDAY_NCOLS, '',
                        )
                    for c in range(1, BDAY_NCOLS + 1):
                        ws_m.column_dimensions[get_column_letter(c)].width = 16

            elif report_key == 'family':
                headers = [_('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                          _('Familiar'), _('Parentesco'), _('Fecha Nac.'), _('Edad'), _('Domicilio'), _('Ocupación')]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for emp in employees:
                    family = emp.family_ids
                    if not family:
                        vals = self._common_cols(emp) + ['', '', '', '', '', '']
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for fam in family:
                            vals = self._common_cols(emp) + [
                                fam.name or '',
                                fam.relation_id.name if fam.relation_id else '',
                                fam.birthdate.strftime('%Y-%m-%d') if fam.birthdate else '',
                                fam.age if fam.birthdate else '',
                                fam.address or '',
                                fam.occupation or '',
                            ]
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                self._apply_excel_table(ws, 10, '')

            elif report_key == 'beneficiaries':
                BEN_NCOLS = 7
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Parentesco'), _('Nombre beneficiario'), _('Porcentaje'),
                ]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for emp in employees:
                    lines = emp.beneficiary_ids
                    if not lines:
                        vals = self._common_cols(emp) + ['', '', '']
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for ben in lines:
                            vals = self._common_cols(emp) + [
                                ben.relation_id.name if ben.relation_id else '',
                                ben.name or '',
                                ben.percentage if ben.percentage is not None else '',
                            ]
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, BEN_NCOLS, '')

            elif report_key == 'documents':
                headers = [_('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                          _('Tipo Documento'), _('¿Tiene archivo?')]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for emp in employees:
                    docs = emp.document_ids
                    if not docs:
                        vals = self._common_cols(emp) + ['', '']
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for doc in docs:
                            tiene = _('Sí') if doc.document else _('No')
                            vals = self._common_cols(emp) + [
                                doc.document_type_id.name if doc.document_type_id else '',
                                tiene,
                            ]
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                self._apply_excel_table(ws, 6, '')

            elif report_key == 'loans':
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Nombre del Concepto'), _('Fecha Inicio'), _('Fecha Fin'),
                    _('Monto Cuota'), _('Cantidad Cuotas'), _('Monto Total'),
                    _('Monto Pagado'), _('Monto Adeudado'),
                ]
                currency_cols_loans = (8, 10, 11, 12)  # Monto Cuota, Total, Pagado, Adeudado
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                currency = self.company_id.currency_id or self.env.company.currency_id
                symbol = (currency.symbol or '').replace('"', '""')
                currency_fmt = '"%s"#,##0.00' % symbol if symbol else '#,##0.00'
                row = 2
                for emp in employees:
                    loans = emp.recurring_concept_ids.filtered(lambda c: c.is_loan and c.active)
                    if not loans:
                        continue
                    for loan in loans:
                        monto_pagado = loan.amount * loan.paid_installments
                        vals = self._common_cols(emp) + [
                            loan.name or '',
                            loan.date_start.strftime('%Y-%m-%d') if loan.date_start else '',
                            loan.date_end.strftime('%Y-%m-%d') if loan.date_end else '',
                            loan.amount,
                            loan.installments,
                            loan.total_amount,
                            monto_pagado,
                            loan.remaining_amount,
                        ]
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                            if c in currency_cols_loans and isinstance(val, (int, float)):
                                cell.number_format = currency_fmt
                        row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, 12, '')

            elif report_key == 'contracts':
                SEX_LABELS = {'male': _('Masculino'), 'female': _('Femenino'), 'other': _('Otro')}
                MARITAL_LABELS = {
                    'single': _('Soltero(a)'), 'married': _('Casado(a)'), 'cohabitant': _('Unión Libre'),
                    'widower': _('Viudo(a)'), 'divorced': _('Divorciado(a)'), 'other': _('Otro'),
                }
                MONTHS_ES_FULL = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
                                  'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')
                orange_fill = PatternFill(start_color='E67E22', end_color='E67E22', fill_type='solid')
                blue_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
                white_font = Font(bold=True, color='FFFFFF')
                black_font = Font(bold=True, color='000000', size=14)
                center_align = Alignment(horizontal='center', vertical='center')
                row_alt_fill = PatternFill(start_color='FFFACD', end_color='FFFACD', fill_type='solid')
                employees._write_registro_contratos_sheet(
                    ws, employees, SEX_LABELS, MARITAL_LABELS, MONTHS_ES_FULL,
                    orange_fill, blue_fill, white_font, black_font, center_align, row_alt_fill
                )
                if ws.max_row >= 3:
                    self._apply_excel_table(
                        ws, 17, '', header_row=3,
                    )
                by_contract_type = {}
                for emp in employees:
                    version = emp.current_version_id.sudo() if emp.current_version_id else False
                    ctype_name = version.contract_type_id.name if version and version.contract_type_id else _('Sin tipo')
                    if ctype_name not in by_contract_type:
                        by_contract_type[ctype_name] = self.env['hr.employee']
                    by_contract_type[ctype_name] |= emp
                for ctype_name, emps in sorted(by_contract_type.items(), key=lambda x: x[0]):
                    if not emps:
                        continue
                    sheet_name_ct = _sanitize_sheet_name(ctype_name)
                    if sheet_name_ct != _sanitize_sheet_name(sheet_name):
                        ws_ct = wb.create_sheet(sheet_name_ct, len(wb.worksheets))
                        emps._write_registro_contratos_sheet(
                            ws_ct, emps, SEX_LABELS, MARITAL_LABELS, MONTHS_ES_FULL,
                            orange_fill, blue_fill, white_font, black_font, center_align, row_alt_fill
                        )
                        if ws_ct.max_row >= 3:
                            self._apply_excel_table(
                                ws_ct, 17, '', header_row=3,
                            )
                        for c in range(1, ws_ct.max_column + 1):
                            ws_ct.column_dimensions[get_column_letter(c)].width = 22

            elif report_key == 'salary_increases':
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Fecha'), _('Salario Anterior'), _('Salario Nuevo'),
                    _('Monto Aumento'), _('% Aumento'), _('Puesto Anterior'), _('Puesto Nuevo'),
                ]
                currency_cols_inc = (6, 7, 8)
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                currency = self.company_id.currency_id or self.env.company.currency_id
                symbol = (currency.symbol or '').replace('"', '""')
                currency_fmt = '"%s"#,##0.00' % symbol if symbol else '#,##0.00'
                emp_ids = employees.ids
                historical = self.env['hr.historical.salary.increase'].search([
                    ('employee_id', 'in', emp_ids),
                    ('employee_id.company_id', '=', self.company_id.id),
                ], order='date desc, id desc')
                row = 2
                for rec in historical:
                    emp = rec.employee_id
                    vals = self._common_cols(emp) + [
                        rec.date.strftime('%Y-%m-%d') if rec.date else '',
                        rec.previous_wage,
                        rec.new_wage,
                        rec.increase_amount,
                        rec.increase_percentage,
                        rec.previous_job_id.name if rec.previous_job_id else '',
                        rec.new_job_id.name if rec.new_job_id else '',
                    ]
                    for c, val in enumerate(vals, 1):
                        cell = ws.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                        if c in currency_cols_inc and isinstance(val, (int, float)):
                            cell.number_format = currency_fmt
                    row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, 11, '')

            elif report_key == 'seniority':
                today = fields.Date.context_today(self)
                SEN_NCOLS = 10
                self._write_seniority_sheet(ws, employees, header_fill, header_font, gray_border)
                if ws.max_row > 1:
                    self._apply_excel_table(ws, SEN_NCOLS, '')
                for c in range(1, SEN_NCOLS + 1):
                    letter = get_column_letter(c)
                    ws.column_dimensions[letter].width = 26 if letter in ('I',) else 16
                month_ids = []
                for e in employees:
                    fi = self._seniority_fecha_ingreso(e)
                    if fi and fi.month == today.month:
                        month_ids.append(e.id)
                month_anniv = self.env['hr.employee'].browse(month_ids).sorted(
                    key=lambda em: (
                        self._seniority_fecha_ingreso(em).day,
                        em.name or '',
                    )
                )
                ws_m = wb.create_sheet(
                    _sanitize_sheet_name(_('Aniversarios del mes')),
                    len(wb.worksheets),
                )
                self._write_seniority_sheet(ws_m, month_anniv, header_fill, header_font, gray_border)
                if ws_m.max_row > 1:
                    self._apply_excel_table(ws_m, SEN_NCOLS, '')
                for c in range(1, SEN_NCOLS + 1):
                    ws_m.column_dimensions[get_column_letter(c)].width = (
                        26 if get_column_letter(c) == 'I' else 16
                    )

            elif report_key == 'headcount':
                headers_dept = [_('Departamento'), _('Cantidad')]
                headers_asig = [_('Asignación'), _('Cantidad')]
                headers_proj = [_('Proyecto'), _('Cantidad')]
                by_dept = {}
                by_asig = {}
                by_proj = {}
                for emp in employees:
                    dept = emp.department_id.name if emp.department_id else _('Sin departamento')
                    by_dept[dept] = by_dept.get(dept, 0) + 1
                    asig = emp.assignment_id.name if emp.assignment_id else _('Sin asignación')
                    by_asig[asig] = by_asig.get(asig, 0) + 1
                    proj = emp.project_id.name if emp.project_id else _('Sin proyecto')
                    by_proj[proj] = by_proj.get(proj, 0) + 1
                for col, h in enumerate(headers_dept, 1):
                    cell = ws.cell(row=1, column=col, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for dept_name in sorted(by_dept.keys()):
                    ws.cell(row=row, column=1, value=dept_name).border = gray_border
                    ws.cell(row=row, column=2, value=by_dept[dept_name]).border = gray_border
                    row += 1
                ws_asig = wb.create_sheet(_sanitize_sheet_name(_('Headcount Asignación')), len(wb.worksheets))
                for col, h in enumerate(headers_asig, 1):
                    cell = ws_asig.cell(row=1, column=col, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for asig_name in sorted(by_asig.keys()):
                    ws_asig.cell(row=row, column=1, value=asig_name).border = gray_border
                    ws_asig.cell(row=row, column=2, value=by_asig[asig_name]).border = gray_border
                    row += 1
                ws_proj = wb.create_sheet(_sanitize_sheet_name(_('Headcount Proyecto')), len(wb.worksheets))
                for col, h in enumerate(headers_proj, 1):
                    cell = ws_proj.cell(row=1, column=col, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for proj_name in sorted(by_proj.keys()):
                    ws_proj.cell(row=row, column=1, value=proj_name).border = gray_border
                    ws_proj.cell(row=row, column=2, value=by_proj[proj_name]).border = gray_border
                    row += 1
                if ws.max_row >= 1:
                    self._apply_excel_table(ws, 2, '')
                if ws_asig.max_row >= 1:
                    self._apply_excel_table(ws_asig, 2, '')
                if ws_proj.max_row >= 1:
                    self._apply_excel_table(ws_proj, 2, '')

            elif report_key == 'references':
                if 'reference.personal' not in self.env:
                    raise UserError(_(
                        'El modelo de referencias personales no está disponible. '
                        'Instale el módulo Honduras Recruitment (Experiencia y Referencias).'
                    ))
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Referencia (nombre)'), _('Dirección'), _('Teléfono'),
                    _('Ocupación'), _('Tiempo'),
                ]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for emp in employees:
                    refs = emp.hn_reference_ids if 'hn_reference_ids' in emp._fields else self.env['reference.personal'].browse()
                    if not refs:
                        vals = self._common_cols(emp) + ['', '', '', '', '']
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for ref in refs:
                            vals = self._common_cols(emp) + [
                                ref.name or '',
                                ref.address or '',
                                ref.phone or '',
                                ref.occupation or '',
                                ref.time or '',
                            ]
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, 9, '')

            elif report_key == 'health_status':
                HEALTH_LABELS = {
                    'good': _('Buena'), 'regular': _('Regular'), 'bad': _('Mala'),
                }
                HEALTH_NCOLS = 20
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Altura (cm)'), _('Peso (kg)'), _('Tipo de Sangre'),
                    _('Estado de Salud'), _('Practica Deportes'), _('Deportes'),
                    _('Visita Club Social'), _('Pasatiempos'),
                    _('¿Tiene enfermedad crónica?'),
                    _('Enfermedad (crónica)'), _('Fecha'), _('En tratamiento'),
                    _('Lugar de tratamiento'),
                    _('Medicamentos Continuos'), _('Medicamentos'), _('Cirugías'),
                ]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row = 2
                for emp in employees:
                    lines = emp.chronic_disease_line_ids
                    base = self._common_cols(emp) + [
                        emp.height or '',
                        emp.weight or '',
                        emp.blood_type_id.name if emp.blood_type_id else '',
                        HEALTH_LABELS.get(emp.health_status, emp.health_status or ''),
                        _('Sí') if emp.practice_sports else _('No'),
                        ', '.join(emp.sport_ids.mapped('name')) if emp.sport_ids else '',
                        _('Sí') if emp.visit_social_club else _('No'),
                        ', '.join(emp.hobby_ids.mapped('name')) if emp.hobby_ids else '',
                        _('Sí') if emp.has_chronic_disease else _('No'),
                    ]
                    med_tail = [
                        _('Sí') if emp.continuous_medication else _('No'),
                        ', '.join(emp.medication_ids.mapped('name')) if emp.medication_ids else '',
                        ', '.join(emp.surgery_ids.mapped('name')) if emp.surgery_ids else '',
                    ]
                    if not lines:
                        vals = base + ['', '', '', ''] + med_tail
                        for c, val in enumerate(vals, 1):
                            cell = ws.cell(row=row, column=c, value=val)
                            cell.border = gray_border
                        row += 1
                    else:
                        for line in lines:
                            vals = base + [
                                line.disease_id.name if line.disease_id else '',
                                line.disease_date.strftime('%Y-%m-%d') if line.disease_date else '',
                                _('Sí') if line.under_treatment else _('No'),
                                line.treatment_place or '',
                            ] + med_tail
                            for c, val in enumerate(vals, 1):
                                cell = ws.cell(row=row, column=c, value=val)
                                cell.border = gray_border
                            row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, HEALTH_NCOLS, '')

            elif report_key == 'warnings':
                warning_types = self.env['hr.employee.warning.type'].search([
                    ('active', '=', True),
                ], order='sequence, id')
                warning_defs = []
                for wt in warning_types:
                    if wt.employee_field:
                        warning_defs.append((wt.employee_field, wt.name))
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                ] + [name for _, name in warning_defs]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                alert_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
                alert_font = Font(bold=True, color='CC0000')
                center = Alignment(horizontal='center', vertical='center')
                row = 2
                for emp in employees:
                    warn_vals = []
                    has_any = False
                    for field_name, _name in warning_defs:
                        if getattr(emp, field_name, False):
                            has_any = True
                            warn_vals.append('X')
                        else:
                            warn_vals.append('')
                    if not has_any:
                        continue
                    vals = self._common_cols(emp) + warn_vals
                    for c, val in enumerate(vals, 1):
                        cell = ws.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                        if c > 4 and val:
                            cell.fill = alert_fill
                            cell.font = alert_font
                            cell.alignment = center
                    row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, len(headers), '')

            elif report_key == 'trial_period':
                trial_days = self.company_id._get_localization_trial_period_days()
                today = fields.Date.context_today(self)
                headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Fecha Ingreso'), _('Días Transcurridos'), _('Días Restantes'),
                    _('Fecha Fin Prueba'), _('Puesto'),
                ]
                for c, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                trial_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
                urgent_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
                row = 2
                for emp in employees:
                    start = emp._get_first_contract_date(no_gap=False)
                    if not start:
                        continue
                    elapsed = (today - start).days
                    if elapsed < 0 or elapsed > trial_days:
                        continue
                    remaining = trial_days - elapsed
                    end_date = start + relativedelta(days=trial_days)
                    vals = self._common_cols(emp) + [
                        start.strftime('%d/%m/%Y'),
                        elapsed,
                        remaining,
                        end_date.strftime('%d/%m/%Y'),
                        emp.job_id.name if emp.job_id else '',
                    ]
                    fill = urgent_fill if remaining <= 7 else trial_fill if remaining <= 15 else None
                    for c, val in enumerate(vals, 1):
                        cell = ws.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                        if fill:
                            cell.fill = fill
                    row += 1
                if ws.max_row > 1:
                    self._apply_excel_table(ws, 8, '')

            elif report_key == 'hires_departures':
                date_from = self.date_from
                date_to = self.date_to
                if not date_from or not date_to:
                    continue

                ws.title = _sanitize_sheet_name(_('Altas'))
                hire_headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Puesto'), _('Fecha Ingreso'), _('Salario'),
                ]
                for c, h in enumerate(hire_headers, 1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                hire_fill = PatternFill(start_color='D9EAD3', end_color='D9EAD3', fill_type='solid')
                row = 2
                for emp in employees:
                    start = emp._get_first_contract_date(no_gap=False)
                    if not start or start < date_from or start > date_to:
                        continue
                    version = emp.current_version_id.sudo() if emp.current_version_id else False
                    vals = self._common_cols(emp) + [
                        emp.job_id.name if emp.job_id else '',
                        start.strftime('%d/%m/%Y'),
                        version.wage if version else 0.0,
                    ]
                    for c, val in enumerate(vals, 1):
                        cell = ws.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                        if c <= 4:
                            cell.fill = hire_fill
                    row += 1

                ws_dep = wb.create_sheet(title=_sanitize_sheet_name(_('Bajas')))
                dep_headers = [
                    _('Nombre'), _('Departamento'), _('Asignación'), _('DNI'),
                    _('Puesto'), _('Fecha Ingreso'), _('Fecha Baja'),
                    _('Motivo de Baja'),
                ]
                for c, h in enumerate(dep_headers, 1):
                    cell = ws_dep.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                dep_fill = PatternFill(start_color='F4CCCC', end_color='F4CCCC', fill_type='solid')
                departed = self.env['hr.employee'].with_context(active_test=False).search([
                    ('company_id', '=', self.company_id.id),
                    ('active', '=', False),
                    ('departure_date', '>=', date_from),
                    ('departure_date', '<=', date_to),
                ], order='departure_date')
                if self.employee_ids:
                    departed = departed.filtered(lambda e: e.id in self.employee_ids.ids)
                row = 2
                for emp in departed:
                    start = emp._get_first_contract_date(no_gap=False)
                    version = (emp.current_version_id.sudo() if emp.current_version_id
                               else (emp.version_ids.sudo()[:1] if emp.version_ids else False))
                    vals = [
                        emp.name or '',
                        emp.department_id.name if emp.department_id else '',
                        emp.assignment_id.name if emp.assignment_id else '',
                        emp.identification_id or '',
                        emp.job_id.name if emp.job_id else '',
                        start.strftime('%d/%m/%Y') if start else '',
                        emp.departure_date.strftime('%d/%m/%Y') if emp.departure_date else '',
                        version.departure_reason_id.name if version and version.departure_reason_id else '',
                    ]
                    for c, val in enumerate(vals, 1):
                        cell = ws_dep.cell(row=row, column=c, value=val)
                        cell.border = gray_border
                        if c <= 4:
                            cell.fill = dep_fill
                    row += 1

                for extra_ws in [ws, ws_dep]:
                    for c in range(1, extra_ws.max_column + 1):
                        extra_ws.column_dimensions[get_column_letter(c)].width = 18

                rot_headers = [
                    _('Año'),
                    _('Mes (yyyy-mm)'),
                    _('Mes (mmm)'),
                    _('Población inicial'),
                    _('Altas'),
                    _('Bajas'),
                    _('Población final'),
                    _('% Rotación acumulada'),
                ]
                ws_rot = wb.create_sheet(
                    _sanitize_sheet_name(_('Rotación de personal')),
                    len(wb.worksheets),
                )
                for c, h in enumerate(rot_headers, 1):
                    cell = ws_rot.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border

                monthly_rows = []
                cur_m = date_from.replace(day=1)
                end_m = date_to.replace(day=1)
                while cur_m <= end_m:
                    prev_last = cur_m - relativedelta(days=1)
                    month_last = cur_m + relativedelta(months=1) - relativedelta(days=1)
                    d_ini = self._hd_count_active_headcount_on_date(prev_last)
                    e_alt = self._hd_count_contract_starts(cur_m, month_last)
                    f_baj = self._hd_count_contract_ends_union(cur_m, month_last)
                    g_fin = d_ini + e_alt - f_baj
                    h_pct = (f_baj / g_fin) if g_fin else 0.0
                    monthly_rows.append({
                        'year': cur_m.year,
                        'month_first': cur_m,
                        'yyyy_mm': cur_m.strftime('%Y-%m'),
                        'mmm': self._MONTH_ABBR_ES[cur_m.month - 1],
                        'd': d_ini,
                        'e': e_alt,
                        'f': f_baj,
                        'g': g_fin,
                        'h': h_pct,
                    })
                    cur_m += relativedelta(months=1)

                pct_fmt = '0.00%'
                row_r = 2
                for mr in monthly_rows:
                    ws_rot.cell(row=row_r, column=1, value=mr['month_first'].year).border = gray_border
                    ws_rot.cell(row=row_r, column=2, value=mr['yyyy_mm']).border = gray_border
                    ws_rot.cell(row=row_r, column=3, value=mr['mmm']).border = gray_border
                    ws_rot.cell(row=row_r, column=4, value=mr['d']).border = gray_border
                    ws_rot.cell(row=row_r, column=5, value=mr['e']).border = gray_border
                    ws_rot.cell(row=row_r, column=6, value=mr['f']).border = gray_border
                    ws_rot.cell(row=row_r, column=7, value=mr['g']).border = gray_border
                    c_h = ws_rot.cell(row=row_r, column=8, value=mr['h'])
                    c_h.border = gray_border
                    c_h.number_format = pct_fmt
                    row_r += 1
                for c in range(1, 9):
                    ws_rot.column_dimensions[get_column_letter(c)].width = 18

                by_year = defaultdict(lambda: {'e': 0, 'f': 0, 'months': []})
                for mr in monthly_rows:
                    y = mr['year']
                    by_year[y]['e'] += mr['e']
                    by_year[y]['f'] += mr['f']
                    by_year[y]['months'].append(mr)

                ws_rot_y = wb.create_sheet(
                    _sanitize_sheet_name(_('Rotación Anual')),
                    len(wb.worksheets),
                )
                for c, h in enumerate(rot_headers, 1):
                    cell = ws_rot_y.cell(row=1, column=c, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = gray_border
                row_y = 2
                for y in sorted(by_year.keys()):
                    months_y = by_year[y]['months']
                    d_y = months_y[0]['d']
                    e_y = by_year[y]['e']
                    f_y = by_year[y]['f']
                    g_y = d_y + e_y - f_y
                    h_y = (f_y / g_y) if g_y else 0.0
                    ws_rot_y.cell(row=row_y, column=1, value=y).border = gray_border
                    ws_rot_y.cell(row=row_y, column=2, value='').border = gray_border
                    ws_rot_y.cell(row=row_y, column=3, value=_('Anual')).border = gray_border
                    ws_rot_y.cell(row=row_y, column=4, value=d_y).border = gray_border
                    ws_rot_y.cell(row=row_y, column=5, value=e_y).border = gray_border
                    ws_rot_y.cell(row=row_y, column=6, value=f_y).border = gray_border
                    ws_rot_y.cell(row=row_y, column=7, value=g_y).border = gray_border
                    c_hy = ws_rot_y.cell(row=row_y, column=8, value=h_y)
                    c_hy.border = gray_border
                    c_hy.number_format = pct_fmt
                    row_y += 1
                for c in range(1, 9):
                    ws_rot_y.column_dimensions[get_column_letter(c)].width = 18

                ws_summary = wb.create_sheet(
                    _sanitize_sheet_name(_('Resumen')),
                    len(wb.worksheets),
                )
                years = list(range(date_from.year, date_to.year + 1))
                reasons = self.env['hr.departure.reason'].search([], order='name')
                h1 = ws_summary.cell(row=1, column=1, value=_('Motivos de salida'))
                h1.fill = header_fill
                h1.font = header_font
                h1.border = gray_border
                for idx, yr in enumerate(years, start=2):
                    hc = ws_summary.cell(row=1, column=idx, value=yr)
                    hc.fill = header_fill
                    hc.font = header_font
                    hc.border = gray_border
                rsum = 2
                for reason in reasons:
                    ws_summary.cell(row=rsum, column=1, value=reason.name).border = gray_border
                    for ci, yr in enumerate(years, start=2):
                        cnt = self._hd_count_departures_by_reason_year(yr, reason.id)
                        ws_summary.cell(row=rsum, column=ci, value=cnt).border = gray_border
                    rsum += 1
                ws_summary.cell(row=rsum, column=1, value=_('Sin motivo')).border = gray_border
                for ci, yr in enumerate(years, start=2):
                    cnt = self._hd_count_departures_by_reason_year(yr, None)
                    ws_summary.cell(row=rsum, column=ci, value=cnt).border = gray_border
                for c in range(1, len(years) + 2):
                    ws_summary.column_dimensions[get_column_letter(c)].width = 18

                if ws.max_row >= 1:
                    self._apply_excel_table(ws, 7, '')
                if ws_dep.max_row >= 1:
                    self._apply_excel_table(ws_dep, 8, '')
                if ws_rot.max_row >= 1:
                    self._apply_excel_table(ws_rot, 8, '')
                if ws_rot_y.max_row >= 1:
                    self._apply_excel_table(ws_rot_y, 8, '')
                if ws_summary.max_row >= 1:
                    self._apply_excel_table(
                        ws_summary, len(years) + 1, '',
                    )

            if report_key not in ('headcount', 'contracts', 'hires_departures'):
                for c in range(1, ws.max_column + 1):
                    ws.column_dimensions[get_column_letter(c)].width = 16
            elif report_key == 'headcount':
                for c in range(1, ws.max_column + 1):
                    ws.column_dimensions[get_column_letter(c)].width = 18
                for ws_extra in [ws_asig, ws_proj]:
                    for c in range(1, ws_extra.max_column + 1):
                        ws_extra.column_dimensions[get_column_letter(c)].width = 18

        self._sanitize_workbook_cells(wb)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        safe_name = re.sub(r'[^\w\s\-]', '', _('Reporte_Empleados'))
        filename = '%s_%s.xlsx' % (safe_name.replace(' ', '_'), datetime.now().strftime('%Y%m%d_%H%M%S'))
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'datas': base64.b64encode(buffer.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'self',
        }

    def _get_employees_with_warnings(self, employees):
        """Filtra empleados que tienen al menos un warning activo (en Python, no SQL)."""
        warning_fields = [
            'has_missing_documents', 'has_expiring_licenses',
            'contract_expiring_soon', 'has_loan_missing_dates_warning',
            'has_parent_type_warning', 'has_trial_period_warning',
        ]
        return employees.filtered(
            lambda emp: any(getattr(emp, f, False) for f in warning_fields)
        )

    def action_view_onscreen(self):
        """Abrir el reporte en pantalla (vista lista) según el tipo seleccionado"""
        self.ensure_one()
        employees = self._get_employees()
        if not employees:
            raise UserError(_('No hay empleados que coincidan con el filtro.'))

        report_types = [(k, name, active) for k, name, active in self._get_report_types() if active]
        if not report_types:
            raise UserError(_('Debe seleccionar al menos un tipo de reporte.'))

        # Si hay varios reportes, abrimos el primero
        report_key = report_types[0][0]
        emp_ids = employees.ids
        if report_key == 'birthdays':
            birthday_emps = self._filter_employees_birthday_range(employees)
            if self.birthday_range_from and self.birthday_range_to and not birthday_emps:
                raise UserError(_('No hay empleados con cumpleaños en el rango seleccionado.'))
            emp_ids = birthday_emps.ids

        view_tallas = self.env.ref(
            'l10n_hn_hr_payroll.view_hr_employee_list_report_tallas', raise_if_not_found=False
        )
        view_birthdays = self.env.ref(
            'l10n_hn_hr_payroll.view_hr_employee_list_report_birthdays', raise_if_not_found=False
        )
        view_seniority = self.env.ref(
            'l10n_hn_hr_payroll.view_hr_employee_list_report_seniority', raise_if_not_found=False
        )

        actions = {
            'tallas': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', emp_ids)],
                'context': {'search_default_active': 1},
                'name': _('Tallas (camisa, pantalón, calzado)'),
                'view_id': view_tallas.id if view_tallas else False,
            },
            'emergency': {
                'res_model': 'hr.employee.emergency.contact',
                'view_mode': 'list',
                'domain': [('employee_id', 'in', emp_ids)],
                'name': _('Contactos de emergencia'),
            },
            'licenses': {
                'res_model': 'hr.employee.license',
                'view_mode': 'list',
                'domain': [('employee_id', 'in', emp_ids)],
                'name': _('Licencias de conducir'),
            },
            'birthdays': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', emp_ids)],
                'context': {'search_default_active': 1},
                'name': _('Cumpleaños'),
                'view_id': view_birthdays.id if view_birthdays else False,
            },
            'family': {
                'res_model': 'hr.employee.family',
                'view_mode': 'list',
                'domain': [('employee_id', 'in', emp_ids)],
                'name': _('Datos familiares'),
            },
            'beneficiaries': {
                'res_model': 'hr.employee.beneficiary',
                'view_mode': 'list',
                'domain': [('employee_id', 'in', emp_ids)],
                'name': _('Beneficiarios'),
            },
            'documents': {
                'res_model': 'hr.employee.document',
                'view_mode': 'list',
                'domain': [('employee_id', 'in', emp_ids)],
                'name': _('Documentos de entrada'),
            },
            'loans': {
                'res_model': 'hr.employee.recurring.concept',
                'view_mode': 'list',
                'domain': [
                    ('employee_id', 'in', emp_ids),
                    ('is_loan', '=', True),
                    ('active', '=', True),
                ],
                'name': _('Préstamos (detalle)'),
            },
            'contracts': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', emp_ids)],
                'context': {'search_default_active': 1},
                'name': _('Registro de contratos'),
            },
            'salary_increases': {
                'res_model': 'hr.historical.salary.increase',
                'view_mode': 'list',
                'domain': [
                    ('employee_id', 'in', emp_ids),
                    ('employee_id.company_id', '=', self.company_id.id),
                ],
                'name': _('Historial de aumentos salariales'),
            },
            'seniority': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', emp_ids)],
                'context': {'search_default_active': 1},
                'name': _('Antigüedad'),
                'view_id': view_seniority.id if view_seniority else False,
            },
            'headcount': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', emp_ids)],
                'context': {'search_default_active': 1, 'group_by': ['department_id']},
                'name': _('Headcount (depto/asignación/proyecto)'),
            },
            'references': {
                'res_model': 'reference.personal',
                'view_mode': 'list',
                'domain': [('employee_id', 'in', emp_ids)],
                'name': _('Referencias'),
            },
            'health_status': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', emp_ids)],
                'context': {'search_default_active': 1},
                'name': _('Estado de Salud'),
            },
            'warnings': {
                'res_model': 'hr.employee',
                'view_mode': 'list',
                'domain': [('id', 'in', self._get_employees_with_warnings(employees).ids)],
                'context': {'search_default_active': 1},
                'name': _('Empleados con Alertas Activas'),
                'view_id': self.env.ref(
                    'l10n_hn_hr_payroll.view_hr_employee_list_report_warnings',
                    raise_if_not_found=False,
                ).id,
            },
        }

        act = actions.get(report_key)
        if not act:
            raise UserError(_('Tipo de reporte no soportado para vista en pantalla: %s') % report_key)

        if report_key == 'references' and 'reference.personal' not in self.env:
            raise UserError(_(
                'El listado de referencias requiere el módulo Honduras Recruitment '
                '(registros en Experiencia y Referencias).'
            ))

        return {
            'type': 'ir.actions.act_window',
            'target': 'current',
            **act,
        }
