# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik

"""
Wizard para importación masiva de HHEE (Horas Extra) y TE (Turnos Extra).
Modelo destino: l10n.hn.payroll.line.he y l10n.hn.payroll.line.te dentro de la planilla.
"""

import base64
import io
import re
from datetime import datetime

from odoo import api, fields, models, _
from odoo.exceptions import UserError

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# Índices de columnas en la plantilla (0-based)
COL_EMPLOYEE_ID = 0
COL_IDENTIDAD = 1
COL_NOMBRE = 2
COL_COMPANIA = 3
COL_DEPARTAMENTO = 4
COL_PERIODO = 5
COL_HE_25 = 6
COL_HE_50 = 7
COL_HE_75 = 8
COL_HE_100 = 9
COL_TE_A = 10
COL_TE_B = 11
COL_TE_C = 12
COL_OBSERVACIONES = 13

HEADERS = [
    'employee_id',
    'Identidad',
    'Nombre Colaborador',
    'Compañía',
    'Departamento',
    'Periodo',
    'HE 25%',
    'HE 50%',
    'HE 75%',
    'HE 100%',
    'Turno A',
    'Turno B',
    'Turno C',
    'Observaciones',
]

# Columnas editables (1-based para openpyxl)
EDITABLE_COLS = (7, 8, 9, 10, 11, 12, 13, 14)  # HE 25%-100%, Turno A-C, Observaciones


class HrHeTeImportWizard(models.TransientModel):
    _name = 'hr.he.te.import.wizard'
    _description = 'Importar HHEE y TE'

    # Contexto: planilla desde la que se abre (opcional)
    payroll_sheet_id = fields.Many2one(
        'l10n.hn.payroll.sheet', string='Planilla',
        help='Si se abre desde una planilla, se pre-llena. Si no, se crea/usa una según período.',
        ondelete='set null',
    )
    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True,
        default=lambda self: self.env.company,
    )
    payroll_period_id = fields.Many2one(
        'l10n.hn.payroll.period', string='Período',
        domain="[('company_id', '=', company_id)]",
        help='Período de nómina (quincena) para la importación.',
    )
    department_id = fields.Many2one(
        'hr.department', string='Departamento',
        domain="['|', ('company_id', '=', False), ('company_id', '=', company_id)]",
        help='Opcional: filtrar empleados por departamento en la plantilla.',
    )
    employee_ids = fields.Many2many(
        'hr.employee', string='Empleados',
        domain="[('active', '=', True), ('company_id', '=', company_id)]",
        help='Opcional: limitar plantilla a estos empleados.',
    )
    file_data = fields.Binary(string='Archivo Excel', attachment=False)
    file_name = fields.Char(string='Nombre archivo')
    skip_empty = fields.Boolean(
        string='Omitir filas con HE y TE vacíos',
        default=True,
        help='Si está activo, las filas con todas las HE y Turnos en 0 se omiten.',
    )
    # Resultados
    result_summary = fields.Text(string='Resumen', readonly=True)
    result_log = fields.Binary(string='Log de importación', readonly=True, attachment=False)
    result_log_filename = fields.Char(string='Nombre log', readonly=True)
    state = fields.Selection([
        ('draft', 'Configuración'),
        ('done', 'Completado'),
    ], default='draft', required=True)

    @api.onchange('payroll_sheet_id')
    def _onchange_payroll_sheet_id(self):
        if self.payroll_sheet_id:
            self.company_id = self.payroll_sheet_id.company_id
            self.payroll_period_id = self.payroll_sheet_id.payroll_period_id
            self.department_id = self.payroll_sheet_id.department_id

    def _get_employees_for_template(self):
        """Empleados a incluir en la plantilla según filtros."""
        domain = [
            ('active', '=', True),
            ('company_id', '=', self.company_id.id),
        ]
        if self.department_id:
            domain.append(('department_id', '=', self.department_id.id))
        if self.employee_ids:
            domain.append(('id', 'in', self.employee_ids.ids))
        employees = self.env['hr.employee'].search(domain, order='name')
        if not employees:
            raise UserError(_('No hay empleados que coincidan con los filtros.'))
        # Filtrar por versión activa en el período (usar período de planilla si aplica)
        period = self.payroll_period_id or (self.payroll_sheet_id.payroll_period_id if self.payroll_sheet_id else None)
        if period and period.date_start and period.date_end:
            versions = self.env['hr.version'].search([
                ('employee_id', 'in', employees.ids),
                ('contract_date_start', '<=', period.date_end),
                '|', ('contract_date_end', '=', False),
                ('contract_date_end', '>=', period.date_start),
            ])
            emp_ids = versions.mapped('employee_id').ids
            employees = employees.filtered(lambda e: e.id in emp_ids)
        return employees

    def action_download_template(self):
        """Generar y descargar plantilla Excel."""
        if not OPENPYXL_AVAILABLE:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))
        period = self.payroll_period_id or (self.payroll_sheet_id.payroll_period_id if self.payroll_sheet_id else None)
        if not period:
            raise UserError(_('Seleccione el período o una planilla antes de descargar la plantilla.'))
        employees = self._get_employees_for_template()
        wb = self._build_template_workbook(employees)
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        period_name = self.payroll_period_id.name or 'plantilla'
        filename = f'Importacion_HHEE_TE_{period_name.replace(" ", "_")}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def _build_template_workbook(self, employees):
        """Construir workbook con formato, freeze panes, autofilter."""
        wb = Workbook()
        ws = wb.active
        ws.title = 'Importación'
        # Estilos
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        editable_fill = PatternFill(start_color='E8F4EA', end_color='E8F4EA', fill_type='solid')  # Verde claro
        readonly_fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')  # Gris
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        # Encabezados (fila 1)
        for c, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
        # Datos (fila 2 en adelante)
        period = self.payroll_period_id or (self.payroll_sheet_id.payroll_period_id if self.payroll_sheet_id else None)
        period_str = f'{period.date_start} - {period.date_end}' if period and period.date_start else ''
        for r, emp in enumerate(employees, 2):
            row_data = [
                emp.id,
                emp.identification_id or '',
                emp.name or '',
                emp.company_id.name or '',
                emp.department_id.name or '',
                period_str,
                0.0, 0.0, 0.0, 0.0,  # HE 25%, 50%, 75%, 100%
                0, 0, 0,  # Turno A, B, C
                '',
            ]
            for c, val in enumerate(row_data, 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.border = thin_border
                if c in EDITABLE_COLS:
                    cell.fill = editable_fill
                    if 7 <= c <= 10:  # HE: decimal
                        cell.number_format = '0.00'
                    elif 11 <= c <= 13:  # Turnos: entero
                        cell.number_format = '0'
                else:
                    cell.fill = readonly_fill
        # Freeze panes
        ws.freeze_panes = 'A2'
        # Autofilter
        if employees:
            ws.auto_filter.ref = ws.dimensions
        # Ajustar anchos
        widths = [12, 15, 30, 20, 20, 22, 10, 10, 10, 10, 10, 10, 10, 25]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        # Hoja Catálogos (opcional)
        ws_cat = wb.create_sheet('Catálogos')
        ws_cat['A1'] = _('Períodos disponibles')
        ws_cat['A1'].font = Font(bold=True)
        periods = self.env['l10n.hn.payroll.period'].search([
            ('company_id', '=', self.company_id.id),
        ], order='year desc, month desc', limit=12)
        for r, p in enumerate(periods, 2):
            ws_cat.cell(row=r, column=1, value=p.name or '')
            ws_cat.cell(row=r, column=2, value=f'{p.date_start} - {p.date_end}' if p.date_start else '')
        ws_cat.column_dimensions['A'].width = 35
        ws_cat.column_dimensions['B'].width = 25
        return wb

    def _get_or_create_sheet(self):
        """Obtener o crear planilla en borrador para el período."""
        if self.payroll_sheet_id:
            if self.payroll_sheet_id.state != 'draft':
                raise UserError(_('La planilla debe estar en borrador para importar.'))
            return self.payroll_sheet_id
        if not self.payroll_period_id:
            raise UserError(_('Seleccione el período o una planilla.'))
        sheet = self.env['l10n.hn.payroll.sheet'].search([
            ('payroll_period_id', '=', self.payroll_period_id.id),
            ('company_id', '=', self.company_id.id),
            ('state', '=', 'draft'),
        ], limit=1)
        if not sheet:
            sheet = self.env['l10n.hn.payroll.sheet'].create({
                'payroll_period_id': self.payroll_period_id.id,
                'company_id': self.company_id.id,
            })
        return sheet

    def _resolve_employee(self, row_data, row_num):
        """Resolver empleado desde fila. Prioridad: employee_id, luego Identidad."""
        emp_id_val = row_data.get(COL_EMPLOYEE_ID)
        identidad_val = str(row_data.get(COL_IDENTIDAD) or '').strip()
        if emp_id_val is not None and emp_id_val != '':
            try:
                eid = int(float(emp_id_val))
                emp = self.env['hr.employee'].browse(eid)
                if emp.exists():
                    return emp
            except (ValueError, TypeError):
                pass
        if identidad_val:
            identidad_clean = re.sub(r'[^0-9]', '', identidad_val)
            emp = self.env['hr.employee'].search([
                ('company_id', '=', self.company_id.id),
                '|', ('identification_id', '=', identidad_val),
                ('identification_id', '=', identidad_clean),
            ], limit=1)
            if emp:
                return emp
        return None

    def _parse_numeric(self, val, col_name):
        """Parsear valor numérico. Retorna float o None."""
        if val is None or val == '':
            return 0.0
        try:
            return float(str(val).replace(',', '.').strip())
        except (ValueError, TypeError):
            return None

    def action_import_file(self):
        """Procesar archivo Excel subido."""
        if not self.file_data:
            raise UserError(_('Suba un archivo Excel antes de importar.'))
        if not OPENPYXL_AVAILABLE:
            raise UserError(_('Se requiere openpyxl. Instale: pip install openpyxl'))
        sheet = self._get_or_create_sheet()
        data = base64.b64decode(self.file_data)
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        if not ws:
            raise UserError(_('El archivo no contiene hojas válidas.'))
        created_he = created_te = updated_he = updated_te = skipped = 0
        errors = []
        log_lines = []
        # Empezar en fila 2 (1=headers)
        num_cols = 14
        for row_num in range(2, ws.max_row + 1):
            row_data = []
            for c in range(1, num_cols + 1):
                cell = ws.cell(row=row_num, column=c)
                row_data.append(cell.value)
            emp = self._resolve_employee(dict(zip(range(num_cols), row_data)), row_num)
            if not emp:
                nombre = str(row_data[COL_NOMBRE] or '').strip() if len(row_data) > COL_NOMBRE else ''
                nombre_str = f' "{nombre}"' if nombre else ''
                err = _('Fila %d%s: No se encontró empleado (employee_id o Identidad inválidos)') % (row_num, nombre_str)
                errors.append(err)
                log_lines.append(err)
                continue
            # Parsear HE 25%, 50%, 75%, 100%
            he_vals = []
            he_ok = True
            for i, col in enumerate([COL_HE_25, COL_HE_50, COL_HE_75, COL_HE_100]):
                v = self._parse_numeric(row_data[col] if len(row_data) > col else None, f'HE')
                if v is None:
                    err = _('Fila %d: HE 25%%/50%%/75%%/100%% deben ser numéricos') % row_num
                    errors.append(err)
                    log_lines.append(err)
                    he_ok = False
                    break
                he_vals.append(v)
            if not he_ok or len(he_vals) != 4:
                continue
            # Parsear Turno A, B, C (enteros)
            te_vals = []
            te_ok = True
            for i, col in enumerate([COL_TE_A, COL_TE_B, COL_TE_C]):
                v = self._parse_numeric(row_data[col] if len(row_data) > col else None, f'Turno')
                if v is None:
                    err = _('Fila %d: Turno A/B/C deben ser numéricos') % row_num
                    errors.append(err)
                    log_lines.append(err)
                    te_ok = False
                    break
                te_vals.append(int(round(v)))
            if not te_ok or len(te_vals) != 3:
                continue
            # Omitir si skip_empty y todo en 0
            if self.skip_empty and sum(he_vals) == 0 and sum(te_vals) == 0:
                skipped += 1
                continue
            # Crear/actualizar HE
            if sum(he_vals) != 0:
                line_he = self.env['l10n.hn.payroll.line.he'].search([
                    ('parent_id', '=', sheet.id),
                    ('employee_id', '=', emp.id),
                ], limit=1)
                vals_he = {'he_25': he_vals[0], 'he_50': he_vals[1], 'he_75': he_vals[2], 'he_100': he_vals[3]}
                if line_he:
                    line_he.write(vals_he)
                    updated_he += 1
                else:
                    self.env['l10n.hn.payroll.line.he'].create({
                        'parent_id': sheet.id,
                        'employee_id': emp.id,
                        **vals_he,
                    })
                    created_he += 1
            # Crear/actualizar TE
            if sum(te_vals) != 0:
                line_te = self.env['l10n.hn.payroll.line.te'].search([
                    ('parent_id', '=', sheet.id),
                    ('employee_id', '=', emp.id),
                ], limit=1)
                vals_te = {'turn_a': te_vals[0], 'turn_b': te_vals[1], 'turn_c': te_vals[2]}
                if line_te:
                    line_te.write(vals_te)
                    updated_te += 1
                else:
                    self.env['l10n.hn.payroll.line.te'].create({
                        'parent_id': sheet.id,
                        'employee_id': emp.id,
                        **vals_te,
                    })
                    created_te += 1
        wb.close()
        # Resumen
        summary = _('Importación completada.\n\n')
        summary += _('Creados HE: %d\n') % created_he
        summary += _('Actualizados HE: %d\n') % updated_he
        summary += _('Creados TE: %d\n') % created_te
        summary += _('Actualizados TE: %d\n') % updated_te
        summary += _('Omitidos: %d\n') % skipped
        summary += _('Errores: %d\n') % len(errors)
        if errors:
            summary += _('\n--- Errores ---\n') + '\n'.join(errors[:20])
            if len(errors) > 20:
                summary += _('\n... y %d más (ver log)') % (len(errors) - 20)
        log_content = '\n'.join(log_lines) if log_lines else summary
        self.write({
            'state': 'done',
            'result_summary': summary,
            'result_log': base64.b64encode(log_content.encode('utf-8')),
            'result_log_filename': f'import_he_te_log_{datetime.now().strftime("%Y%m%d_%H%M%S")}.txt',
        })
        sheet.message_post(body=_('Importación HE/TE: %s') % summary)
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_download_log(self):
        """Descargar log de importación."""
        if not self.result_log:
            raise UserError(_('No hay log para descargar.'))
        attachment = self.env['ir.attachment'].create({
            'name': self.result_log_filename or 'import_log.txt',
            'type': 'binary',
            'datas': self.result_log,
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }
