# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik

import base64
import io
from calendar import monthrange
from datetime import date

from markupsafe import escape

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    get_column_letter = None  # type: ignore

IMPORT_SHEET = 'Historial'
IMPORT_HEADERS = [
    'identidad',
    'codigo_empleado',
    'año',
    'mes',
    'neto_nómina',
    'comisiones',
]


class L10nHnMonthlySalaryHistoryImportWizard(models.TransientModel):
    _name = 'l10n.hn.monthly.salary.history.import.wizard'
    _description = 'Importar historial salarial mensual'

    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True, default=lambda self: self.env.company,
    )
    file_data = fields.Binary(string='Archivo Excel')
    file_name = fields.Char(string='Nombre del archivo')
    overwrite_existing = fields.Boolean(
        string='Sobrescribir registros existentes',
        default=False,
        help='Si está marcado, reemplaza filas del mismo empleado/mes aunque provengan de nómina o importación previa.',
    )
    result_message = fields.Html(string='Resultado', readonly=True)
    state = fields.Selection([('draft', 'Borrador'), ('done', 'Hecho')], default='draft')

    def _normalize_ident(self, val):
        if val is None:
            return ''
        s = str(val).strip().replace('-', '').replace(' ', '')
        if s.endswith('.0'):
            s = s[:-2]
        return s

    def _row_employee(self, company, ident, code):
        Employee = self.env['hr.employee'].sudo()
        domain = [('company_id', 'child_of', company.id)]
        emp = Employee.browse()
        if ident:
            ident = self._normalize_ident(ident)
            emp = Employee.search(domain + [('identification_id', '=', ident)], limit=1)
        if not emp and code:
            code = str(code).strip()
            emp = Employee.search(domain + [('employee_code', '=', code)], limit=1)
        return emp

    def action_download_template(self):
        self.ensure_one()
        if not OPENPYXL_AVAILABLE:
            raise UserError(_('Instale openpyxl.'))
        wb = Workbook(write_only=False)
        ws = wb.active
        ws.title = IMPORT_SHEET

        hdr = [
            'identidad',
            'codigo_empleado',
            'nombre',
            'departamento',
            'puesto',
            'fecha_ingreso',
            'año',
            'mes',
            'neto_nómina',
            'comisiones',
        ]
        header_fill = PatternFill(fill_type='solid', fgColor='FF366092')
        header_font = Font(bold=True, color='FFFFFF')
        header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin = Side(style='thin', color='CCCCCC')
        cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        num_align = Alignment(horizontal='right', vertical='center')

        for col, title in enumerate(hdr, 1):
            c = ws.cell(row=1, column=col, value=title)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = cell_border

        today = date.today()
        year_d, month_d = today.year, today.month
        Employee = self.env['hr.employee'].sudo()
        employees = Employee.search([
            ('company_id', 'child_of', self.company_id.id),
            ('active', '=', True),
        ], order='name')

        date_align = Alignment(horizontal='center', vertical='center')
        row = 2
        for emp in employees:
            dept = emp.department_id.name if emp.department_id else ''
            job = emp.job_id.name if emp.job_id else ''
            hire = emp._get_first_contract_date(no_gap=False)
            ws.cell(row=row, column=1, value=emp.identification_id or '').border = cell_border
            ws.cell(row=row, column=2, value=emp.employee_code or '').border = cell_border
            ws.cell(row=row, column=3, value=emp.name or '').border = cell_border
            ws.cell(row=row, column=4, value=dept).border = cell_border
            ws.cell(row=row, column=5, value=job).border = cell_border
            if hire:
                c_hire = ws.cell(row=row, column=6, value=hire)
                c_hire.number_format = 'dd/mm/yyyy'
            else:
                c_hire = ws.cell(row=row, column=6, value='')
            c_hire.border = cell_border
            c_hire.alignment = date_align
            c_year = ws.cell(row=row, column=7, value=year_d)
            c_year.border = cell_border
            c_year.alignment = num_align
            c_month = ws.cell(row=row, column=8, value=month_d)
            c_month.border = cell_border
            c_month.alignment = num_align
            c_net = ws.cell(row=row, column=9, value=0.0)
            c_net.border = cell_border
            c_net.alignment = num_align
            c_comm = ws.cell(row=row, column=10, value=0.0)
            c_comm.border = cell_border
            c_comm.alignment = num_align
            row += 1

        ws.freeze_panes = 'A2'
        if row > 2:
            ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(hdr)), row - 1)
        elif not employees:
            merge_to = '%s2' % get_column_letter(len(hdr))
            ws.merge_cells('A2:%s' % merge_to)
            c = ws.cell(row=2, column=1, value=_('No hay empleados activos en esta compañía. Complete el archivo manualmente.'))
            c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            c.border = cell_border
        widths = [14, 14, 32, 24, 24, 14, 8, 6, 14, 12]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 22

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        safe_company = (self.company_id.name or 'empresa').replace('/', '-')[:40]
        filename = 'plantilla_historial_salarial_%s.xlsx' % safe_company
        att = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(buf.read()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % att.id,
            'target': 'self',
        }

    def action_import(self):
        self.ensure_one()
        if not OPENPYXL_AVAILABLE:
            raise UserError(_('Instale openpyxl.'))
        if not self.file_data:
            raise UserError(_('Seleccione un archivo.'))
        History = self.env['l10n.hn.employee.monthly.salary.history'].sudo()
        try:
            data = base64.b64decode(self.file_data)
            wb = load_workbook(filename=io.BytesIO(data), data_only=True)
        except Exception as err:
            raise UserError(_('No se pudo leer el Excel: %s') % err) from err
        if IMPORT_SHEET not in wb.sheetnames:
            raise UserError(_('Falta la hoja "%s". Hojas: %s') % (IMPORT_SHEET, ', '.join(wb.sheetnames)))
        ws = wb[IMPORT_SHEET]
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        if not rows:
            raise UserError(_('El archivo está vacío.'))
        header = [str(h).lower().strip() if h is not None else '' for h in rows[0]]
        # header flex: require identidad OR codigo, año, mes, neto
        idx = {h: i for i, h in enumerate(header)}
        required = ['año', 'mes', 'neto_nómina']
        for r in required:
            if r not in idx:
                raise UserError(_('Falta la columna requerida: %s') % r)
        if 'identidad' not in idx and 'codigo_empleado' not in idx:
            raise UserError(_('Debe existir la columna identidad y/o codigo_empleado'))

        ok, errs = 0, []
        for rnum, row in enumerate(rows[1:], start=2):
            if not row or all(v is None or str(v).strip() == '' for v in row):
                continue
            def gv(key):
                i = idx.get(key)
                if i is None or i >= len(row):
                    return None
                return row[i]

            ident = gv('identidad')
            code = gv('codigo_empleado')
            try:
                year = int(gv('año'))
                month = int(gv('mes'))
            except (TypeError, ValueError):
                errs.append(_('Fila %s: año/mes inválidos') % rnum)
                continue
            if month < 1 or month > 12:
                errs.append(_('Fila %s: mes fuera de rango') % rnum)
                continue
            net = gv('neto_nómina')
            comm = gv('comisiones') or 0.0
            try:
                net = float(net)
                comm = float(comm)
            except (TypeError, ValueError):
                errs.append(_('Fila %s: neto/comisiones deben ser numéricos') % rnum)
                continue

            emp = self._row_employee(self.company_id, ident, code)
            if not emp:
                errs.append(_('Fila %s: empleado no encontrado (identidad/código)') % rnum)
                continue
            if emp.company_id != self.company_id:
                errs.append(
                    _('Fila %s: el empleado pertenece a otra compañía (%s)') % (
                        rnum, emp.company_id.display_name,
                    )
                )
                continue

            existing = History.search([
                ('employee_id', '=', emp.id),
                ('year', '=', year),
                ('month', '=', month),
            ], limit=1)
            if existing and not self.overwrite_existing:
                errs.append(_('Fila %s: ya existe registro para %s %04d-%02d (use sobrescribir)') % (
                    rnum, emp.name, year, month,
                ))
                continue
            vals = {
                'employee_id': emp.id,
                'company_id': self.company_id.id,
                'year': year,
                'month': month,
                'amount_net': net,
                'amount_commission': comm,
                'data_source': 'import',
                'payslip_ids': [(5, 0, 0)],
            }
            if existing:
                existing.write(vals)
                existing._apply_snapshot_from_employee(emp)
            else:
                rec = History.create(vals)
                rec._apply_snapshot_from_employee(emp)
            ok += 1

        msg = '<p><strong>%s</strong> filas importadas.</p>' % ok
        if errs:
            msg += '<p><strong>Advertencias / errores:</strong></p><ul>%s</ul>' % ''.join(
                '<li>%s</li>' % escape(str(e)) for e in errs[:200]
            )
            if len(errs) > 200:
                msg += '<p>… y %s más.</p>' % (len(errs) - 200)
        self.result_message = msg
        self.state = 'done'
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }


class L10nHnMonthlySalaryHistoryGenerateWizard(models.TransientModel):
    _name = 'l10n.hn.monthly.salary.history.generate.wizard'
    _description = 'Generar historial salarial desde nómina'

    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True, default=lambda self: self.env.company,
    )
    year_from = fields.Integer(string='Año desde', required=True, default=lambda self: date.today().year)
    month_from = fields.Integer(string='Mes desde', required=True, default=lambda self: date.today().month)
    year_to = fields.Integer(string='Año hasta', required=True, default=lambda self: date.today().year)
    month_to = fields.Integer(string='Mes hasta', required=True, default=lambda self: date.today().month)
    employee_ids = fields.Many2many(
        'hr.employee', string='Empleados',
        domain="[('company_id', 'child_of', company_id)]",
        help='Vacío = todos los empleados de la compañía.',
    )
    payslip_states = fields.Selection(
        [
            ('validated', 'Validados'),
            ('paid', 'Pagados'),
            ('both', 'Validados y pagados'),
        ],
        string='Recibos a incluir',
        default='both',
        required=True,
    )
    overwrite_import = fields.Boolean(
        string='Sobrescribir filas importadas/manuales',
        default=False,
        help='Si no está marcado, no modifica registros cuyo origen sea importación o manual.',
    )
    result_message = fields.Text(string='Resultado', readonly=True)
    state = fields.Selection([('draft', 'Borrador'), ('done', 'Hecho')], default='draft')

    @api.constrains('year_from', 'month_from', 'year_to', 'month_to')
    def _check_range(self):
        for wiz in self:
            d1 = date(wiz.year_from, wiz.month_from, 1)
            d2 = date(wiz.year_to, wiz.month_to, 1)
            if d1 > d2:
                raise ValidationError(_('El rango de fechas es inválido (desde > hasta).'))

    def action_generate(self):
        self.ensure_one()
        if self.payslip_states == 'validated':
            states = ('validated',)
        elif self.payslip_states == 'paid':
            states = ('paid',)
        else:
            states = ('validated', 'paid')
        emps = self.employee_ids if self.employee_ids else None
        n = self.env['l10n.hn.employee.monthly.salary.history'].regenerate_range(
            self.company_id,
            self.year_from,
            self.month_from,
            self.year_to,
            self.month_to,
            employee_ids=emps,
            states=states,
            overwrite_import=self.overwrite_import,
        )
        self.result_message = _(
            'Proceso terminado. Registros generados o actualizados (con recibos en el período): %s'
        ) % n
        self.state = 'done'
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
