# -*- coding: utf-8 -*-
# © 2025 Grupo Nahuiik
# License: LGPL-3

import base64
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:
    openpyxl = None


class HnInitialBalanceLine(models.TransientModel):
    _name = 'hn.initial.balance.line'
    _description = 'Línea de saldo inicial vacaciones HN'

    wizard_id = fields.Many2one(
        'hn.initial.balance.wizard',
        string='Wizard',
        required=True,
        ondelete='cascade',
    )
    employee_id = fields.Many2one(
        'hr.employee',
        string='Empleado',
        required=True,
    )
    vacation_taken = fields.Float(
        string='Días ya gozados',
        default=0,
        help='Días de vacaciones tomados antes de la fecha de corte',
    )


class HnInitialBalanceWizard(models.TransientModel):
    _name = 'hn.initial.balance.wizard'
    _description = 'Cargar saldos iniciales de vacaciones HN'

    balance_date = fields.Date(
        string='Fecha de corte',
        required=True,
        default=fields.Date.context_today,
        help='Fecha a la cual corresponde el saldo inicial',
    )
    excel_file = fields.Binary('Archivo Excel')
    excel_filename = fields.Char()
    import_mode = fields.Boolean('Modo importación Excel', default=False)
    line_ids = fields.One2many(
        'hn.initial.balance.line',
        'wizard_id',
        string='Empleados',
    )

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        if 'line_ids' in fields_list and self.env.context.get('active_model') == 'hr.employee':
            employee_ids = self.env.context.get('active_ids', [])
            if employee_ids:
                employees = self.env['hr.employee'].browse(employee_ids)
                res['line_ids'] = [
                    (0, 0, {
                        'employee_id': emp.id,
                        'vacation_taken': emp.hn_initial_vacation_taken_display or 0,
                    })
                    for emp in employees
                ]
        return res

    def action_import_excel(self):
        """Parsea Excel y carga line_ids.
        Columnas: A=Código, B=Departamento, C=DNI, D=Nombre, E=Puesto, F=Días ya tomados."""
        self.ensure_one()
        if not self.excel_file:
            raise UserError(_('Seleccione un archivo Excel.'))
        if not openpyxl:
            raise UserError(_('El módulo openpyxl no está instalado.'))
        data = base64.b64decode(self.excel_file)
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        ws = wb.active
        lines = []
        employees = self.env['hr.employee'].search([
            ('active', '=', True),
            ('company_id', 'in', self.env.companies.ids),
        ])
        emp_by_code = {}
        emp_by_dni = {}
        for e in employees:
            code = getattr(e, 'employee_code', None) or getattr(e, 'registration_number', None) or ''
            if code:
                emp_by_code[str(code).strip()] = e
            dni = getattr(e, 'identification_id', None) or ''
            if dni:
                emp_by_dni[str(dni).strip()] = e
        emp_by_code.update({str(e.id): e for e in employees})
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 2:
                continue
            # A=0 Código, B=1 Depto, C=2 DNI, D=3 Nombre, E=4 Puesto, F=5 Días
            key = str(row[0]).strip() if row[0] else ''
            if not key and len(row) > 2:
                key = str(row[2]).strip() if row[2] else ''  # DNI como alternativa
            days_col = row[5] if len(row) > 5 else (row[3] if len(row) > 3 else row[1])
            try:
                days = float(days_col or 0)
            except (TypeError, ValueError):
                continue
            if not key:
                continue
            emp = emp_by_code.get(key) or emp_by_dni.get(key)

            if emp:
                lines.append((0, 0, {'employee_id': emp.id, 'vacation_taken': days}))
        self.line_ids = [(5, 0, 0)] + lines
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_download_template(self):
        """Genera y descarga plantilla Excel con todos los empleados activos."""
        self.ensure_one()
        if not openpyxl:
            raise UserError(_('El módulo openpyxl no está instalado.'))
        from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
        from openpyxl.utils import get_column_letter

        employees = self.env['hr.employee'].search([
            ('active', '=', True),
            ('company_id', 'in', self.env.companies.ids),
        ], order='name')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Empleados'
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        headers = ['Código empleado', 'Departamento', 'DNI', 'Nombre empleado', 'Puesto', 'Días ya tomados']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        for row_idx, emp in enumerate(employees, 2):
            code = getattr(emp, 'employee_code', None) or getattr(emp, 'registration_number', None) or ''
            dept = emp.department_id.name if emp.department_id else ''
            dni = getattr(emp, 'identification_id', None) or ''
            puesto = emp.job_id.name if emp.job_id else ''
            for col_idx, val in enumerate([
                code,
                dept,
                dni,
                emp.name or '',
                puesto,
                emp.hn_initial_vacation_taken_display or 0,
            ], 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 10
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    val_len = len(str(cell.value or ''))
                    if val_len > max_len:
                        max_len = val_len
            ws.column_dimensions[col_letter].width = min(max_len + 2, 50)
        ws2 = wb.create_sheet('Instrucciones', 1)
        ws2['A1'] = 'Instrucciones de importación'
        ws2['A2'] = '1. Use la columna A "Código empleado" O columna C "DNI" para identificar al empleado.'
        ws2['A3'] = '2. Complete la columna F "Días ya tomados" con los días de vacaciones tomados antes de la fecha de corte.'
        ws2['A4'] = '3. Guarde el archivo y use "Importar desde Excel" en el wizard.'
        ws2['A5'] = '4. Puede eliminar filas de empleados que no desee cargar.'
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        data = base64.b64encode(buffer.getvalue())
        attachment = self.env['ir.attachment'].create({
            'name': 'Plantilla_saldos_iniciales_vacaciones.xlsx',
            'type': 'binary',
            'datas': data,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true&filename={attachment.name}',
            'target': 'self',
        }

    def action_apply(self):
        self.ensure_one()
        for line in self.line_ids:
            line.employee_id.write({
                'hn_initial_vacation_taken': line.vacation_taken,
                'hn_initial_balance_date': self.balance_date,
                'hn_initial_balance_loaded': True,
            })
            # Usar valor efectivo tras write (puede ser 0 si contrato no permanente)
            taken = line.employee_id.hn_initial_vacation_taken or 0
            line.employee_id.message_post(
                body=f'Saldo inicial de vacaciones cargado: {taken} días ya gozados '
                     f'(fecha de corte: {self.balance_date})',
                subject='Saldo inicial vacaciones',
            )
        return {'type': 'ir.actions.act_window_close'}
