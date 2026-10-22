# -*- coding: utf-8 -*-
"""Excel .xlsx (openpyxl) — extractos Banco Atlántida, Banpaís, Occidente HN."""

import datetime
import io

from odoo import _
from odoo.exceptions import UserError


def _require_openpyxl():
    try:
        from openpyxl import load_workbook  # noqa: F401,WPS433
    except ImportError as err:
        raise UserError(
            _('Falta «openpyxl» en el servidor. '
              'Instálela o añádala en external_dependencies del módulo gn_qif_openai_reconcile.'),
        ) from err


def workbook_from_bytes(raw_bytes):
    """openpyxl Workbook desde bytes (.xlsx)."""
    _require_openpyxl()
    from openpyxl import load_workbook

    bio = io.BytesIO(raw_bytes or b'')
    try:
        return load_workbook(bio, data_only=True, read_only=False)
    except Exception as err:
        raise UserError(_('No se pudo abrir el archivo Excel (.xlsx): %s') % err) from err


def detect_xlsx_extract_format(raw_bytes):
    """Devuelve clave técnica o cadena vacía si no coincide con ningún layout conocido."""
    wb = workbook_from_bytes(raw_bytes)
    try:
        sh = wb.active
        r2a = norm_cell(sh.cell(2, 1).value).upper().replace('Á', 'A')
        if 'ATLANTIDA' in r2a:
            return 'xlsx_atlantida'
        if norm_cell(sh.cell(1, 1).value) == 'Cuenta:':
            r7 = norm_cell(sh.cell(7, 1).value).upper()
            if r7 == 'FECHA':
                return 'xlsx_banpais'
        mr = min(22, sh.max_row or 0)
        for r in range(1, mr + 1):
            fb = norm_cell(sh.cell(r, 2).value).lower()
            debit_header = ''
            if sh.max_column >= 16:
                debit_header = norm_cell(sh.cell(r, 16).value).lower()
            credit_header = ''
            if sh.max_column >= 18:
                credit_header = norm_cell(sh.cell(r, 18).value).lower()
            saldo_header = ''
            if sh.max_column >= 20:
                saldo_header = norm_cell(sh.cell(r, 20).value).lower()
            if fb == 'fecha' and 'bit' in debit_header and ('redit' in credit_header or credit_header.startswith('cred')):
                if 'aldo' in saldo_header:
                    return 'xlsx_occidente_hn'
        return ''
    finally:
        wb.close()


def norm_cell(val):
    if val is None:
        return ''
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        return str(val)
    return str(val).strip()


def _parse_money(val):
    if val is None or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    txt = norm_cell(val).replace(',', '').replace("'", '')
    if not txt:
        return 0.0
    try:
        return float(txt)
    except ValueError:
        return 0.0


_DDMYYYY = (
    '%d/%m/%Y',
    '%d-%m-%Y',
    '%Y-%m-%d %H:%M:%S',
    '%Y-%m-%d',
)


def _parse_date_cell(val):
    if val is None or val == '':
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    s = norm_cell(val)
    for fmt in _DDMYYYY:
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_atlantida_xlsx(raw_bytes):
    wb = workbook_from_bytes(raw_bytes)
    try:
        sh = wb.active
        header_row = None
        lim = min(25, sh.max_row or 0)
        for r in range(1, lim + 1):
            b = norm_cell(sh.cell(r, 2).value).lower()
            g = norm_cell(sh.cell(r, 7).value).lower()
            if b == 'fecha' and ('débito' in g or 'debito' in g):
                header_row = r
                break
        if header_row is None:
            raise UserError(_('No se reconoce la tabla de movimientos (Banco Atlántida .xlsx).'))

        account_number = ''
        for r in range(1, header_row):
            if norm_cell(sh.cell(r, 1).value) == 'Cuenta:' and norm_cell(sh.cell(r, 2).value):
                account_number = norm_cell(sh.cell(r, 2).value)

        titular_hint = ''
        if norm_cell(sh.cell(1, 1).value):
            titular_hint = norm_cell(sh.cell(1, 1).value).replace('\n', ' ')[:256]

        currency = ''
        txt_all = ''
        for r in range(1, header_row):
            txt_all += norm_cell(sh.cell(r, 1).value) + norm_cell(sh.cell(r, 2).value)

        txt_u = txt_all.upper()
        if 'LEMPIRAS' in txt_u or 'HNL' in txt_u or 'LPS' in txt_u:
            currency = 'HNL'

        balance_start = 0.0
        transactions = []
        dates = []

        data_max = sh.max_row or header_row

        def row_amount(ridx):
            dbt = _parse_money(sh.cell(ridx, 7).value)
            crd = _parse_money(sh.cell(ridx, 8).value)
            return crd - dbt

        for r in range(header_row + 1, data_max + 1):
            date_str_try = _parse_date_cell(sh.cell(r, 2).value)
            if not date_str_try:
                continue
            row_desc = norm_cell(sh.cell(r, 4).value)

            amt = row_amount(r)
            if amt == 0.0 and 'saldo inicial' in row_desc.lower():
                balance_start = _parse_money(sh.cell(r, 9).value)
                continue
            desc = norm_cell(sh.cell(r, 4).value)
            ref_cell = norm_cell(sh.cell(r, 5).value)
            ref_ag = norm_cell(sh.cell(r, 6).value)

            amt = row_amount(r)
            if amt == 0.0 and not desc.strip():
                continue

            tail = ''
            if ref_ag and ref_ag.upper() != 'CASA MATRIZ':
                tail = ref_ag[:120]
            elif ref_cell and '/' not in ref_cell[:5]:
                tail = ''
            payment = desc
            if tail and tail.lower() != desc.lower()[: len(tail)]:
                payment = f'{desc} — {tail}'

            dates.append(date_str_try)
            transactions.append({
                'date': date_str_try,
                'amount': amt,
                'payment_ref': (payment or desc)[:1024],
                'ref': (ref_cell or '')[:256],
                'sequence': len(transactions) + 1,
            })

        balance_end = 0.0
        if transactions and sh.max_column >= 9:
            br = header_row + 1 + len(transactions) - 1
            if br <= data_max:
                balance_end = _parse_money(sh.cell(br, 9).value)
        elif data_max >= header_row + 2:
            balance_end = _parse_money(sh.cell(data_max, 9).value)

        total_amount = sum(t['amount'] for t in transactions)
        meta = {
            'min_date': min(dates) if dates else None,
            'max_date': max(dates) if dates else None,
            'total_amount': total_amount,
            'count': len(transactions),
            'balance_start': balance_start,
            'balance_end': balance_end if balance_end else (balance_start + total_amount),
            'titular': titular_hint[:256],
            'account_number': account_number[:64],
            'alias': '',
            'bank_name': _('Banco Atlántida'),
            'currency': currency or '',
        }
        return transactions, meta
    finally:
        wb.close()


def parse_banpais_xlsx(raw_bytes):
    wb = workbook_from_bytes(raw_bytes)
    try:
        sh = wb.active

        balance_start = 0.0
        balance_reported_end = None
        currency = ''

        mr = min(25, sh.max_row or 0)
        for r in range(1, mr + 1):
            lbl = norm_cell(sh.cell(r, 1).value)
            lbl_l = lbl.lower()
            cell_b = norm_cell(sh.cell(r, 2).value)
            if 'saldo inicial' in lbl_l and cell_b:
                balance_start = _parse_money(cell_b)
            if 'saldo final' in lbl_l and cell_b:
                balance_reported_end = _parse_money(cell_b)
            if lbl == 'Moneda:' and cell_b:
                cu = cell_b.upper()
                if cu in {'L', 'LPS'}:
                    currency = 'HNL'

        header_row = None
        lim = min(25, sh.max_row or 0)
        for r in range(1, lim + 1):
            if norm_cell(sh.cell(r, 1).value).upper() == 'FECHA' and norm_cell(sh.cell(r, 4).value):
                cand = norm_cell(sh.cell(r, 4).value).upper()
                if 'DESC' in cand:
                    header_row = r
                    break

        if header_row is None:
            raise UserError(_('No se reconoce la tabla de movimientos (Banpaís .xlsx).'))

        account_number = ''
        for r in range(1, header_row):
            if norm_cell(sh.cell(r, 1).value).startswith('Cuenta'):
                account_number = norm_cell(sh.cell(r, 2).value)

        company = ''
        for r in range(1, header_row):
            v1 = norm_cell(sh.cell(r, 1).value)
            if 'nombre de cuenta' in v1.lower() or 'nombre cuenta' in v1.lower():
                company = norm_cell(sh.cell(r, 2).value)

        transactions = []
        dates = []

        data_max = sh.max_row or header_row
        for r in range(header_row + 1, data_max + 1):
            d = _parse_date_cell(sh.cell(r, 1).value)
            if not d:
                continue

            trx = norm_cell(sh.cell(r, 2).value)
            ref = norm_cell(sh.cell(r, 3).value)
            desc = norm_cell(sh.cell(r, 4).value)
            if not desc and not ref and trx:
                continue
            debit = _parse_money(sh.cell(r, 5).value)
            credit = _parse_money(sh.cell(r, 6).value)

            amt = credit - debit
            dates.append(d)
            pay = desc
            if trx and trx not in desc[:12]:
                pay = f'{desc} ({trx})' if desc else trx
            transactions.append({
                'date': d,
                'amount': amt,
                'payment_ref': pay[:1024],
                'ref': (ref or '')[:256],
                'sequence': len(transactions) + 1,
            })

        total_amount = sum(t['amount'] for t in transactions)
        balance_end = balance_reported_end
        if balance_end is None and transactions:
            balance_end = _parse_money(sh.cell(header_row + len(transactions), 7).value)
        meta = {
            'min_date': min(dates) if dates else None,
            'max_date': max(dates) if dates else None,
            'total_amount': total_amount,
            'count': len(transactions),
            'balance_start': balance_start,
            'balance_end': balance_end if balance_end is not None else (balance_start + total_amount),
            'titular': company[:256],
            'account_number': account_number[:64],
            'alias': '',
            'bank_name': _('Banpaís'),
            'currency': currency or '',
        }
        return transactions, meta
    finally:
        wb.close()


def parse_occidente_hn_xlsx(raw_bytes):
    wb = workbook_from_bytes(raw_bytes)
    try:
        sh = wb.active

        header_row = None
        lim = min(25, sh.max_row or 0)
        for r in range(1, lim + 1):
            b = norm_cell(sh.cell(r, 2).value).lower()
            d16 = ''
            if sh.max_column >= 16:
                d16 = norm_cell(sh.cell(r, 16).value).lower()
            if b == 'fecha' and ('bit' in d16):  # Débito / Debito
                header_row = r
                break
        if header_row is None:
            raise UserError(_('No se reconoce la tabla de movimientos (Occidente HN .xlsx).'))

        account_number = ''
        company_hint = ''
        balance_start = 0.0
        balance_claimed_sheet = None
        currency = ''

        for r in range(1, header_row):
            lbl = norm_cell(sh.cell(r, 2).value).lower().strip()
            col_d = ''
            if sh.max_column >= 4:
                col_d = norm_cell(sh.cell(r, 4).value)

            if lbl == 'cuenta:' and col_d:
                account_number = col_d
            if lbl == 'moneda:' and col_d.upper() in {'LPS', 'LEMP', 'HN'}:
                currency = 'HNL'

        for r in range(6, min(header_row, 11)):
            l12 = ''
            if sh.max_column >= 12:
                l12 = norm_cell(sh.cell(r, 12).value).lower()
            if 'saldo inicial' in l12 and sh.max_column >= 14:
                balance_start = _parse_money(sh.cell(r, 14).value)
            q17 = ''
            if sh.max_column >= 17:
                q17 = norm_cell(sh.cell(r, 17).value).lower()
            if ('saldo final' in l12 or 'saldo final' in q17) and sh.max_column >= 19:
                balance_claimed_sheet = _parse_money(sh.cell(r, 19).value)

        if sh.max_column >= 12 and norm_cell(sh.cell(6, 12).value):
            company_hint = norm_cell(sh.cell(6, 12).value)

        transactions = []
        dates = []
        running_saldo = []

        last_data_row = header_row
        max_r = sh.max_row or header_row

        for r in range(header_row + 1, max_r + 1):
            d = _parse_date_cell(sh.cell(r, 2).value)
            if not d:
                continue
            desc = norm_cell(sh.cell(r, 11).value)
            ref_cell = ''
            if sh.max_column >= 10:
                ref_cell = norm_cell(sh.cell(r, 10).value)
            debit = _parse_money(sh.cell(r, 16).value) if sh.max_column >= 16 else 0.0
            credit = _parse_money(sh.cell(r, 18).value) if sh.max_column >= 18 else 0.0
            saldo_running = None
            if sh.max_column >= 20:
                sr = sh.cell(r, 20).value
                if sr not in (None, ''):
                    saldo_running = _parse_money(sr)

            amt = credit - debit
            if amt == 0.0 and not desc and not ref_cell:
                continue

            dates.append(d)
            if saldo_running is not None:
                running_saldo.append(saldo_running)
            transactions.append({
                'date': d,
                'amount': amt,
                'payment_ref': desc[:1024],
                'ref': (ref_cell or '')[:256],
                'sequence': len(transactions) + 1,
            })
            last_data_row = r

        total_amount = sum(t['amount'] for t in transactions)
        balance_end = running_saldo[-1] if running_saldo else (balance_claimed_sheet or 0.0)
        meta = {
            'min_date': min(dates) if dates else None,
            'max_date': max(dates) if dates else None,
            'total_amount': total_amount,
            'count': len(transactions),
            'balance_start': balance_start,
            'balance_end': balance_end,
            'titular': company_hint[:256],
            'account_number': account_number[:64],
            'alias': '',
            'bank_name': _('Banco de Occidente'),
            'currency': currency or 'HNL',
        }
        return transactions, meta
    finally:
        wb.close()


def parse_xlsx_by_format(raw_bytes, fmt_key):
    if fmt_key == 'xlsx_atlantida':
        return parse_atlantida_xlsx(raw_bytes)
    if fmt_key == 'xlsx_banpais':
        return parse_banpais_xlsx(raw_bytes)
    if fmt_key == 'xlsx_occidente_hn':
        return parse_occidente_hn_xlsx(raw_bytes)
    raise UserError(_('Formato .xlsx no soportado: %s') % fmt_key)


def bank_label_for_format(fmt_key):
    return {
        'xlsx_atlantida': _('Banco Atlántida'),
        'xlsx_banpais': _('Banpaís'),
        'xlsx_occidente_hn': _('Occidente (HN) Excel'),
        'xls_ficohsa': _('FICOHSA'),
        'xls_davivienda_hn': _('Davivienda Honduras'),
        'pdf_occident': _('PDF Occidente HN'),
    }.get(fmt_key, _('Extracto'))
