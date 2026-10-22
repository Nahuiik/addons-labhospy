# -*- coding: utf-8 -*-
import base64
import binascii
import hashlib
import io
import unicodedata
from collections import defaultdict
from datetime import date, datetime

from odoo import Command, _, api, fields, models
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - dependencia declarada en el manifest
    load_workbook = None


def _normalize_account_code(code):
    if not code:
        return ""
    return "".join(c for c in str(code) if c.isalnum())


def _norm(value):
    """Normaliza texto de cabecera: minúsculas, sin acentos, espacios colapsados."""
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    return " ".join(s.split())


def _binary_data_to_bytes(data):
    """Convierte un valor Binary de Odoo en bytes crudos del fichero."""
    if not data:
        return b""
    if isinstance(data, bytes):
        try:
            return base64.b64decode(data)
        except Exception:
            return data
    s = data if isinstance(data, str) else data.decode()
    try:
        return base64.b64decode(s.encode("ascii"))
    except (ValueError, binascii.Error):
        return base64.b64decode(s)


def _to_float(cell):
    """Importes de FICOHSA: pueden venir como número (openpyxl) o como texto."""
    if cell is None:
        return 0.0
    if isinstance(cell, (int, float)):
        return float(cell)
    s = str(cell).strip().replace(",", "").replace('"', "")
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_ficohsa_date(cell):
    """Fecha de Pago del reporte FICOHSA (formato dd-mm-yyyy o celda fecha de Excel)."""
    if cell is None:
        return None
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    s = str(cell).strip()
    if not s:
        return None
    if " " in s:
        s = s.split()[0]
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


class FicohsaPosImportWizard(models.TransientModel):
    _name = "ficohsa.pos.import.wizard"
    _description = "Importar liquidaciones FICOHSA POS desde Excel"

    # Cabeceras esperadas (normalizadas) -> clave interna. Se admiten variantes.
    _HEADER_CANDIDATES = {
        "pay_date": ("fecha de pago",),
        "net": ("monto neto liquidacion",),
        "commission": ("comision",),
        "isv": ("isv retenido (-)", "isv retenido"),
        "total": ("monto total ventas",),
    }

    import_mode = fields.Selection(
        [
            ("single", "Un solo archivo"),
            ("batch", "Por lotes (varios Excel)"),
        ],
        string="Modo de importación",
        default="single",
        required=True,
        help="Por cada archivo Excel: las transacciones se agrupan por la Fecha de Pago y se "
        "genera una liquidación por cada día con importe dentro de ese archivo. En modo lote, "
        "cada archivo se traza aparte: no se mezclan importes entre archivos.",
    )

    attachment_ids = fields.Many2many(
        "ir.attachment",
        "ficohsa_pos_import_wizard_attachment_rel",
        "wizard_id",
        "attachment_id",
        string="Archivos Excel (lote)",
        help="Suba varios archivos a la vez; no es necesario agregar fila por fila.",
    )

    bank_journal_id = fields.Many2one(
        "account.journal",
        string="Diario banco depósito",
        domain=[("type", "=", "bank")],
        help="Seleccione el diario del banco donde se acreditan los depósitos netos de FICOHSA.",
    )

    liquidation_journal_id = fields.Many2one(
        "account.journal",
        string="Diario a liquidar (POS / TC)",
        required=True,
        domain=[("type", "in", ["bank", "cash"])],
        help="Siempre TC POS.",
    )

    commission_account_id = fields.Many2one(
        "account.account",
        string="Cuenta comisión POS",
        compute="_compute_commission_account_id",
        readonly=True,
    )

    data_file = fields.Binary(string="Archivo FICOHSA (Excel)")
    file_name = fields.Char(string="Nombre archivo")

    progress = fields.Float(string="Progreso %", digits=(16, 1), default=0.0)
    result_message = fields.Text(string="Resultado", readonly=True)
    # Resumen estructurado para el panel OWL de resultados.
    result_data = fields.Json(string="Datos resultado", readonly=True)

    processed_days = fields.Integer(default=0)
    total_days = fields.Integer(default=0)

    state = fields.Selection(
        [("draft", "Pendiente"), ("running", "Procesando"), ("done", "Terminado")],
        default="draft",
    )

    # Nombres del diario de liquidación POS de FICOHSA (en orden de preferencia).
    _LIQUIDATION_JOURNAL_NAMES = ("POS FICOHSA", "TC POS")

    @api.model
    def _tc_pos_journal(self):
        """Diario de liquidación POS del adquirente FICOHSA."""
        company = self.env.company
        Journal = self.env["account.journal"].with_company(company)
        for name in self._LIQUIDATION_JOURNAL_NAMES:
            for op in ("=", "ilike"):
                j = Journal.search(
                    [("name", op, name), ("type", "in", ("bank", "cash")),
                     ("company_id", "=", company.id)], limit=1)
                if j:
                    return j
        return Journal.browse()

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        lj = self._tc_pos_journal()
        if lj:
            res["liquidation_journal_id"] = lj.id
        return res

    @api.model
    def _find_account(self, codes, needle, names):
        """Resuelve una cuenta por código (variantes), código normalizado o nombre."""
        company = self.env.company
        Account = self.env["account.account"].with_company(company)
        # Odoo 19: account.account es multi-compañía (company_ids) y ya no existe
        # el campo 'deprecated'.
        base = [("company_ids", "in", company.id)]
        for code in codes:
            acc = Account.search(base + [("code", "=", code)], limit=1)
            if acc:
                return acc
        for acc in Account.search(base):
            if _normalize_account_code(acc.code) == needle:
                return acc
        for name_tpl in names:
            acc = Account.search(base + [("name", "ilike", name_tpl)], limit=1)
            if acc:
                return acc
        return self.env["account.account"]

    @api.model
    def _resolve_commission_account(self):
        """Cuenta comisión POS: 7010101.02 «Comisiones Cobro POS»."""
        return self._find_account(
            ("7010101.02", "7010101-02", "701010102"),
            "701010102",
            ("Comisiones Cobro POS", "Comisión Cobro POS", "Comision Cobro POS"),
        )

    @api.model
    def _resolve_isv_account(self):
        """Cuenta ISV/retención POS: 1030101.10 «Retención Por Cobros TC»."""
        return self._find_account(
            ("1030101.10", "1030101-10", "103010110"),
            "103010110",
            ("Retención Por Cobros TC", "Retencion Por Cobros TC"),
        )

    @api.depends("import_mode")
    def _compute_commission_account_id(self):
        for wiz in self:
            acc = wiz._resolve_commission_account()
            wiz.commission_account_id = acc if acc else False

    @api.onchange("import_mode")
    def _onchange_import_mode(self):
        if self.import_mode == "single":
            self.attachment_ids = [Command.clear()]
        else:
            self.data_file = False
            self.file_name = False

    def _bytes_from_attachment(self, attachment):
        """Lee fichero crudo; evita bin_size que sustituye datas por texto de tamaño."""
        att = attachment.with_context(bin_size=False)
        raw = getattr(att, "raw", None)
        if raw:
            return raw
        return _binary_data_to_bytes(att.datas)

    def _open_worksheet(self, raw):
        """Devuelve la primera hoja del libro Excel a partir de los bytes crudos."""
        if load_workbook is None:
            raise UserError(
                _(
                    "Falta la librería de Python «openpyxl» para leer archivos Excel. "
                    "Instálela en el servidor: pip install openpyxl."
                )
            )
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001 - mensaje claro al usuario
            raise UserError(
                _(
                    "No se pudo abrir el archivo como Excel (.xlsx). Verifique que sea el "
                    "reporte «Cierre por Terminal» de FICOHSA. Detalle: %s"
                )
                % (exc,)
            )
        return wb.worksheets[0] if wb.worksheets else None

    def _locate_columns(self, rows):
        """Encuentra la fila de cabecera y mapea claves internas a índices de columna."""
        for header_idx, row in enumerate(rows):
            norms = {}
            for idx, cell in enumerate(row):
                key = _norm(cell)
                if key:
                    norms[key] = idx
            if "fecha de pago" in norms and "monto neto liquidacion" in norms:
                colmap = {}
                for key, candidates in self._HEADER_CANDIDATES.items():
                    for cand in candidates:
                        if cand in norms:
                            colmap[key] = norms[cand]
                            break
                return header_idx, colmap
        return None, {}

    def _aggregate_xlsx_into(self, by_date, raw):
        """Acumula importes por Fecha de Pago dentro de un único Excel."""
        ws = self._open_worksheet(raw)
        if ws is None:
            return
        rows = list(ws.iter_rows(values_only=True))
        header_idx, colmap = self._locate_columns(rows)
        if header_idx is None or "pay_date" not in colmap or "net" not in colmap:
            raise UserError(
                _(
                    "No se encontraron las columnas esperadas (Fecha de Pago, Monto Neto "
                    "Liquidacion, Comision) en el Excel de FICOHSA."
                )
            )
        c_date = colmap["pay_date"]
        c_net = colmap["net"]
        c_com = colmap.get("commission")
        c_isv = colmap.get("isv")
        max_col = max(c for c in colmap.values() if c is not None)
        for row in rows[header_idx + 1:]:
            if not row or all(c in (None, "") for c in row):
                continue
            if any(_norm(c) in ("suma:", "suma", "total:", "totales") for c in row):
                break
            if len(row) <= max_col:
                continue
            fecha = _parse_ficohsa_date(row[c_date])
            if not fecha:
                continue
            agg = by_date[fecha]
            agg["bank_net"] += _to_float(row[c_net])
            if c_com is not None:
                agg["commission"] += _to_float(row[c_com])
            if c_isv is not None:
                agg["isv"] += _to_float(row[c_isv])

    def _liquidation_description(self, fdate, entry):
        """Texto único por (archivo, día) para el documento y la comprobación de duplicados."""
        nm = (entry.get("name") or _("archivo.xlsx")).strip() or _("archivo.xlsx")
        if len(nm) > 120:
            nm = nm[:117] + "..."
        return "LIQUIDACION POS FICOHSA %s — %s [%s]" % (
            fdate.strftime("%d/%m/%Y"),
            nm,
            entry["hash"][:8],
        )

    def _get_commission_account(self):
        acc = self._resolve_commission_account()
        if not acc:
            raise UserError(
                _(
                    "No se encontró la cuenta de comisión POS (código 7010101.02 "
                    "«Comisiones Cobro POS») en la compañía actual."
                )
            )
        return acc

    def _get_isv_account(self):
        acc = self._resolve_isv_account()
        if not acc:
            raise UserError(
                _(
                    "No se encontró la cuenta de ISV/retención POS (código 1030101.10 "
                    "«Retención Por Cobros TC») en la compañía actual."
                )
            )
        return acc

    def action_import(self):
        self.ensure_one()
        if not self.bank_journal_id:
            raise UserError(_("Seleccione el diario banco depósito."))

        lj = self._tc_pos_journal()
        if not lj:
            raise UserError(
                _("No existe un diario llamado «TC POS» (tipo banco o caja) en esta compañía.")
            )
        self.liquidation_journal_id = lj

        if self.import_mode == "single":
            if not self.data_file:
                raise UserError(_("Selecciona un archivo Excel."))
            raw = _binary_data_to_bytes(self.with_context(bin_size=False).data_file)
            if not raw:
                raise UserError(_("El archivo Excel está vacío o no se pudo leer."))
            files_payload = [(self.file_name or _("archivo.xlsx"), raw)]
        else:
            atts = self.attachment_ids
            if not atts:
                raise UserError(
                    _("En modo por lotes, suba uno o más archivos Excel (área de adjuntos).")
                )
            files_payload = []
            for att in atts:
                nm = att.name or _("lote.xlsx")
                raw = self._bytes_from_attachment(att)
                if not raw:
                    raise UserError(_("El adjunto «%s» está vacío o no se pudo leer.") % (nm,))
                files_payload.append((nm, raw))

        file_entries = []
        for fname, raw in files_payload:
            file_hash = hashlib.sha256(raw).hexdigest()
            file_entries.append({"name": fname, "raw": raw, "hash": file_hash})

        hashes_batch = [e["hash"] for e in file_entries]
        if len(hashes_batch) != len(set(hashes_batch)):
            raise UserError(
                _("Hay el mismo archivo más de una vez en este lote (contenido idéntico).")
            )

        Log = self.env["ficohsa.pos.import.file.log"]
        company = self.env.company
        existing_hashes = set(
            Log.search(
                [("company_id", "=", company.id), ("file_hash", "in", list(set(hashes_batch)))]
            ).mapped("file_hash")
        )
        dup_names = [e["name"] for e in file_entries if e["hash"] in existing_hashes]
        if dup_names:
            raise UserError(
                _(
                    "Este contenido de archivo Excel ya fue importado antes (huella digital "
                    "idéntica). Archivos rechazados: %(files)s"
                )
                % {"files": ", ".join(dup_names)}
            )

        work_units = []
        skipped_empty_files = []
        for entry in file_entries:
            by_date = defaultdict(lambda: {"bank_net": 0.0, "commission": 0.0, "isv": 0.0})
            self._aggregate_xlsx_into(by_date, entry["raw"])
            if not by_date:
                skipped_empty_files.append(entry["name"])
                continue
            for fdate in sorted(by_date.keys()):
                agg = by_date[fdate]
                if round(agg["bank_net"] + agg["commission"] + agg["isv"], 2) <= 0:
                    continue
                work_units.append((entry, fdate, agg))

        if not work_units:
            if skipped_empty_files and len(skipped_empty_files) == len(file_entries):
                raise UserError(
                    _("Ningún archivo contenía movimientos de detalle válidos: %(files)s")
                    % {"files": ", ".join(skipped_empty_files)}
                )
            raise UserError(
                _("No se encontraron movimientos de detalle en los Excel (modo %(mode)s).")
                % {"mode": _("lote") if self.import_mode == "batch" else _("archivo único")}
            )

        bank_acc = self.bank_journal_id.default_account_id
        if not bank_acc:
            raise UserError(
                _("Configure la cuenta por defecto (liquidez) en el diario banco %s.")
                % self.bank_journal_id.display_name
            )

        needs_commission = any(round(u[2]["commission"], 2) > 0 for u in work_units)
        needs_isv = any(round(u[2]["isv"], 2) > 0 for u in work_units)
        commission_acc = self._get_commission_account() if needs_commission else None
        isv_acc = self._get_isv_account() if needs_isv else None

        DebitCredit = self.env["debit.credit"]
        created = self.env["debit.credit"]

        total = len(work_units)
        self.write(
            {
                "state": "running",
                "total_days": total,
                "processed_days": 0,
                "progress": 0.0,
                "result_message": False,
            }
        )

        use_web_progress = callable(getattr(self, "web_progress_iter", None))
        iter_units = work_units
        if use_web_progress:
            iter_units = self.web_progress_iter(
                work_units,
                total=total,
                msg=_("Importando liquidaciones FICOHSA POS…"),
                cancellable=True,
            )

        for idx, unit in enumerate(iter_units, start=1):
            entry, fdate, agg = unit
            bank_amt = round(agg["bank_net"], 2)
            commission_amt = round(agg["commission"], 2)
            isv_amt = round(agg["isv"], 2)
            grand_total = round(bank_amt + commission_amt + isv_amt, 2)

            if not use_web_progress:
                pct = round(100.0 * idx / total, 1) if total else 100.0
                self.write({"progress": pct, "processed_days": idx})

            description = self._liquidation_description(fdate, entry)

            dup = DebitCredit.search(
                [
                    ("description", "=", description),
                    ("date", "=", fields.Date.to_string(fdate)),
                    ("journal_id", "=", lj.id),
                    ("move_type", "=", "credit_card"),
                    ("company_id", "=", self.env.company.id),
                    ("state", "in", ("draft", "validated")),
                ],
                limit=1,
            )
            if dup:
                continue

            debit_cmds = [
                Command.create(
                    {
                        "account_id": bank_acc.id,
                        "amount": bank_amt,
                        "type": "debit",
                        "name": _("LIQUIDACION DE POS"),
                    }
                ),
            ]
            if commission_amt:
                debit_cmds.append(
                    Command.create(
                        {
                            "account_id": commission_acc.id,
                            "amount": commission_amt,
                            "type": "debit",
                            "name": _("COMISION DE POS"),
                        }
                    )
                )
            if isv_amt:
                debit_cmds.append(
                    Command.create(
                        {
                            "account_id": isv_acc.id,
                            "amount": isv_amt,
                            "type": "debit",
                            "name": _("ISV RETENIDO POS"),
                        }
                    )
                )

            dc = DebitCredit.create(
                {
                    "journal_id": lj.id,
                    "date": fdate,
                    "total": grand_total,
                    "description": description,
                    "move_type": "credit_card",
                    "currency_id": self.env.company.currency_id.id,
                    "company_id": self.env.company.id,
                    "currency_rate": 1.0,
                    "debit_line_ids": debit_cmds,
                }
            )

            dc.action_validate()
            created |= dc

        nfiles = len(file_entries)
        msg_txt = _(
            "Importación finalizada (%(mode)s, %(nfiles)s archivo(s)). "
            "Por archivo: una liquidación por cada Fecha de Pago con importe. "
            "Nuevos asientos: %(cnt)s (operaciones día/archivo previstas: %(ops)s)."
        ) % {
            "mode": _("lote") if self.import_mode == "batch" else _("un archivo"),
            "nfiles": nfiles,
            "cnt": len(created),
            "ops": len(work_units),
        }
        if skipped_empty_files:
            msg_txt += "\n" + _(
                "Archivos sin líneas de detalle importables: %(files)s"
            ) % {"files": ", ".join(skipped_empty_files)}

        files_summary = "\n".join(e["name"] for e in file_entries)
        hist = self.env["ficohsa.pos.import.history"].create_from_wizard(
            self.import_mode,
            self.bank_journal_id,
            file_entries,
            len(created),
            len(work_units),
            msg_txt,
            files_summary,
        )

        Log.create(
            [
                {
                    "company_id": company.id,
                    "file_hash": e["hash"],
                    "file_name": e["name"],
                    "history_id": hist.id,
                }
                for e in file_entries
            ]
        )

        # Resumen estructurado para el componente OWL.
        result_data = {
            "mode": "batch" if self.import_mode == "batch" else "single",
            "created": len(created),
            "planned": len(work_units),
            "files": [e["name"] for e in file_entries],
            "skipped": skipped_empty_files,
            "bank_journal": self.bank_journal_id.display_name,
        }
        self.write(
            {
                "result_message": msg_txt,
                "result_data": result_data,
                "state": "done",
                "progress": 100.0,
            }
        )

        if len(created) == 1:
            return {
                "type": "ir.actions.act_window",
                "res_model": "debit.credit",
                "view_mode": "form",
                "target": "current",
                "res_id": created.id,
                "views": [(False, "form")],
            }
        if created:
            return {
                "type": "ir.actions.act_window",
                "name": _("Liquidaciones FICOHSA importadas"),
                "res_model": "debit.credit",
                "view_mode": "list,form",
                "domain": [("id", "in", created.ids)],
            }
        return {
            "type": "ir.actions.act_window",
            "res_model": "ficohsa.pos.import.wizard",
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
            "views": [(False, "form")],
        }
